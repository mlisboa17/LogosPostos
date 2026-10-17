"""Ponto de composição do robô noturno dos módulos de caixa e placar."""

from __future__ import annotations

import argparse
import asyncio
import json
import logging
import os
from datetime import date
from decimal import Decimal
from pathlib import Path
from tempfile import NamedTemporaryFile

from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font

from src.modules.cash_reconciliation.adapters import pendencias as repositorio_pendencias
from src.modules.cash_reconciliation.adapters import persistencia as persistencia_caixa
from src.modules.cash_reconciliation.adapters import saude_robo
from src.modules.cash_reconciliation.application.resumo_diario import (
    abas_conferencia_recebimentos, conferencia_para_registro, gerar_mensagem_curta,
    identificar_caixa,
)
from src.modules.cash_reconciliation.config import carregar_unidades, prazos_pendencia_dias
from src.modules.cash_reconciliation.domain.tempo import agora, formatar_data, ontem, para_local
from src.modules.cash_reconciliation.domain.execucao import ResultadoDiario
from src.modules.cash_reconciliation.domain.integridade_execucao import ExecucaoNoturna, FonteExecucao
from src.modules.cash_reconciliation.jobs.noturno import (
    FormatoLogBrasil,
    _dia,
    executar,
    validar_registros,
    registros_para_publicacao,
)
from src.modules.commercial_performance.adapters.persistencia import salvar_placar
from src.modules.commercial_performance.application.placar import obter_placar
from src.modules.commercial_performance.config import POSTOS
from src.modules.webposto_integration.funcionarios import buscar_nomes_funcionarios
from src.modules.webposto_integration.http import WebPostoErro
from src.shared.notificacao.telegram import (
    ResultadoDestinatario,
    enviar_documento,
    enviar_mensagem,
)

logger = logging.getLogger("src.modules.cash_reconciliation.jobs.noturno")
DIRETORIO_RESUMOS = Path(__file__).resolve().parents[2] / "data" / "resumos"
CAMINHO_ENV = Path(__file__).resolve().parents[2] / ".env"


def _ler_envios_telegram(caminho: Path) -> dict:
    try:
        conteudo = caminho.read_text(encoding="utf-8")
    except FileNotFoundError:
        return {}
    except OSError as exc:
        raise OSError(type(exc).__name__) from None
    try:
        dados = json.loads(conteudo)
    except json.JSONDecodeError as exc:
        raise ValueError(type(exc).__name__) from None
    if not isinstance(dados, dict):
        raise ValueError("Registro de envios inválido.")
    return dados


def _registrar_envio_telegram(
    caminho: Path, dia: date, sucesso: bool, execucao: ExecucaoNoturna,
) -> None:
    dados = _ler_envios_telegram(caminho)
    dados[dia.isoformat()] = {
        "registrado_em": agora().isoformat(),
        "situacao": "ok" if sucesso else "falhou",
        "run_id": execucao.run_id,
        "audit_status": execucao.audit_status,
        "delivery_status": execucao.delivery_status,
        "document_delivery_status": execucao.document_delivery_status,
    }
    temporario: Path | None = None
    try:
        caminho.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=caminho.parent,
            suffix=".tmp",
            delete=False,
        ) as arquivo:
            temporario = Path(arquivo.name)
            json.dump(dados, arquivo, ensure_ascii=False, indent=2)
            arquivo.write("\n")
            arquivo.flush()
            os.fsync(arquivo.fileno())
        os.replace(temporario, caminho)
    finally:
        if temporario is not None and temporario.exists():
            temporario.unlink()


def _coletar_dados_telegram(
    dia: date, resultados: list[ResultadoDiario], execucao: ExecucaoNoturna,
) -> dict:
    unidades = carregar_unidades()
    validar_registros(execucao, resultados)
    if dia != execucao.dia or not execucao.permite_publicacao():
        raise ValueError("Auditoria atual incompleta; composição financeira bloqueada.")
    resultados = registros_para_publicacao(execucao, resultados)
    unidades_com_falha: set[int] = set()
    abertas = {
        unidade: repositorio_pendencias.contar_abertas_somente_leitura(unidade=unidade)
        for unidade in unidades
    }
    prazos = prazos_pendencia_dias()
    vencidas = {
        unidade: repositorio_pendencias.contar_vencidas_somente_leitura(
            unidade=unidade,
            prazos_dias=prazos,
        )
        for unidade in unidades
    }
    codigos_por_unidade = {}
    for codigo, resultado in ((item.empresa_codigo, item) for item in resultados):
        if resultado.fechamento is None:
            continue
        codigos = {
            caixa.caixa.funcionario_codigo
            for caixa in resultado.fechamento.caixas
            if caixa.caixa.funcionario_codigo is not None
        }
        codigos.update(
            vale.funcionario_codigo
            for caixa in resultado.fechamento.caixas
            for vale in caixa.vales
            if vale.funcionario_codigo is not None
        )
        codigos_por_unidade[codigo] = codigos

    async def consultar_nomes() -> dict[int, dict[int, str]]:
        async def consultar_unidade(codigo: int, codigos: set[int]) -> tuple[int, dict[int, str]]:
            if not codigos:
                return codigo, {}
            try:
                nomes = await buscar_nomes_funcionarios(unidades[codigo], dia, dia)
            except WebPostoErro as exc:
                logger.warning(
                    "Telegram nomes indisponiveis unidade=%s motivo=%s",
                    codigo,
                    type(exc).__name__,
                )
                return codigo, {}
            return codigo, {chave: nome for chave, nome in nomes.items() if chave in codigos}

        encontrados = await asyncio.gather(
            *(consultar_unidade(codigo, codigos) for codigo, codigos in codigos_por_unidade.items())
        )
        return dict(encontrados)

    nomes_funcionarios = asyncio.run(consultar_nomes())
    return {
        "unidades": unidades,
        "resultados": resultados,
        "unidades_com_falha": unidades_com_falha,
        "abertas": abertas,
        "vencidas": vencidas,
        "nomes_funcionarios": nomes_funcionarios,
    }


def _mensagens_telegram(dia: date, dados: dict) -> list[str]:
    unidades = dados["unidades"]
    resultados = dados["resultados"]
    unidades_com_falha = dados["unidades_com_falha"]
    abertas = dados["abertas"]
    vencidas = dados["vencidas"]
    return gerar_mensagem_curta(
        dia,
        resultados,
        list(unidades),
        unidades_com_falha,
        abertas,
        vencidas,
        dados["nomes_funcionarios"],
    )


def _criar_planilha_telegram(dia: date, dados: dict, diretorio: Path | None = None) -> Path:
    destino = diretorio or DIRETORIO_RESUMOS
    destino.mkdir(parents=True, exist_ok=True)
    caminho = destino / f"{dia.isoformat()}.xlsx"
    workbook = Workbook()
    resumo = workbook.active
    resumo.title = "Resumo"
    caixas_sheet = workbook.create_sheet("Caixas")
    despesas_sheet = workbook.create_sheet("Despesas")
    vales_sheet = workbook.create_sheet("Vales")

    cabecalhos = {
        "Resumo": [
            "Filial",
            "Caixas",
            "Quebra",
            "Faltas sem desconto (R$)",
            "Despesas (qtd)",
            "Despesas (R$)",
            "Vales (qtd)",
            "Vales (R$)",
            "Pendências abertas",
            "Pendências vencidas",
        ],
        "Caixas": [
            "Filial",
            "Caixa",
            "Operador",
            "Abertura",
            "Fechamento",
            "Consolidado",
            "Quebra",
            "Falta sem desconto",
            "Código interno",
        ],
        "Despesas": [
            "Filial",
            "Caixa",
            "Valor",
            "Descrição",
            "Plano de contas",
            "Alerta",
            "Código interno",
        ],
        "Vales": ["Filial", "Caixa", "Funcionário", "Valor", "Origem", "Código interno"],
    }
    for planilha, cabecalho in zip(
        (resumo, caixas_sheet, despesas_sheet, vales_sheet),
        cabecalhos.values(),
    ):
        planilha.append(cabecalho)

    por_unidade = {resultado.empresa_codigo: resultado for resultado in dados["resultados"]}
    conferencias = [conferencia_para_registro(r) for r in dados["resultados"] if r.conferencia_recebimentos]
    if conferencias:
        abas_conferencia_recebimentos(workbook, conferencias, dados["nomes_funcionarios"])
    for codigo, unidade in dados["unidades"].items():
        registro = por_unidade.get(codigo)
        fechamento = registro.fechamento if registro else None
        auditorias = list(fechamento.caixas) if fechamento else []
        apelido = unidade.apelido or unidade.nome
        faltas = [
            alerta
            for auditoria in auditorias
            for alerta in auditoria.alertas
            if alerta.codigo == "FALTA_SEM_DESCONTO"
        ]
        despesas = [despesa for auditoria in auditorias for despesa in auditoria.despesas]
        vales = [vale for auditoria in auditorias for vale in auditoria.vales]
        resumo.append(
            [
                apelido,
                len(auditorias),
                float(fechamento.quebra_total) if fechamento else 0,
                float(sum((abs(item.valor or Decimal(0)) for item in faltas), Decimal(0))),
                len(despesas),
                float(sum((item.valor for item in despesas), Decimal(0))),
                len(vales),
                float(sum((item.valor for item in vales), Decimal(0))),
                dados["abertas"].get(codigo),
                dados["vencidas"].get(codigo),
            ]
        )
        nomes = dados["nomes_funcionarios"].get(codigo, {})
        for auditoria in auditorias:
            caixa = auditoria.caixa
            falta_caixa = sum(
                (
                    abs(alerta.valor or Decimal(0))
                    for alerta in auditoria.alertas
                    if alerta.codigo == "FALTA_SEM_DESCONTO" and alerta.referencia == caixa.codigo
                ),
                Decimal(0),
            )
            abertura = para_local(caixa.abertura).replace(tzinfo=None)
            fechamento_caixa = (
                para_local(caixa.fechamento).replace(tzinfo=None)
                if caixa.fechamento is not None
                else None
            )
            rotulo_caixa = identificar_caixa(apelido, caixa, [a.caixa for a in auditorias], nomes)
            caixas_sheet.append(
                [
                    apelido,
                    rotulo_caixa,
                    nomes.get(caixa.funcionario_codigo) or f"funcionário {caixa.funcionario_codigo}",
                    abertura,
                    fechamento_caixa,
                    "sim" if caixa.consolidado else "não",
                    float(auditoria.quebra),
                    float(falta_caixa),
                    caixa.codigo,
                ]
            )
            for despesa in auditoria.despesas:
                plano = " ".join(
                    parte
                    for parte in (despesa.plano_conta_codigo, despesa.plano_conta_nome)
                    if parte
                )
                despesas_sheet.append(
                    [
                        apelido,
                        rotulo_caixa,
                        float(despesa.valor),
                        despesa.descricao or "",
                        plano,
                        "SEM PLANO" if not despesa.plano_conta_codigo else "",
                        caixa.codigo,
                    ]
                )
            rotulos_origem = {"C": "sobra de caixa", "D": "falta de caixa"}
            for vale in auditoria.vales:
                codigo_funcionario = vale.funcionario_codigo
                vales_sheet.append(
                    [
                        apelido,
                        rotulo_caixa,
                        nomes.get(
                            codigo_funcionario,
                            (
                                f"func. {codigo_funcionario}"
                                if codigo_funcionario is not None
                                else "indisponível"
                            ),
                        ),
                        float(vale.valor),
                        rotulos_origem.get(vale.origem.upper(), "vale"),
                        caixa.codigo,
                    ]
                )

    formatos_monetarios = {
        "Resumo": ("C", "D", "F", "H"),
        "Caixas": ("G", "H"),
        "Despesas": ("C",),
        "Vales": ("D",),
        "Ações de recebimento": ("D",),
        "Correções prontas": ("F",),
        "Detalhes por pagamento": ("E", "N", "R"),
    }
    for planilha in workbook.worksheets:
        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions
        for celula in planilha[1]:
            celula.font = Font(bold=True)
        colunas_monetarias: tuple[str | int, ...] = formatos_monetarios.get(planilha.title, ())
        if planilha.title.startswith("Dia "):
            colunas_monetarias = tuple(range(2, planilha.max_column + 1))
        elif planilha.title.startswith("Formas "):
            colunas_monetarias = (4, 6, 7)
        for coluna in colunas_monetarias:
            for linha in range(2, planilha.max_row + 1):
                celula = planilha.cell(linha, coluna) if isinstance(coluna, int) else planilha[f"{coluna}{linha}"]
                celula.number_format = '"R$" #,##0.00'
        if planilha.title == "Caixas":
            for linha in range(2, planilha.max_row + 1):
                planilha[f"D{linha}"].number_format = "dd/mm/yyyy hh:mm"
                planilha[f"E{linha}"].number_format = "dd/mm/yyyy hh:mm"
        for coluna in planilha.columns:
            letra = coluna[0].column_letter
            largura = max((len(str(celula.value or "")) for celula in coluna), default=0)
            planilha.column_dimensions[letra].width = min(max(largura + 2, 12), 60)

    workbook.save(caminho)
    return caminho


def _artefatos_telegram(
    dia: date, resultados: list[ResultadoDiario], execucao: ExecucaoNoturna,
    diretorio: Path | None = None,
) -> tuple[list[str], Path]:
    dados = _coletar_dados_telegram(dia, resultados, execucao)
    mensagens = _mensagens_telegram(dia, dados)
    planilha = _criar_planilha_telegram(dia, dados, diretorio)
    return mensagens, planilha


def _tentar_envio_telegram(
    dia: date,
    *,
    dia_manual: bool,
    enviar_manual: bool,
    reenviar: bool,
    caminho_registro: Path | None = None,
    execucao: ExecucaoNoturna,
    resultados: list[ResultadoDiario],
) -> None:
    validar_registros(execucao, resultados)
    if os.getenv("TELEGRAM_ENVIO", "0").strip() != "1":
        execucao.delivery_status = "SKIPPED"
        return
    if dia != ontem() and not (dia_manual and enviar_manual):
        execucao.delivery_status = "SKIPPED"
        return
    if dia_manual and not enviar_manual:
        execucao.delivery_status = "SKIPPED"
        return

    registro = caminho_registro or DIRETORIO_RESUMOS / "envios_telegram.json"
    try:
        envios = _ler_envios_telegram(registro)
        if dia.isoformat() in envios and not reenviar:
            logger.info("Telegram envio=ignorado motivo=ja_enviado")
            execucao.delivery_status = "SKIPPED"
            return
        token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
        chat_ids = list(
            dict.fromkeys(
                item.strip()
                for item in os.getenv("TELEGRAM_CHAT_IDS", "").split(",")
                if item.strip()
            )
        )
        if not token or not chat_ids:
            logger.warning("Telegram envio=falhou motivo=ConfiguracaoAusente")
            execucao.delivery_status = "DELIVERY_FAILED"
            return
        completo = dia == execucao.dia and execucao.permite_publicacao()
        dados = _coletar_dados_telegram(dia, resultados, execucao) if completo else None
        mensagens = _mensagens_telegram(dia, dados) if dados is not None else [
            execucao.aviso_operacional({
                codigo: u.apelido or u.nome for codigo, u in carregar_unidades().items()
            })
        ]
        saude_robo.salvar_resumo(dia, "\n".join(mensagens) + "\n")
        planilha = None
        if dados is not None:
            try:
                planilha = _criar_planilha_telegram(dia, dados)
            except (OSError, ValueError) as exc:
                execucao.document_delivery_status = "DELIVERY_FAILED"
                logger.warning("Telegram planilha indisponivel motivo=%s", type(exc).__name__)
        resultados_envio = [
            resultado
            for texto in mensagens
            for resultado in enviar_mensagem(
                texto,
                token=token,
                chat_ids=chat_ids,
                parse_mode="HTML" if completo else None,
            )
        ]
        sucesso = bool(resultados_envio) and all(resultado.sucesso for resultado in resultados_envio)
        execucao.delivery_status = "DELIVERY_SUCCESS" if sucesso else "DELIVERY_FAILED"
        resultados_documento = []
        if planilha is not None:
            execucao.document_delivery_status = "DELIVERY_FAILED"
            resultados_documento = enviar_documento(
                planilha,
                f"Detalhe do dia {formatar_data(dia)}",
                token=token,
                chat_ids=chat_ids,
            )
            if not resultados_documento or any(not item.sucesso for item in resultados_documento):
                execucao.document_delivery_status = "DELIVERY_FAILED"
                logger.warning("Telegram planilha envio=falhou")
            else:
                execucao.document_delivery_status = "DELIVERY_SUCCESS"
        if sucesso:
            logger.info("Telegram envio=ok destinatarios=%s", len(chat_ids))
        else:
            falha = next(
                (resultado for resultado in resultados_envio if not resultado.sucesso),
                ResultadoDestinatario(chat_id="", sucesso=False, tipo_erro="sem destinatário"),
            )
            motivo = (
                f"HTTP {falha.status_http}"
                if falha.status_http is not None
                else falha.tipo_erro or "erro desconhecido"
            )
            logger.warning("Telegram envio=falhou motivo=%s", motivo)
        _registrar_envio_telegram(registro, dia, sucesso, execucao)
    except Exception as exc:
        execucao.erro_publicacao = f"Publicação não concluída ({type(exc).__name__})."
        if execucao.delivery_status == "NOT_ATTEMPTED":
            execucao.delivery_status = "DELIVERY_FAILED"
        logger.warning("Telegram envio=falhou motivo=%s", type(exc).__name__)


async def _executar_placar(
    dia: date,
    timeout: float,
    diretorio: Path | None,
    execucao: ExecucaoNoturna | None = None,
) -> tuple[int, set[int]]:
    falhas = 0
    unidades_com_falha: set[int] = set()
    for posto in POSTOS.values():
        if execucao is not None:
            execucao.registrar(FonteExecucao(nome="placar", unidade=posto.empresa_codigo, dia=dia))
        try:
            placar = await asyncio.wait_for(
                obter_placar(posto.empresa_codigo, dia.strftime("%Y-%m"), dia, forcar_coleta=True),
                timeout=timeout,
            )
            if placar.posto != posto.empresa_codigo or placar.dia != dia:
                raise ValueError("Placar fora do escopo solicitado.")
            if execucao is not None:
                fonte = FonteExecucao(
                    nome="placar", unidade=posto.empresa_codigo, dia=dia, run_id=execucao.run_id,
                    coletado_em=placar.proveniencia.executado_em, cobertura_inicio=dia.replace(day=1),
                    cobertura_fim=dia, estado="SUCCESS",
                )
                execucao.registrar(fonte)
                execucao.avaliar()
                if next(f for f in execucao.fontes if f.nome == "placar"
                        and f.unidade == posto.empresa_codigo).estado == "STALE":
                    raise ValueError("Placar sem coleta corrente.")
                placar = placar.model_copy(update={"proveniencia": placar.proveniencia.model_copy(
                    update={"execucao_id": execucao.run_id},
                )})
            salvar_placar(placar, diretorio=diretorio)
            logger.info(
                "Placar posto=%s mes=%s dia=%s litros=%s",
                posto.empresa_codigo,
                placar.mes,
                formatar_data(placar.dia),
                placar.acumulado,
            )
        except Exception as exc:
            if execucao is not None:
                execucao.registrar(FonteExecucao(
                    nome="placar", unidade=posto.empresa_codigo, dia=dia, estado="FAILED",
                    motivo=f"Coleta comercial não concluída ({type(exc).__name__}).",
                ))
            logger.error(
                "Placar nao concluido posto=%s mes=%s tipo=%s",
                posto.empresa_codigo,
                dia.strftime("%Y-%m"),
                type(exc).__name__,
            )
            falhas += 1
            unidades_com_falha.add(posto.empresa_codigo)
    return falhas, unidades_com_falha


def main() -> int:
    parser = argparse.ArgumentParser(description="Auditoria noturna somente leitura.")
    parser.add_argument(
        "--dia",
        type=_dia,
        default=None,
        help="Dia em dd/mm/aaaa ou aaaa-mm-dd; padrão: ontem em Recife.",
    )
    parser.add_argument("--enviar-telegram", action="store_true")
    parser.add_argument("--reenviar", action="store_true")
    args = parser.parse_args()
    dia = args.dia or ontem()
    execucao = ExecucaoNoturna(dia=dia)
    resultados: list[ResultadoDiario] = []
    logging.getLogger("httpx").setLevel(logging.WARNING)
    logging.getLogger("httpcore").setLevel(logging.WARNING)
    diretorio = persistencia_caixa.DIRETORIO
    diretorio.mkdir(parents=True, exist_ok=True)
    handler = logging.FileHandler(diretorio / "noturno.log", encoding="utf-8")
    handler.setFormatter(FormatoLogBrasil("%(asctime)s %(levelname)s %(message)s"))
    modulo_logger = logging.getLogger("src.modules.cash_reconciliation")
    modulo_logger.setLevel(logging.INFO)
    modulo_logger.addHandler(handler)
    try:
        saude_robo.salvar_execucao(execucao)
        try:
            load_dotenv(CAMINHO_ENV, override=False)
            codigo_saida = asyncio.run(executar(
                dia, executar_placar=_executar_placar, execucao=execucao,
                resultados_execucao=resultados,
            ))
        except Exception as exc:
            logger.error("Auditoria interrompida run=%s tipo=%s", execucao.run_id, type(exc).__name__)
            execucao.registrar(FonteExecucao(
                nome="execucao", dia=dia, estado="FAILED",
                motivo=f"Execução interrompida ({type(exc).__name__}); consultar diagnóstico local.",
            ))
            execucao.finalizar()
            saude_robo.salvar_resumo(dia, execucao.aviso_operacional({}))
            saude_robo.salvar_estado(saude_robo.EstadoRobo(
                ultima_execucao=agora(), dia_processado=dia, unidades_processadas=(),
                unidades_com_falha=(), resumo_gerado=False, execucao=execucao,
            ))
            codigo_saida = 1
        saude_robo.salvar_execucao(execucao)
        _tentar_envio_telegram(
            dia,
            dia_manual=args.dia is not None,
            enviar_manual=args.enviar_telegram,
            reenviar=args.reenviar,
            execucao=execucao, resultados=resultados,
        )
        saude_robo.salvar_execucao(execucao)
        return 1 if (codigo_saida or not execucao.permite_publicacao()
                     or execucao.delivery_status == "DELIVERY_FAILED"
                     or execucao.document_delivery_status == "DELIVERY_FAILED"
                     or execucao.erro_publicacao) else 0
    finally:
        modulo_logger.removeHandler(handler)
        handler.close()


if __name__ == "__main__":
    raise SystemExit(main())
