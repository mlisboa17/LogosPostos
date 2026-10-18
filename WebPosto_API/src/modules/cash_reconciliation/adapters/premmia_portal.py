"""Coleta Premmia somente leitura; conciliacao sempre le a planilha mensal."""

from __future__ import annotations

import argparse
import hashlib
import asyncio
import json
import logging
import os
import re
import time
import unicodedata
import uuid
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Iterable
from urllib.parse import parse_qs, unquote_plus, urlencode, urlsplit, urlunsplit

from openpyxl import Workbook, load_workbook
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

from src.shared.segredos import obter_segredo

from ..domain.cartoes import TransacaoAdquirente
from ..domain.models import Unidade
from ..domain.tempo import FUSO, agora, ontem

logger = logging.getLogger(__name__)
DIRETORIO = Path(__file__).resolve().parents[4] / "data" / "premmia"
LOGIN = "https://cn.vibraenergia.com.br/premmia/revendedor/login"
PAGAMENTOS = "https://cn.vibraenergia.com.br/premmia/revendedor/conferencia/pagamentos"
TIMEOUT = 60_000
TIMEOUT_LOGIN = 90_000
COLUNAS = (
    "CPF",
    "Nome",
    "Produto",
    "Valor líquido",
    "Data/Hora da transação",
    "Código Transação",
    "Forma de Pagamento",
    "Status",
)
COLUNAS_COLETAS = ("Dia", "Quantidade", "Total", "Coletado em")
AVISO_HUMANO = "Premmia pediu verificação humana / credencial recusada"


class PremmiaErro(RuntimeError):
    def __init__(self, mensagem: str, *, etapa: str | None = None):
        super().__init__(mensagem)
        self.etapa = etapa


class PremmiaVerificacaoHumana(PremmiaErro):
    def __init__(self):
        super().__init__(AVISO_HUMANO)


class PremmiaDiaNaoColetado(PremmiaErro):
    pass


class PremmiaIntegridadeAviso(PremmiaErro):
    pass


def chave_linha(registro: dict) -> tuple[str, str, Decimal, str]:
    return (
        str(registro["Código Transação"]).strip(),
        data_hora(registro["Data/Hora da transação"]).isoformat(),
        dinheiro(registro["Valor líquido"]),
        str(registro["Forma de Pagamento"]).strip(),
    )


def verificar_integridade(
    registros: list[dict], gravadas: list[dict], dias: Iterable[date],
) -> None:
    faltantes = 0
    for dia in set(dias):
        esperadas = {chave_linha(r) for r in registros if data_hora(r["Data/Hora da transação"]).date() == dia}
        atuais = {chave_linha(r) for r in gravadas if data_hora(r["Data/Hora da transação"]).date() == dia}
        faltantes += len(esperadas ^ atuais)
    if faltantes:
        mensagem = f"Premmia: {faltantes} linhas do portal não entraram na planilha"
        logger.warning("%s", mensagem)
        raise PremmiaIntegridadeAviso(mensagem)


def normalizar(texto: str) -> str:
    return "".join(
        letra
        for letra in unicodedata.normalize("NFKD", texto.casefold())
        if not unicodedata.combining(letra)
    )


def caminho_mensal(unidade: Unidade, dia: date, diretorio: Path | None = None) -> Path:
    apelido = re.sub(r"[^a-z0-9]", "", normalizar(unidade.apelido or unidade.nome))
    return (diretorio or DIRETORIO) / f"posto{apelido}_{dia:%m_%Y}.xlsx"


def dinheiro(valor) -> Decimal:
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        resultado = Decimal(texto)
    except InvalidOperation:
        raise PremmiaErro("Valor monetário inválido no XLSX.") from None
    if not resultado.is_finite():
        raise PremmiaErro("Valor monetário não finito no XLSX.")
    return resultado


def data_hora(valor) -> datetime:
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=FUSO) if valor.tzinfo is None else valor.astimezone(FUSO)
    texto = str(valor).strip()
    try:
        momento = datetime.fromisoformat(texto.replace("Z", "+00:00"))
    except ValueError:
        for formato in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
            try:
                return datetime.strptime(texto, formato).replace(tzinfo=FUSO)
            except ValueError:
                continue
        raise PremmiaErro("Data inválida no XLSX.") from None
    return momento.replace(tzinfo=FUSO) if momento.tzinfo is None else momento.astimezone(FUSO)


def ler_original(caminho: Path, inicio: date, fim: date) -> list[dict]:
    workbook = load_workbook(caminho, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = (
            workbook["Conferência"]
            if "Conferência" in workbook.sheetnames
            else workbook.worksheets[0]
        )
        linhas = iter(sheet.iter_rows(values_only=True))
        indices = None
        for _, linha in zip(range(30), linhas):
            rotulos = [normalizar(str(v or "").strip()) for v in linha]
            if all(normalizar(c) in rotulos for c in COLUNAS):
                indices = {c: rotulos.index(normalizar(c)) for c in COLUNAS}
                break
        if indices is None:
            raise PremmiaErro("Layout do XLSX inesperado; original preservado.")
        registros = []
        for linha in linhas:
            if not any(v is not None and str(v).strip() for v in linha):
                continue
            if len(linha) <= max(indices.values()):
                raise PremmiaErro("Linha incompleta no XLSX.")
            registro = {c: linha[i] for c, i in indices.items()}
            momento = data_hora(registro["Data/Hora da transação"])
            if not inicio <= momento.date() <= fim:
                raise PremmiaErro("Exportação fora do período filtrado.")
            if not str(registro["Código Transação"] or "").strip():
                raise PremmiaErro("Código Transação ausente no XLSX.")
            if not registro["Forma de Pagamento"] or not registro["Status"]:
                raise PremmiaErro("Forma/status ausente no XLSX.")
            dinheiro(registro["Valor líquido"])
            registro["Data/Hora da transação"] = momento.isoformat()
            registros.append(registro)
        return registros
    finally:
        workbook.close()


def ler_mensal(caminho: Path) -> tuple[list[dict], list[dict]]:
    if not caminho.exists():
        return [], []
    workbook = load_workbook(caminho, read_only=True, data_only=True, keep_links=False)
    try:
        resultado = []
        for aba, colunas in (
            ("Transações", (*COLUNAS, "Dia", "Coletado em")),
            ("Coletas", COLUNAS_COLETAS),
        ):
            if aba not in workbook.sheetnames:
                raise PremmiaErro("Planilha mensal sem abas obrigatórias.")
            linhas = iter(workbook[aba].iter_rows(values_only=True))
            if tuple(next(linhas, ())) != tuple(colunas):
                raise PremmiaErro("Layout da planilha mensal inválido.")
            resultado.append([dict(zip(colunas, linha)) for linha in linhas if any(linha)])
        return resultado[0], resultado[1]
    finally:
        workbook.close()


def dias_coletados(caminho: Path) -> set[date]:
    _, coletas = ler_mensal(caminho)
    return {date.fromisoformat(str(c["Dia"])) for c in coletas}


def dias_faltantes(unidade: Unidade, ate: date, diretorio: Path | None = None) -> list[date]:
    inicio_mes = ate.replace(day=1)
    anterior = inicio_mes - timedelta(days=1)
    meses = [(anterior.replace(day=1), anterior), (inicio_mes, ate)]
    faltantes = []
    for inicio, fim in meses:
        caminho = caminho_mensal(unidade, inicio, diretorio)
        # Completa o mes anterior quando ja monitorado; na virada, mesmo sem planilha.
        if inicio != inicio_mes and not caminho.exists() and ate.day != 1:
            continue
        coletados = dias_coletados(caminho)
        faltantes.extend(
            dia
            for dia in (inicio + timedelta(days=i) for i in range((fim - inicio).days + 1))
            if dia not in coletados
        )
    return faltantes


def acrescentar_coleta(
    unidade: Unidade,
    registros: list[dict],
    dias: list[date],
    diretorio: Path | None = None,
) -> None:
    por_mes = defaultdict(list)
    for dia in sorted(set(dias)):
        por_mes[(dia.year, dia.month)].append(dia)
    instante = agora().isoformat()
    for dias_mes in por_mes.values():
        caminho = caminho_mensal(unidade, dias_mes[0], diretorio)
        antigas, coletas = ler_mensal(caminho)
        coletados = {date.fromisoformat(str(c["Dia"])) for c in coletas}
        novos_dias = set(dias_mes) - coletados
        if not novos_dias:
            verificar_integridade(registros, antigas, dias_mes)
            continue
        vistos = {chave_linha(r) for r in antigas}
        novas = []
        for registro in registros:
            momento = data_hora(registro["Data/Hora da transação"])
            codigo = str(registro["Código Transação"] or "").strip()
            if not codigo:
                raise PremmiaErro("Código Transação ausente na coleta.")
            chave = chave_linha(registro)
            if momento.date() in novos_dias and chave not in vistos:
                novas.append(
                    {
                        **registro,
                        "Código Transação": codigo,
                        "Data/Hora da transação": momento.isoformat(),
                        "Dia": momento.date().isoformat(),
                        "Coletado em": instante,
                    }
                )
                vistos.add(chave)
        for dia in sorted(novos_dias):
            diario = [r for r in novas if r["Dia"] == dia.isoformat()]
            coletas.append(
                {
                    "Dia": dia.isoformat(),
                    "Quantidade": len(diario),
                    "Total": str(sum((dinheiro(r["Valor líquido"]) for r in diario), Decimal(0))),
                    "Coletado em": instante,
                }
            )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Transações"
        sheet.append((*COLUNAS, "Dia", "Coletado em"))
        for registro in antigas + novas:
            sheet.append([registro.get(c) for c in (*COLUNAS, "Dia", "Coletado em")])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet = workbook.create_sheet("Coletas")
        sheet.append(COLUNAS_COLETAS)
        for coleta in coletas:
            sheet.append([coleta[c] for c in COLUNAS_COLETAS])
        caminho.parent.mkdir(parents=True, exist_ok=True)
        temporario = None
        try:
            with NamedTemporaryFile(dir=caminho.parent, suffix=".xlsx", delete=False) as arquivo:
                temporario = Path(arquivo.name)
            workbook.save(temporario)
            os.replace(temporario, caminho)
        finally:
            workbook.close()
            if temporario is not None:
                temporario.unlink(missing_ok=True)
        gravadas, _ = ler_mensal(caminho)
        verificar_integridade(registros, gravadas, dias_mes)


def buscar_transacoes(
    unidade: Unidade,
    dia: date,
    diretorio: Path | None = None,
) -> list[TransacaoAdquirente]:
    caminho = caminho_mensal(unidade, dia, diretorio)
    registros, coletas = ler_mensal(caminho)
    if dia.isoformat() not in {c["Dia"] for c in coletas}:
        raise PremmiaDiaNaoColetado("Premmia: dia ainda não coletado na planilha mensal.")
    return [
        TransacaoAdquirente(
            adquirente="PREMMIA",
            identificador=hashlib.sha256(repr(chave_linha(r)).encode("utf-8")).hexdigest(),
            nsu=str(r["Código Transação"]),
            autorizacao=None,
            valor=dinheiro(r["Valor líquido"]),
            momento=data_hora(r["Data/Hora da transação"]),
            bandeira=str(r["Forma de Pagamento"]),
            situacao=None if normalizar(str(r["Status"])) == "processada" else str(r["Status"]),
        )
        for r in registros
        if r["Dia"] == dia.isoformat()
    ]


def verificar_pagina(page) -> None:
    texto = normalizar(page.locator("body").inner_text(timeout=TIMEOUT))
    desafios = (
        r"captcha|nao sou (?:um )?robo|verifi(?:que|car|cacao).*human",
        r"autenticacao (?:de|em) (?:dois|2)|duplo fator|dois fatores|2fa",
        r"codigo (?:de verificacao|de seguranca|enviado|por sms)|token de acesso",
        r"confirmar.*(?:sms|e-mail)|verificacao.*(?:sms|e-mail)|verifique.*(?:identidade|navegador)",
        r"aceit(?:ar|e).*termos|concordo.*termos|alter(?:ar|e|acao).*senha|"
        r"troca de senha|atualiz(?:ar|e).*cadastro|confirm(?:ar|e).*cadastro|"
        r"senha expirad|primeiro acesso|redefin(?:ir|a).*senha",
        r"senha.*(?:incorret|invalid)|(?:usuario|login).*invalid|"
        r"credencia(?:l|is).*invalid|acesso negado|conta bloqueada|usuario ou senha|login ou senha",
    )
    captcha = page.locator(
        "iframe[src*='recaptcha'], iframe[src*='hcaptcha'], "
        "iframe[src*='challenges.cloudflare'], .g-recaptcha, .h-captcha"
    )
    if any(re.search(padrao, texto) for padrao in desafios) or any(
        captcha.nth(i).is_visible() for i in range(captcha.count())
    ):
        raise PremmiaVerificacaoHumana()


def unico(locator):
    visiveis = [locator.nth(i) for i in range(locator.count()) if locator.nth(i).is_visible()]
    if len(visiveis) != 1:
        raise PremmiaErro("Seletor do portal não confirmado.")
    return visiveis[0]


def aguardar(page, condicao, *, timeout: int = TIMEOUT) -> None:
    limite = time.monotonic() + timeout / 1000
    while time.monotonic() < limite:
        verificar_pagina(page)
        if condicao():
            return
        page.wait_for_timeout(300)
    raise PlaywrightTimeoutError("Tempo limite na navegação Premmia.")


def aguardar_exportacao(page) -> None:
    limite = time.monotonic() + 180
    while time.monotonic() < limite:
        verificar_pagina(page)
        try:
            page.wait_for_load_state("networkidle", timeout=1000)
            botao = page.get_by_role("button", name="Exportar", exact=True)
            botao.wait_for(state="visible", timeout=1000)
        except PlaywrightTimeoutError:
            continue
        if botao.is_enabled():
            verificar_pagina(page)
            return
    raise PlaywrightTimeoutError("Exportar não ficou pronto em 180 segundos.")


def url_periodo(url: str, inicio: date, fim: date) -> str:
    partes = urlsplit(url)
    if partes.scheme != "https" or partes.netloc != "cn.vibraenergia.com.br":
        raise PremmiaErro("Endereço de pagamentos inesperado.")
    parametros = partes.query.split("&")
    produtos = [p for p in parametros if unquote_plus(p.partition("=")[0]) == "products"]
    if not produtos or not all(p.partition("=")[2] for p in produtos):
        raise PremmiaErro("products não retornado pelo portal.")
    preservados = [
        p for p in parametros if unquote_plus(p.partition("=")[0]) not in {"start", "end", "page"}
    ]
    novos = urlencode(
        {"page": 1, "start": inicio.isoformat() + " 00:00", "end": fim.isoformat() + " 23:59"}
    )
    return urlunsplit(
        (
            partes.scheme,
            partes.netloc,
            partes.path,
            "&".join(preservados + [novos]),
            partes.fragment,
        )
    )


def abrir_login(playwright, unidade: Unidade, *, visivel: bool):
    modos = (True,) if visivel else (False, True)
    for modo_visivel in modos:
        browser = playwright.chromium.launch(
            headless=False, args=[] if modo_visivel else ["--headless=new"],
            slow_mo=300 if modo_visivel else 0,
        )
        try:
            versao = browser.version
            context = browser.new_context(
                accept_downloads=True, locale="pt-BR", timezone_id="America/Recife",
                viewport={"width": 1366, "height": 768},
                user_agent=f"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/{versao} Safari/537.36",
            )
            page = context.new_page()
            page.set_default_timeout(TIMEOUT)
            page.set_default_navigation_timeout(TIMEOUT_LOGIN)
            resposta = page.goto(LOGIN, wait_until="commit", timeout=TIMEOUT_LOGIN)
            if resposta is None or resposta.status >= 400:
                raise PremmiaErro("Portal indisponível na página de login.", etapa="abrir login")
            aguardar(page, lambda: page.locator("input[data-testid='Login']").is_visible(), timeout=TIMEOUT_LOGIN)
            verificar_pagina(page)
            logger.info("Premmia unidade=%s etapa=abrir login modo=%s situacao=ok",
                        unidade.empresa_codigo, "visivel" if modo_visivel else "invisivel")
            return browser, page
        except PremmiaVerificacaoHumana:
            browser.close()
            raise
        except (PlaywrightError, PremmiaErro):
            browser.close()
            if modo_visivel:
                raise
            logger.warning("Premmia unidade=%s etapa=abrir login modo=invisivel fallback=visivel; credenciais ainda não enviadas",
                           unidade.empresa_codigo)
        except Exception:
            browser.close()
            raise


def coletar_sync(
    unidade: Unidade,
    dias: list[date],
    *,
    visivel: bool = False,
    diretorio: Path | None = None,
) -> None:
    if not dias:
        return
    usuario = obter_segredo(unidade.premmia_usuario_env or "")
    senha = obter_segredo(unidade.premmia_senha_env or "")
    if not usuario or not senha:
        erro = PremmiaVerificacaoHumana()
        erro.etapa = "validar credenciais"
        raise erro
    tentativa_login = (diretorio or DIRETORIO) / "logins" / f"{unidade.empresa_codigo}_{agora().date().isoformat()}.json"
    if tentativa_login.exists():
        raise PremmiaErro("Premmia: login já tentado nesta noite; sem nova tentativa.", etapa="login já tentado")
    pasta = (diretorio or DIRETORIO) / "originais"
    pasta.mkdir(parents=True, exist_ok=True)
    etapa = "iniciar navegador"
    try:
        with sync_playwright() as playwright:
            etapa = "abrir login"
            browser, page = abrir_login(playwright, unidade, visivel=visivel)
            try:
                page.set_default_navigation_timeout(TIMEOUT)
                unico(page.locator("input[data-testid='Login']")).fill(usuario)
                unico(page.locator("input[data-testid='Senha']")).fill(senha)
                verificar_pagina(page)
                etapa = "enviar login"
                tentativa_login.parent.mkdir(parents=True, exist_ok=True)
                try:
                    with tentativa_login.open("x", encoding="utf-8") as arquivo:
                        json.dump({"registrado_em": agora().isoformat()}, arquivo)
                        arquivo.flush()
                        os.fsync(arquivo.fileno())
                except FileExistsError:
                    raise PremmiaErro("Premmia: login já tentado nesta noite; sem nova tentativa.", etapa=etapa) from None
                unico(page.get_by_role("button", name="Entrar", exact=True)).click()
                etapa = "confirmar login"
                aguardar(page, lambda: "/login" not in page.url)
                etapa = "abrir pagamentos/products"
                resposta = page.goto(PAGAMENTOS, wait_until="commit")
                verificar_pagina(page)
                if resposta is None or resposta.status >= 400 or "/login" in page.url:
                    raise PremmiaErro("Pagamentos indisponíveis.")
                aguardar(page, lambda: bool(parse_qs(urlsplit(page.url).query).get("products")))
                url_original = page.url
                intervalos = [(min(dias), max(dias))]
                if min(dias) != max(dias):
                    intervalos.append((min(dias), min(dias)))
                for tentativa, (inicio, fim) in enumerate(intervalos):
                    try:
                        etapa = "filtro pela URL"
                        logger.info("Premmia unidade=%s etapa=%s", unidade.empresa_codigo, etapa)
                        filtrada = url_periodo(url_original, inicio, fim)
                        resposta = page.goto(filtrada, wait_until="commit")
                        if resposta is None or resposta.status >= 400:
                            raise PremmiaErro("Filtro por URL indisponível.")
                        etapa = "aguardar Exportar"
                        aguardar_exportacao(page)
                        atuais, esperados = parse_qs(urlsplit(page.url).query), parse_qs(
                            urlsplit(filtrada).query
                        )
                        if any(
                            atuais.get(k) != esperados.get(k)
                            for k in ("start", "end", "page", "products")
                        ):
                            raise PremmiaErro("Portal não preservou o filtro.")
                        verificar_pagina(page)
                        etapa = "download XLSX"
                        with page.expect_download(timeout=180_000) as download_info:
                            unico(page.get_by_role("button", name="Exportar", exact=True)).click()
                        verificar_pagina(page)
                        download = download_info.value
                        if Path(download.suggested_filename).suffix.lower() != ".xlsx":
                            raise PremmiaErro("Exportação diferente de XLSX.")
                        original = (
                            pasta
                            / f"{unidade.empresa_codigo}_{inicio}_{fim}_{uuid.uuid4().hex}.xlsx"
                        )
                        download.save_as(original)
                        break
                    except PlaywrightTimeoutError:
                        verificar_pagina(page)
                        if tentativa + 1 == len(intervalos):
                            raise
                        logger.warning(
                            "Premmia unidade=%s timeout etapa=%s; fallback de um dia, mesma sessão",
                            unidade.empresa_codigo,
                            etapa,
                        )
                etapa = "validar/persistir XLSX"
                registros = ler_original(original, inicio, fim)
                acrescentar_coleta(
                    unidade, registros, [d for d in dias if inicio <= d <= fim], diretorio
                )
            finally:
                browser.close()
    except PlaywrightError as exc:
        logger.warning("Premmia unidade=%s falha etapa=%s tipo=%s", unidade.empresa_codigo, etapa, type(exc).__name__)
        raise PremmiaErro(f"Premmia: {type(exc).__name__} na etapa {etapa}.", etapa=etapa) from None
    except OSError as exc:
        logger.warning("Premmia unidade=%s falha etapa=%s tipo=%s", unidade.empresa_codigo, etapa, type(exc).__name__)
        raise PremmiaErro(f"Premmia: falha de arquivo na etapa {etapa}.", etapa=etapa) from None
    except PremmiaVerificacaoHumana as exc:
        exc.etapa = etapa
        logger.warning(
            "Premmia unidade=%s etapa=%s: %s", unidade.empresa_codigo, etapa, AVISO_HUMANO
        )
        raise
    except PremmiaIntegridadeAviso as exc:
        logger.warning("Premmia unidade=%s etapa=%s: %s", unidade.empresa_codigo, etapa, exc)
        raise
    except PremmiaErro as exc:
        exc.etapa = exc.etapa or etapa
        logger.error("Premmia unidade=%s falha etapa=%s", unidade.empresa_codigo, etapa)
        raise


async def coletar(
    unidade: Unidade,
    *,
    ate: date | None = None,
    visivel: bool = False,
    diretorio: Path | None = None,
) -> None:
    if not unidade.premmia_ativo:
        return
    alvo = ate or ontem()
    dias = dias_faltantes(unidade, alvo, diretorio)
    if dias:
        aviso = (diretorio or DIRETORIO) / "avisos" / f"{unidade.empresa_codigo}.json"
        if aviso.exists():
            try:
                estado = json.loads(aviso.read_text(encoding="utf-8"))
            except (ValueError, OSError):
                raise PremmiaErro(
                    "Não foi possível ler o bloqueio de verificação humana."
                ) from None
            if not isinstance(estado, dict) or "dia_alvo" not in estado:
                raise PremmiaErro("Bloqueio de verificação humana inválido.")
            if estado["dia_alvo"] == alvo.isoformat():
                raise PremmiaVerificacaoHumana()
        try:
            await asyncio.to_thread(
                coletar_sync,
                unidade,
                dias,
                visivel=visivel,
                diretorio=diretorio,
            )
        except PremmiaVerificacaoHumana:
            aviso.parent.mkdir(parents=True, exist_ok=True)
            temporario = None
            try:
                with NamedTemporaryFile(
                    mode="w",
                    encoding="utf-8",
                    dir=aviso.parent,
                    suffix=".json",
                    delete=False,
                ) as arquivo:
                    temporario = Path(arquivo.name)
                    json.dump(
                        {"dia_alvo": alvo.isoformat(), "registrado_em": agora().isoformat()},
                        arquivo,
                    )
                os.replace(temporario, aviso)
            finally:
                if temporario is not None:
                    temporario.unlink(missing_ok=True)
            raise


def main() -> int:
    from dotenv import load_dotenv
    from ..config import carregar_unidades

    parser = argparse.ArgumentParser(description="Coleta mensal Premmia somente leitura.")
    parser.add_argument("--unidade", type=int, required=True)
    parser.add_argument("--visivel", action="store_true")
    args = parser.parse_args()
    load_dotenv(Path(__file__).resolve().parents[4] / ".env")
    if os.getenv("WEBPOSTO_WRITES") != "0":
        parser.error("WEBPOSTO_WRITES deve ser 0.")
    unidade = carregar_unidades().get(args.unidade)
    if unidade is None or not unidade.premmia_ativo:
        parser.error("Unidade sem coleta Premmia ativa.")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    try:
        asyncio.run(coletar(unidade, visivel=args.visivel))
    except PremmiaVerificacaoHumana:
        print(AVISO_HUMANO)
        return 0
    except PremmiaIntegridadeAviso as exc:
        logger.warning("Premmia unidade=%s: %s", args.unidade, exc)
        return 0
    except (PremmiaErro, OSError) as exc:
        logger.error("Coleta não concluída unidade=%s tipo=%s", args.unidade, type(exc).__name__)
        return 1
    restantes = dias_faltantes(unidade, ontem())
    if restantes:
        print(
            f"Coleta parcial unidade={args.unidade}; {len(restantes)} dias ainda faltantes "
            "(fallback de um dia); sem novo login."
        )
    else:
        print(f"Coleta concluída unidade={args.unidade}; planilha mensal preservada.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
