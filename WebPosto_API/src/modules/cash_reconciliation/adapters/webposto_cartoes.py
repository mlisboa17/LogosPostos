"""Adapters webPosto V1 para a conciliacao de recebimentos eletronicos.

CARTOES (recebimentos lancados) · ABASTECIMENTOS (frentista via identfid) ·
VENDAS/ITENS + VENDAS_FORMA_PAGAMENTO (abastecimento -> venda -> forma de pagamento).
"""
from __future__ import annotations

from datetime import date
from collections import defaultdict
import logging
from pathlib import Path
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

import httpx

from ..domain.cartoes import Abastecimento, CartaoErp
from ..domain.models import Unidade
from ..domain.tempo import ler_data_hora
from .webposto_http import paginar

TIPO_DINHEIRO = "D"
logger = logging.getLogger(__name__)
DIRETORIO_CONFERENCIA = Path(__file__).resolve().parents[4] / "data" / "conferencia"


def _dec(v: Any) -> Decimal:
    return Decimal(str(v or 0))


def _texto(v: Any) -> str | None:
    s = str(v).strip() if v is not None else ""
    return s or None


def para_cartao(r: dict[str, Any]) -> CartaoErp:
    return CartaoErp(
        codigo=r["cartaoCodigo"],
        empresa_codigo=r["empresaCodigo"],
        venda_codigo=r.get("vendaCodigo"),
        valor=_dec(r.get("valor")),
        momento=ler_data_hora(f"{r['dataMovimento']}T{r.get('horaMovimento') or '00:00:00'}"),
        administradora=str(r.get("adiministradoraDescricao") or ""),  # sic: grafia da API
        nsu=_texto(r.get("nsu")),
        nsu_tef=_texto(r.get("nsuTef")),
        autorizacao=_texto(r.get("autorizacao")),
        tipo_inclusao=_texto(r.get("tipoInclusao")),
    )


def para_abastecimento(r: dict[str, Any]) -> Abastecimento:
    momento = ler_data_hora(r["dataHoraAbastecimento"])
    return Abastecimento(
        codigo=r["abastecimentoCodigo"],
        empresa_codigo=r["empresaCodigo"],
        momento=momento,
        bico=r.get("codigoBico"),
        valor=_dec(r.get("valorTotal")),
        frentista=r.get("codigoFrentista"),
        venda_item_codigo=r.get("vendaItemCodigo"),
        litros=_dec(r.get("quantidade")) if r.get("quantidade") is not None else None,
    )


def _periodo(unidade: Unidade, inicio: date, fim: date) -> dict[str, Any]:
    return {"dataInicial": inicio.isoformat(), "dataFinal": fim.isoformat(), "empresaCodigo": unidade.empresa_codigo}


async def buscar_cartoes(unidade: Unidade, inicio: date, fim: date, *, client: httpx.AsyncClient | None = None) -> list[CartaoErp]:
    linhas = await paginar(unidade, "/INTEGRACAO/V1/CARTOES", _periodo(unidade, inicio, fim), client=client)
    return [para_cartao(r) for r in linhas if r.get("empresaCodigo") == unidade.empresa_codigo]


async def buscar_abastecimentos(unidade: Unidade, inicio: date, fim: date, *, client: httpx.AsyncClient | None = None) -> list[Abastecimento]:
    linhas = await paginar(unidade, "/INTEGRACAO/V1/ABASTECIMENTOS", _periodo(unidade, inicio, fim), client=client)
    return [para_abastecimento(r) for r in linhas
            if r.get("empresaCodigo") == unidade.empresa_codigo and not r.get("afericao") and r.get("dataHoraAbastecimento")]


TIPO_PIX = "B"


def ler_pix_cnpj_gravado(empresa_codigo: int, dia: date) -> list[CartaoErp]:
    """Importa somente dados operacionais; descarta cliente e documento fiscal."""
    from datetime import datetime
    import re
    from openpyxl import load_workbook
    from ..domain.tempo import FUSO

    if empresa_codigo != 74014:
        return []
    caminho = DIRETORIO_CONFERENCIA / f"pix_cnpj_doze_{dia:%m_%Y}.xlsx"
    if not caminho.is_file():
        return []
    from zipfile import BadZipFile
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        workbook = load_workbook(caminho, read_only=True, data_only=True)
    except (OSError, BadZipFile, InvalidFileException):
        logger.exception("Falha ao abrir planilha local PIX CNPJ unidade=%s", empresa_codigo)
        raise ValueError("Falha ao abrir planilha local PIX CNPJ.") from None
    try:
        sheet = workbook["PIX CNPJ Doze"]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        if len(headers) != 12 or tuple(headers[i] for i in (0, 2, 6, 8)) != ("Data", "Valor", "Venda", "Caixa"):
            raise ValueError("Layout PIX CNPJ inválido.")
        saida = []
        for row in rows:
            if row[0] in (None, "TOTAL"):
                continue
            data = datetime.strptime(str(row[0]), "%d/%m/%Y").date()
            if data != dia:
                continue
            if row[2] is None or not _dec(row[2]).is_finite() or _dec(row[2]) <= 0:
                raise ValueError("Valor PIX CNPJ inválido.")
            if row[10] or row[11]:
                logger.warning("PIX CNPJ unidade=%s: anotacoes da planilha nao sao conferencia auditada", empresa_codigo)
            operador = str(row[9] or "")
            match = re.fullmatch(r"(.*)\s+\((\d+)\)", operador)
            momento = datetime.strptime(f"{row[0]} {row[1]}", "%d/%m/%Y %H:%M").replace(tzinfo=FUSO)
            saida.append(CartaoErp(
                codigo=int(row[5]), empresa_codigo=empresa_codigo, venda_codigo=int(row[6]) if row[6] else None,
                valor=_dec(row[2]), momento=momento, administradora="PIX CNPJ",
                nsu=None, nsu_tef=None, autorizacao=None, tipo_inclusao=str(row[7] or "") or None,
                caixa_codigo=int(row[8]) if row[8] else None, frentista=int(match[2]) if match else None,
                nome_frentista=match[1] if match else None, motivo_frentista="operador do caixa na planilha local",
            ))
        return saida
    except (KeyError, ValueError, TypeError, StopIteration, InvalidOperation):
        logger.exception("Planilha local PIX CNPJ invalida unidade=%s", empresa_codigo)
        raise ValueError("Planilha local PIX CNPJ inválida.") from None
    finally:
        workbook.close()


def _cancelada(valor: Any) -> bool:
    if valor is None:
        return False
    marcador = str(valor).strip().upper()
    if marcador in {"N", "FALSE", "0"}:
        return False
    if marcador in {"S", "TRUE", "1"}:
        return True
    raise ValueError("Marcador de cancelamento de venda inválido.")


async def buscar_contexto_vendas(
    unidade: Unidade,
    inicio: date,
    fim: date,
    *,
    client: httpx.AsyncClient | None = None,
    contexto: dict[int, dict] | None = None,
    abastecimentos: list[Abastecimento] | None = None,
    cartoes_existentes: Iterable[CartaoErp] = (),
) -> tuple[set[int], list[CartaoErp]]:
    """Retorna itens em dinheiro e pagamentos de venda que não têm cartão detalhado.

    PIX não entra em V1/CARTOES; Cartão POS sem detalhe é mantido como alerta, não como cartão lançado.
    """
    p = _periodo(unidade, inicio, fim)
    itens = await paginar(unidade, "/INTEGRACAO/V1/VENDAS/ITENS", p, client=client)
    formas = await paginar(unidade, "/INTEGRACAO/V1/VENDAS_FORMA_PAGAMENTO", p, client=client)
    vendas = await paginar(unidade, "/INTEGRACAO/V1/VENDAS", p, client=client)
    if abastecimentos is not None:
        por_item_venda = {
            item["vendaItemCodigo"]: item.get("vendaCodigo")
            for item in itens if item.get("empresaCodigo") == unidade.empresa_codigo
        }
        abastecimentos[:] = [
            a.model_copy(update={"venda_codigo": por_item_venda.get(a.venda_item_codigo)})
            for a in abastecimentos
        ]
    if contexto is not None:
        por_item = defaultdict(set)
        for abastecimento in abastecimentos or []:
            if abastecimento.frentista is not None:
                por_item[abastecimento.venda_item_codigo].add(abastecimento.frentista)
        por_venda = defaultdict(set)
        for item in itens:
            if item.get("empresaCodigo") == unidade.empresa_codigo:
                por_venda[item.get("vendaCodigo")].update(por_item[item.get("vendaItemCodigo")])
        caixas = await paginar(unidade, "/INTEGRACAO/V1/CAIXAS", p, client=client)
        operadores = {
            caixa["caixaCodigo"]: caixa.get("funcionarioCodigo")
            for caixa in caixas
            if caixa.get("empresaCodigo") == unidade.empresa_codigo
        }
        for venda in vendas:
            if venda.get("empresaCodigo") != unidade.empresa_codigo:
                continue
            candidatos = por_venda[venda["vendaCodigo"]]
            if len(candidatos) == 1:
                frentista = next(iter(candidatos))
                motivo = "frentista do abastecimento vinculado à venda"
            elif len(candidatos) > 1:
                frentista, motivo = None, "venda com abastecimentos de frentistas diferentes"
            else:
                frentista = operadores.get(venda.get("caixaCodigo"))
                motivo = (
                    "operador do caixa"
                    if frentista is not None
                    else "venda sem abastecimento ou caixa identificável"
                )
            contexto[venda["vendaCodigo"]] = {
                "frentista": frentista,
                "motivo_frentista": motivo,
                "caixa_codigo": venda.get("caixaCodigo"),
            }
    hora = {
        v["vendaCodigo"]: v.get("dataHora")
        for v in vendas
        if v.get("empresaCodigo") == unidade.empresa_codigo and not _cancelada(v.get("cancelada"))
    }
    vendas_dinheiro = {
        f["vendaCodigo"] for f in formas
        if f.get("empresaCodigo") == unidade.empresa_codigo
        and f.get("tipoFormaPagamento") == TIPO_DINHEIRO and f["vendaCodigo"] in hora
    }
    itens_dinheiro = {
        i["vendaItemCodigo"] for i in itens
        if i.get("empresaCodigo") == unidade.empresa_codigo and i.get("vendaCodigo") in vendas_dinheiro
    }
    pix = [
        CartaoErp(
            codigo=f.get("codigo") or f["vendaCodigo"],
            empresa_codigo=f["empresaCodigo"],
            venda_codigo=f["vendaCodigo"],
            valor=_dec(f.get("valorPagamento")),
            momento=ler_data_hora(hora[f["vendaCodigo"]]),
            administradora=f"PIX (venda) {f.get('nomeFormaPagamento') or ''}".strip(),
            nsu=None, nsu_tef=None, autorizacao=None,
        )
        for f in formas
        if f.get("tipoFormaPagamento") == TIPO_PIX and hora.get(f["vendaCodigo"]) and f.get("empresaCodigo") == unidade.empresa_codigo
    ]
    vendas_com_cartao = {
        c.venda_codigo for c in cartoes_existentes if c.venda_codigo is not None
    }
    pos_sem_detalhe = [
        CartaoErp(
            codigo=int(f.get("codigo") or f["vendaCodigo"]),
            empresa_codigo=f["empresaCodigo"],
            venda_codigo=f["vendaCodigo"],
            valor=_dec(f.get("valorPagamento")),
            momento=ler_data_hora(hora[f["vendaCodigo"]]),
            administradora=f"CARTÃO POS SEM DETALHE {f.get('nomeFormaPagamento') or ''}".strip(),
            nsu=None,
            nsu_tef=None,
            autorizacao=None,
            caixa_codigo=(contexto or {}).get(f["vendaCodigo"], {}).get("caixa_codigo"),
            frentista=(contexto or {}).get(f["vendaCodigo"], {}).get("frentista"),
        )
        for f in formas
        if f.get("empresaCodigo") == unidade.empresa_codigo
        and f.get("tipoFormaPagamento") not in {TIPO_DINHEIRO, TIPO_PIX}
        and (
            "CARTAO POS" in str(f.get("nomeFormaPagamento") or "").upper().replace("Ã", "A")
            or str(f.get("tipoFormaPagamento") or "").strip().upper() in {"C", "CARTAO", "CARTÃO"}
        )
        and f.get("vendaCodigo") not in vendas_com_cartao
        and hora.get(f.get("vendaCodigo"))
    ]
    return itens_dinheiro, [*pix, *pos_sem_detalhe]
