"""Compoe o relatorio detalhado a partir de resultados locais ja persistidos."""
from __future__ import annotations

import asyncio
import logging
from datetime import date, datetime
from decimal import Decimal
from statistics import median
from typing import Iterable

from ..adapters.webposto_cartoes import buscar_abastecimentos
from ..adapters.webposto_detalhes import buscar_vales
from ..adapters.webposto_http import WebPostoErro
from ..application.auditar_fechamento import auditar_unidade
from ..application.recebimentos import conciliar_recebimentos
from ..domain.execucao import ResultadoDiario
from ..domain.models import Proveniencia, Unidade
from ..domain.tempo import agora, formatar_data, formatar_data_hora
from ..rules.cartoes import VERSAO as VERSAO_CARTOES
from ..rules.fechamento import VERSAO as VERSAO_FECHAMENTO
from .resumo_diario import _descricao_despesa_segura, identificar_caixa, moeda_br

logger = logging.getLogger(__name__)

GRUPOS_ACAO = {1: "Lançar", 2: "Alterar", 3: "Apurar", 4: "Conferir à mão"}


def _mapa_nomes(registro: ResultadoDiario, nomes: dict[int, str] | None) -> dict[int, str]:
    resultado = dict(nomes or {})

    def visitar(valor) -> None:
        if hasattr(valor, "model_dump"):
            visitar(valor.model_dump(mode="python"))
        elif isinstance(valor, dict):
            codigo = valor.get("frentista") or valor.get("funcionario_codigo")
            nome = valor.get("nome_frentista") or valor.get("nome")
            if codigo is not None and nome:
                resultado.setdefault(codigo, nome)
            for item in valor.values():
                visitar(item)
        elif isinstance(valor, (tuple, list)):
            for item in valor:
                visitar(item)

    visitar(registro)
    return resultado


def _funcionario(codigo: int | None, nomes: dict[int, str]) -> str:
    if codigo is None:
        return "frentista não identificado"
    nome = nomes.get(codigo, "").strip()
    return f"{nome} ({codigo})" if nome else f"funcionário {codigo}"


def _setor(unidade: Unidade, centro_custo: int | None) -> str:
    if unidade.empresa_codigo == 11495:
        return {
            7295: "PISTA",
            21650: "CONVENIÊNCIA",
            22310: "CONVENIÊNCIA",
            24423: "FOOD",
            18713: "ajuste de NF",
        }.get(centro_custo, "setor não identificado")
    if unidade.empresa_codigo == 118508 and centro_custo == 24423:
        return "FOOD"
    return "setor único" if centro_custo is None else f"centro {centro_custo}"


def _data_hora(valor) -> str:
    return formatar_data_hora(valor) if valor else "não informado"


def _nsu_final(nsu: str | None) -> str:
    return f"…{nsu[-4:]}" if nsu else "—"


def _pendencia_info(item) -> tuple[object, bool]:
    if isinstance(item, tuple) and len(item) == 2:
        return item[0], bool(item[1])
    if isinstance(item, dict):
        return item, bool(item.get("vencida"))
    return item, False


def _totais_adquirente(item):
    if item is None or item.situacao != "ok" or not item.evidencias_completas or not item.formas:
        return None
    qtd_portal = sum(f.portal.quantidade for f in item.formas)
    portal = sum((f.portal.total for f in item.formas), Decimal(0))
    qtd_webposto = sum(f.webposto.quantidade for f in item.formas)
    webposto = sum((f.webposto.total for f in item.formas), Decimal(0))
    return qtd_portal, portal, qtd_webposto, webposto


def _bloco_status(
    registro: ResultadoDiario,
    limite_diferenca_alta: Decimal = Decimal("300.00"),
) -> tuple[list[str], int, Decimal, int]:
    adquirentes = {item.adquirente.upper(): item for item in registro.recebimentos.adquirentes}
    linhas = ["## Status da conferência", ""]
    diferencas = 0
    total_diferencas = Decimal(0)
    pendentes = 0
    caixas = registro.fechamento.caixas if registro.fechamento else ()
    caixa_nao_consolidado = any(not caixa.caixa.consolidado for caixa in caixas)

    for chave, rotulo in (("PAGBANK", "PagBank (EDI)"), ("PREMMIA", "Premmia (portal)")):
        item = adquirentes.get(chave)
        totais = _totais_adquirente(item)
        if totais is None:
            detalhe = (item.erro or item.situacao) if item else "fonte não configurada"
            linhas.append(f"- ⏳ **{rotulo}:** não conferido (motivo: {detalhe}).")
            pendentes += 1
            continue
        qtd_fonte, fonte, qtd_erp, erp = totais
        diferenca = fonte - erp
        if diferenca == 0 and qtd_fonte == qtd_erp:
            linhas.append(f"- ✅ **{rotulo}:** {qtd_fonte} · {moeda_br(fonte)}; bate com formas correspondentes no webPosto.")
            continue
        diferencas += 1
        total_diferencas += abs(diferenca)
        duplicada = next((a for a in item.alertas if a.codigo == "CARTAO_COBRANCA_DUPLICADA"), None)
        if duplicada:
            motivo = f"cobrança em duplicidade; {duplicada.mensagem}; verificar estorno ao cliente."
        elif fonte > erp:
            motivo = "adquirente maior que o webPosto → lançar/ajustar no webPosto (grupo 1/2)."
        else:
            motivo = "webPosto maior que a adquirente → apurar lançamento sem pagamento (grupo 3)."
        if abs(diferenca) > limite_diferenca_alta and caixa_nao_consolidado:
            motivo += " Diferença alta após consulta atualizada em caixa não consolidado: a confirmar até a consolidação."
        linhas.append(
            f"- ❌ **{rotulo}:** {qtd_fonte} · {moeda_br(fonte)} × "
            f"{qtd_erp} · {moeda_br(erp)} → diferença {moeda_br(abs(diferenca))}; {motivo}"
        )

    modalidades_dinheiro = [
        m for caixa in caixas for m in caixa.modalidades
        if m.rotulo.strip().casefold() == "dinheiro"
    ]
    if modalidades_dinheiro:
        apresentado = sum((m.apresentado for m in modalidades_dinheiro), Decimal(0))
        apurado = sum((m.apurado for m in modalidades_dinheiro), Decimal(0))
        diferenca = apresentado - apurado
        if diferenca == 0:
            linhas.append(f"- ✅ **Dinheiro:** operador apresentou {moeda_br(apresentado)}; bate com o apurado.")
        else:
            diferencas += 1
            total_diferencas += abs(diferenca)
            falta = "falta" if diferenca < 0 else "sobra"
            confirmar = (
                " Diferença alta em caixa não consolidado: a confirmar até a consolidação."
                if abs(diferenca) > limite_diferenca_alta and caixa_nao_consolidado
                else ""
            )
            linhas.append(
                f"- ❌ **Dinheiro:** apresentado {moeda_br(apresentado)} × apurado "
                f"{moeda_br(apurado)} → {falta} {moeda_br(abs(diferenca))}.{confirmar}"
            )
    else:
        linhas.append("- ⏳ **Dinheiro:** não conferido (motivo: apresentado pelo operador indisponível).")
        pendentes += 1

    transferencia = sum(
        (
            m.apurado
            for caixa in caixas for m in caixa.modalidades
            if "transf" in m.rotulo.casefold() or "banc" in m.rotulo.casefold()
        ),
        Decimal(0),
    )
    mais = adquirentes.get("MAIS_PAGAMENTOS")
    if mais is not None and mais.situacao == "ok" and mais.evidencias_completas:
        linhas.append(f"- ✅ **Mais Pagamentos × Transferência Bancária:** conferido; {moeda_br(transferencia)}.")
    else:
        linhas.append(
            f"- ⏳ **Mais Pagamentos × Transferência Bancária:** {moeda_br(transferencia)} "
            "(motivo: fonte Mais Pagamentos sem coleta/comprovante)."
        )
        pendentes += 1

    pix_cnpj = {
        cartao.codigo: cartao
        for item in registro.recebimentos.adquirentes
        for cartao in (*item.sem_fonte, *item.a_menor)
        if "PIX CNPJ" in cartao.administradora.upper()
    }
    valor_pix = sum((c.valor for c in pix_cnpj.values()), Decimal(0))
    if registro.empresa_codigo == 74014 or pix_cnpj:
        linhas.append(
            f"- ⏳ **PIX CNPJ × extrato:** {len(pix_cnpj)} · {moeda_br(valor_pix)} "
            "(motivo: extrato não anexado para conferência)."
        )
        pendentes += 1

    if diferencas:
        rotulo_diferencas = "diferença" if diferencas == 1 else "diferenças"
        resultado = (
            f"❌ **DIFERENÇAS ENCONTRADAS** — {diferencas} {rotulo_diferencas}, "
            f"{moeda_br(total_diferencas)}"
        )
        if pendentes:
            rotulo_pendentes = "fonte" if pendentes == 1 else "fontes"
            resultado += f" e {pendentes} {rotulo_pendentes} sem conferência."
    elif pendentes:
        resultado = f"⚠ **CONFERÊNCIA PARCIAL** — {pendentes} fonte(s) sem conferência."
    else:
        resultado = "✅ **CONFERÊNCIA REALIZADA — tudo bate**."
    linhas.extend(["", resultado, ""])
    return linhas, diferencas, total_diferencas, pendentes


def resumo_status_conferencia(registro: ResultadoDiario | None) -> dict[str, object]:
    if registro is None:
        return {"selo": "⚠", "situacao": "sem relatório", "diferencas": 0, "fontes_pendentes": 1}
    _, diferencas, total, pendentes = _bloco_status(registro)
    selo = "❌" if diferencas else "⚠" if pendentes else "✅"
    return {
        "selo": selo,
        "situacao": "diferenças encontradas" if diferencas else "conferência parcial" if pendentes else "tudo bate",
        "diferencas": diferencas,
        "diferenca_total": total,
        "fontes_pendentes": pendentes,
    }


def limite_diferenca_alta(unidade: int, dia: date) -> Decimal:
    from ..adapters.persistencia import carregar_dia

    diferencas = []
    for deslocamento in range(1, 31):
        dia_anterior = date.fromordinal(dia.toordinal() - deslocamento)
        registro = carregar_dia(unidade, dia_anterior)
        if registro is not None and registro.fechamento is not None:
            diferencas.append(abs(registro.fechamento.quebra_total))
    mediana = Decimal(str(median(diferencas))) if diferencas else Decimal(0)
    return max(Decimal("300.00"), mediana * 3)


def mudancas_desde_robo(gravado: ResultadoDiario | None, atual: ResultadoDiario) -> tuple[str, ...]:
    if gravado is None:
        return ("Não havia resultado do robô salvo para comparar.",)
    mudancas = []
    dinheiro_antigo = sum(
        (
            m.apurado for caixa in (gravado.fechamento.caixas if gravado.fechamento else ())
            for m in caixa.modalidades if m.rotulo.strip().casefold() == "dinheiro"
        ),
        Decimal(0),
    )
    dinheiro_atual = sum(
        (
            m.apurado for caixa in (atual.fechamento.caixas if atual.fechamento else ())
            for m in caixa.modalidades if m.rotulo.strip().casefold() == "dinheiro"
        ),
        Decimal(0),
    )
    if dinheiro_antigo != dinheiro_atual:
        mudancas.append(
            f"Dinheiro apurado {moeda_br(dinheiro_antigo)} → {moeda_br(dinheiro_atual)}."
        )

    pix_antigo = {
        c.codigo for item in gravado.recebimentos.adquirentes
        for c in (*item.sem_fonte, *item.a_menor)
        if "PIX CNPJ" in c.administradora.upper()
    }
    pix_atual = {
        c.codigo: c for item in atual.recebimentos.adquirentes
        for c in (*item.sem_fonte, *item.a_menor)
        if "PIX CNPJ" in c.administradora.upper()
    }
    for codigo, cartao in pix_atual.items():
        if codigo in pix_antigo:
            continue
        tipo = "manual " if (cartao.tipo_inclusao or "").casefold() == "manualmente" else ""
        mudancas.append(
            f"Novo lançamento PIX CNPJ {moeda_br(cartao.valor)} {tipo}{cartao.momento:%H:%M}."
        )

    for nome, rotulo in (("PAGBANK", "PagBank"), ("PREMMIA", "Premmia")):
        anterior = next(
            (item for item in gravado.recebimentos.adquirentes if item.adquirente.upper() == nome),
            None,
        )
        recente = next(
            (item for item in atual.recebimentos.adquirentes if item.adquirente.upper() == nome),
            None,
        )
        totais_anteriores, totais_recentes = _totais_adquirente(anterior), _totais_adquirente(recente)
        if totais_anteriores and totais_recentes and totais_anteriores != totais_recentes:
            mudancas.append(
                f"{rotulo}: totais da fonte/webPosto atualizados desde a execução do robô."
            )
    return tuple(mudancas)


async def consultar_dados_atuais(unidade: Unidade, dia: date):
    """Consulta fontes operacionais e retorna erros parciais sem marcar fontes como conferidas."""
    from ...webposto_integration.funcionarios import buscar_nomes_funcionarios

    falhas: list[str] = []
    fechamento = None
    erro_fechamento = None
    try:
        fechamento = await auditar_unidade(unidade.empresa_codigo, dia, dia)
    except (WebPostoErro, TimeoutError, OSError, ValueError) as exc:
        logger.warning(
            "Relatorio atual: fechamento indisponivel unidade=%s tipo=%s",
            unidade.empresa_codigo,
            type(exc).__name__,
        )
        erro_fechamento = "coleta falhou"
        falhas.append("Fechamento webPosto: coleta falhou.")

    try:
        recebimentos = await conciliar_recebimentos(unidade.empresa_codigo, dia)
    except (WebPostoErro, TimeoutError, OSError, ValueError) as exc:
        logger.warning(
            "Relatorio atual: recebimentos indisponiveis unidade=%s tipo=%s",
            unidade.empresa_codigo,
            type(exc).__name__,
        )
        from ..domain.recebimentos import ResultadoRecebimentos

        recebimentos = ResultadoRecebimentos(
            empresa_codigo=unidade.empresa_codigo,
            dia=dia,
            adquirentes=(),
        )
        falhas.append("Recebimentos: coleta falhou.")

    try:
        abastecimentos, vales, nomes = await asyncio.gather(
            buscar_abastecimentos(unidade, dia, dia),
            buscar_vales(unidade, dia, dia),
            buscar_nomes_funcionarios(unidade, dia, dia),
        )
    except (WebPostoErro, TimeoutError, OSError, ValueError) as exc:
        logger.warning(
            "Relatorio atual: detalhes indisponiveis unidade=%s tipo=%s",
            unidade.empresa_codigo,
            type(exc).__name__,
        )
        abastecimentos, vales, nomes = [], [], {}
        falhas.append("Abastecimentos/vales/funcionários: coleta falhou.")

    fontes = set(fechamento.proveniencia.outras_fontes if fechamento else ())
    if fechamento and fechamento.proveniencia.fonte_sangrias:
        fontes.add(fechamento.proveniencia.fonte_sangrias)
    for item in recebimentos.adquirentes:
        if item.proveniencia:
            fontes.update(item.proveniencia.outras_fontes)
    registro = ResultadoDiario(
        empresa_codigo=unidade.empresa_codigo,
        dia=dia,
        fechamento=fechamento,
        erro_fechamento=erro_fechamento,
        recebimentos=recebimentos,
        proveniencia=Proveniencia(
            execucao_id=f"relatorio-{unidade.empresa_codigo}-{dia.isoformat()}",
            executado_em=agora(),
            versao_regra="RELATORIO_CAIXA_DADOS_ATUAIS",
            fonte_sangrias=fechamento.proveniencia.fonte_sangrias if fechamento else "",
            extratos=(),
            outras_fontes=tuple(sorted(fontes)),
        ),
        versoes_regras=(VERSAO_FECHAMENTO, VERSAO_CARTOES, "RN-CARTAO-011", "RN-CARTAO-012"),
        conferencia_recebimentos=bool(recebimentos.adquirentes),
    )
    return registro, tuple(abastecimentos), tuple(v for v in vales if v.caixa_codigo == 0), nomes, tuple(falhas)


def gerar_relatorio_caixa_dia(
    registro: ResultadoDiario,
    unidade: Unidade,
    *,
    situacao_robo: str = "não informada",
    nomes_funcionarios: dict[int, str] | None = None,
    pendencias: Iterable = (),
    conferencia=None,
    dados_consultados_em: datetime | None = None,
    mudou_desde_robo: Iterable[str] = (),
    abastecimentos: Iterable = (),
    vales_fora_caixa: Iterable = (),
    falhas_coleta: Iterable[str] = (),
    limite_diferenca_alta: Decimal = Decimal("300.00"),
    incluir_codigo_interno: bool = False,
) -> str:
    """Gera Markdown a partir dos resultados operacionais fornecidos."""
    nomes = _mapa_nomes(registro, nomes_funcionarios)
    fechamento = registro.fechamento
    caixas = fechamento.caixas if fechamento else ()
    nome_unidade = unidade.apelido or unidade.nome
    rotulos_caixas = {
        a.caixa.codigo: identificar_caixa(
            nome_unidade, a.caixa, [c.caixa for c in caixas], nomes, curto=True,
        ) for a in caixas
    }
    pendencias = tuple(_pendencia_info(item) for item in pendencias)
    abastecimentos = tuple(abastecimentos)
    vales_fora_caixa = tuple(vales_fora_caixa)
    abertas = [(item, vencida) for item, vencida in pendencias if getattr(item, "status", item.get("status") if isinstance(item, dict) else None) == "aberta"]
    vencidas = [(item, vencida) for item, vencida in abertas if vencida]
    alertas = [alerta for caixa in caixas for alerta in caixa.alertas]
    adquirentes = registro.recebimentos.adquirentes
    adquirentes_com_falha = [item for item in adquirentes if item.situacao != "ok"]

    status, _, _, _ = _bloco_status(registro, limite_diferenca_alta)
    linhas = [
        f"# Relatório do caixa — {nome_unidade}",
        "",
        f"- Unidade: {nome_unidade} ({unidade.empresa_codigo})",
        f"- Dia: {formatar_data(registro.dia)}",
    ]
    instante_consulta = dados_consultados_em or registro.proveniencia.executado_em or agora()
    linhas.extend([
        f"- Dados {'consultados' if dados_consultados_em else 'registrados'} em: "
        f"{formatar_data_hora(instante_consulta)}",
        f"- Situação do robô: {situacao_robo}",
        f"- Caixas: {len(caixas)}",
        "",
        *status,
    ])
    mudancas = tuple(mudou_desde_robo)
    if mudancas:
        linhas.extend(["## Mudou desde o robô", "", *[f"- {mudanca}" for mudanca in mudancas], ""])
    falhas = tuple(falhas_coleta)
    if falhas:
        linhas.extend(["## Fontes indisponíveis", "", *[f"- {falha}" for falha in falhas], ""])
    linhas.extend([
        "## Resumo",
        "",
        f"- Quebra total: {moeda_br(fechamento.quebra_total if fechamento else Decimal(0))}.",
        f"- Caixas não consolidados: {sum(not caixa.caixa.consolidado for caixa in caixas)}.",
        f"- Alertas de fechamento: {len(alertas)}.",
        f"- Pendências abertas: {len(abertas)}, sendo {len(vencidas)} vencidas.",
        f"- Adquirentes indisponíveis: {', '.join(item.adquirente for item in adquirentes_com_falha) or 'nenhuma'}.",
        "- Abastecimentos completos e litros não constam nos resultados diários persistidos.",
        "",
        "## Caixas",
        "",
        "| Caixa | Operador | Abertura | Fechamento | Consolidado | PDV | Setor |"
        + (" Código interno |" if incluir_codigo_interno else ""),
        "|---:|---|---|---|---|---:|---|" + ("---:|" if incluir_codigo_interno else ""),
    ])
    for auditoria in caixas:
        caixa = auditoria.caixa
        linhas.append(
            f"| {rotulos_caixas[caixa.codigo]} | {_funcionario(caixa.funcionario_codigo, nomes)} | "
            f"{_data_hora(caixa.abertura)} | {_data_hora(caixa.fechamento)} | "
            f"{'sim' if caixa.consolidado else 'não'} | {caixa.pdv_codigo} | "
            f"{_setor(unidade, caixa.centro_custo)} |"
            + (f" {caixa.codigo} |" if incluir_codigo_interno else "")
        )

    linhas.extend([
        "",
        "## Formas de pagamento do caixa",
        "",
        "| Caixa | Forma | Apurado (sistema) | Apresentado (operador) | Diferença |",
        "|---:|---|---:|---:|---:|",
    ])
    for auditoria in caixas:
        for modalidade in auditoria.modalidades:
            linhas.append(
                f"| {rotulos_caixas[auditoria.caixa.codigo]} | {modalidade.rotulo} | {moeda_br(modalidade.apurado)} | "
                f"{moeda_br(modalidade.apresentado)} | {moeda_br(modalidade.diferenca)} |"
            )
    linhas.extend(["", f"Quebra total: **{moeda_br(fechamento.quebra_total if fechamento else Decimal(0))}**.", ""])
    linhas.extend(["| Caixa | Falta sem desconto | Situação |", "|---:|---:|---|"])
    for auditoria in caixas:
        desconto = auditoria.desconto_falta
        alertas_falta = [a for a in auditoria.alertas if a.codigo == "FALTA_SEM_DESCONTO"]
        if (desconto and desconto.situacao == "sem_desconto") or alertas_falta:
            valor = desconto.falta if desconto else abs(alertas_falta[0].valor or Decimal(0))
            situacao = "confirmada" if auditoria.caixa.consolidado else "provisória"
            linhas.append(f"| {rotulos_caixas[auditoria.caixa.codigo]} | {moeda_br(abs(valor))} | {situacao} |")
    if not any("FALTA_SEM_DESCONTO" == a.codigo for a in alertas) and not any(
        a.desconto_falta and a.desconto_falta.situacao == "sem_desconto" for a in caixas
    ):
        linhas.append("| — | R$ 0,00 | sem falta registrada |")

    linhas.extend(["", "## Sangrias", "", "| Caixa | Hora | Valor | Destino | Alterada |", "|---:|---|---:|---|---|"])
    for auditoria in caixas:
        for sangria in auditoria.sangrias:
            if sangria.conta_codigo is not None:
                destino = f"conta webPosto {sangria.conta_codigo}"
            elif unidade.empresa_codigo == 118508:
                destino = "repassada ao Casa Caiada (RN-SANGRIA-002)"
            elif unidade.empresa_codigo == 5555:
                destino = "assumida no cofre (RN-SANGRIA-003)"
            else:
                destino = "sem conta de destino"
            linhas.append(
                f"| {rotulos_caixas[auditoria.caixa.codigo]} | {_data_hora(sangria.momento)} | {moeda_br(sangria.valor)} | "
                f"{destino} | {'sim' if sangria.alterada else 'não'} |"
            )
    if not any(a.sangrias for a in caixas):
        linhas.append("| — | — | R$ 0,00 | nenhuma sangria persistida | — |")

    linhas.extend(["", "## Despesas", "", "| Caixa | Hora | Valor | Descrição | Plano de contas | Centro de custo / setor | Alerta |", "|---:|---|---:|---|---|---|---|"])
    despesas_ajuste_nf = []
    for auditoria in caixas:
        for despesa in auditoria.despesas:
            plano = f"{despesa.plano_conta_codigo or '—'} — {despesa.plano_conta_nome or 'nome não persistido'}"
            if despesa.centro_custo == 18713:
                despesas_ajuste_nf.append(despesa)
                continue
            alerta = "SEM PLANO ⚠" if not despesa.plano_conta_codigo else "—"
            linhas.append(
                f"| {rotulos_caixas[auditoria.caixa.codigo]} | não persistida | {moeda_br(despesa.valor)} | "
                f"{_descricao_despesa_segura(despesa.descricao or 'sem descrição')} | {plano} | "
                f"{_setor(unidade, despesa.centro_custo) if despesa.centro_custo is not None else 'centro de custo não persistido'} | {alerta} |"
            )
    if not any(d.centro_custo != 18713 for a in caixas for d in a.despesas):
        linhas.append("| — | — | R$ 0,00 | nenhuma despesa persistida | — | — | — |")
    linhas.append("\nDespesas sem centro de custo persistido não podem ser separadas por setor com segurança.")
    linhas.extend(["", "### Ajustes de NF (FECHAMENTO)", "", "| Caixa | Valor | Descrição | Plano de contas |", "|---:|---:|---|---|"])
    for despesa in despesas_ajuste_nf:
        plano = f"{despesa.plano_conta_codigo or '—'} — {despesa.plano_conta_nome or 'nome não persistido'}"
        linhas.append(
            f"| {rotulos_caixas.get(despesa.caixa_codigo, 'caixa não identificado')} | {moeda_br(despesa.valor)} | "
            f"{_descricao_despesa_segura(despesa.descricao or 'sem descrição')} | {plano} |"
        )
    if not despesas_ajuste_nf:
        linhas.append("| — | R$ 0,00 | nenhum ajuste FECHAMENTO identificado nos dados persistidos | — |")

    linhas.extend(["", "## Vales de funcionários", "", "| Caixa | Funcionário | Valor | Origem |", "|---:|---|---:|---|"])
    for auditoria in caixas:
        for vale in (*auditoria.vales_falta, *auditoria.vales):
            origem = "falta de caixa" if vale.origem.upper() == "D" else "vale comum"
            linhas.append(f"| {rotulos_caixas[auditoria.caixa.codigo]} | {_funcionario(vale.funcionario_codigo, nomes)} | {moeda_br(vale.valor)} | {origem} |")
    if not any(a.vales_falta or a.vales for a in caixas):
        linhas.append("| — | nenhum vale persistido | R$ 0,00 | — |")
    if vales_fora_caixa:
        linhas.extend(["", "### Vales fora de caixa", ""])
        for vale in vales_fora_caixa:
            linhas.append(
                f"- Vale {moeda_br(vale.valor)} fora de caixa — "
                f"{_funcionario(vale.funcionario_codigo, nomes)}."
            )

    linhas.extend(["", "## Frentistas do dia", "", "| Frentista | Abastecimentos | Litros | Valor | Ocorrências ligadas |", "|---|---:|---:|---:|---|"])
    codigos = set(nomes)
    codigos.update(caixa.caixa.funcionario_codigo for caixa in caixas)
    por_frentista: dict[int, list] = {}
    for abastecimento in abastecimentos:
        if abastecimento.momento.date() == registro.dia and abastecimento.frentista is not None:
            por_frentista.setdefault(abastecimento.frentista, []).append(abastecimento)
            if abastecimento.nome_frentista:
                nomes.setdefault(abastecimento.frentista, abastecimento.nome_frentista)
    codigos.update(por_frentista)
    if codigos:
        for codigo in sorted(codigos):
            abastecimentos_frentista = por_frentista.get(codigo, [])
            if abastecimentos_frentista:
                litros = sum((a.litros or Decimal(0) for a in abastecimentos_frentista), Decimal(0))
                valor = sum((a.valor for a in abastecimentos_frentista), Decimal(0))
                medidas = f"{len(abastecimentos_frentista)} | {litros:.2f} | {moeda_br(valor)}"
            else:
                medidas = "não persistido | não persistido | não persistido"
            linhas.append(f"| {_funcionario(codigo, nomes)} | {medidas} | consultar caixas/recebimentos abaixo |")
    else:
        linhas.append("| Não identificados nos resultados gravados | não persistido | não persistido | não persistido | — |")
    if not abastecimentos:
        linhas.append("\nA lista integral de V1/ABASTECIMENTOS não foi consultada; nenhum total foi estimado.")

    recebimentos = [item for item in adquirentes if item.adquirente.upper() in {"PAGBANK", "PREMMIA"}]
    outros = [item for item in adquirentes if item.adquirente.upper() not in {"PAGBANK", "PREMMIA"}]
    conferidos_qtd = sum(item.casados.quantidade for item in recebimentos if item.casados)
    conferidos_total = sum((item.casados.total for item in recebimentos if item.casados), Decimal(0))
    linhas.extend([
        "", "## Recebimentos", "",
        f"Conferidos: **{conferidos_qtd} · {moeda_br(conferidos_total)}**.",
        "", "| Adquirente | Situação | Casados (N · R$) |", "|---|---|---:|",
    ])
    for item in (*recebimentos, *outros):
        situacao = item.situacao if item.adquirente.upper() in {"PAGBANK", "PREMMIA"} else "pendente"
        casados = f"{item.casados.quantidade} · {moeda_br(item.casados.total)}" if item.casados else "—"
        linhas.append(f"| {item.adquirente} | {situacao} | {casados} |")
    for grupo, rotulo in GRUPOS_ACAO.items():
        itens = [item for item in (conferencia.itens if conferencia else ()) if item.grupo == grupo]
        total = sum((item.valor for item in itens), Decimal(0))
        linhas.extend(["", f"### {rotulo}: {len(itens)} · {moeda_br(total)}"])
        if itens:
            linhas.extend(["", "| Hora | Valor | De → Para | Frentista provável | Caixa | NSU final |", "|---|---:|---|---|---:|---|"])
            for item in itens:
                responsavel = item.responsavel
                linhas.append(
                    f"| {_data_hora(item.momento)} | {moeda_br(item.valor)} | {item.de or '—'} → {item.para or '—'} | "
                    f"{responsavel} | {rotulos_caixas.get(item.caixa, 'caixa não identificado')} | {_nsu_final(item.nsu)} |"
                )
        else:
            linhas.append("\nNenhum item persistido neste grupo.")
    if not conferencia:
        linhas.append("\nClassificação de ações da Sprint 14 não persistida para este dia.")

    linhas.extend(["", "## Pendências do dia", "", "| Status | Tipo | Quantidade · R$ |", "|---|---|---:|"])
    agrupadas: dict[tuple[str, str], list] = {}
    for item, vencida in abertas:
        tipo = item.get("tipo", "não informado") if isinstance(item, dict) else item.tipo
        valor = item.get("valor") if isinstance(item, dict) else item.valor
        chave = ("vencida" if vencida else "aberta", tipo)
        agrupadas.setdefault(chave, []).append(Decimal(str(valor or 0)))
    for (status, tipo), valores in sorted(agrupadas.items()):
        linhas.append(f"| {status} | {tipo} | {len(valores)} · {moeda_br(sum(valores, Decimal(0)))} |")
    if not agrupadas:
        linhas.append("| — | nenhuma pendência aberta persistida para o dia | 0 · R$ 0,00 |")
    return "\n".join(linhas).rstrip() + "\n"


def planilha_relatorio_caixa_dia(markdown: str) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    workbook.remove(workbook.active)
    blocos: list[tuple[str, list[str]]] = [("Resumo", [])]
    for linha in markdown.splitlines():
        if linha.startswith("## "):
            titulo = linha[3:].strip()[:31] or "Bloco"
            blocos.append((titulo, []))
        else:
            blocos[-1][1].append(linha)

    for titulo, linhas_bloco in blocos:
        sheet = workbook.create_sheet(titulo)
        tabela = False
        for linha in linhas_bloco:
            if linha.startswith("|"):
                celulas = [celula.strip() for celula in linha.strip().strip("|").split("|")]
                if all(set(celula) <= {"-", ":"} for celula in celulas):
                    continue
                sheet.append(celulas)
                if not tabela:
                    for cell in sheet[sheet.max_row]:
                        cell.font = Font(bold=True)
                    tabela = True
            elif linha.strip():
                sheet.append([linha.strip()])
        sheet.freeze_panes = "A2" if tabela else "A1"
        for coluna in sheet.columns:
            largura = min(60, max(12, max(len(str(c.value or "")) for c in coluna) + 2))
            sheet.column_dimensions[coluna[0].column_letter].width = largura

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()