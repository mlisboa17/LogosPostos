"""Sondagem unica, visivel e somente leitura do Premmia para o posto VIP."""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import os
import re
import shutil
import sys
import time
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import parse_qs, unquote_plus, urlencode, urlsplit, urlunsplit
from zoneinfo import ZoneInfo

import httpx
from dotenv import load_dotenv
from openpyxl import load_workbook
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
DESTINO = ROOT / "data" / "premmia"
LOGIN = "https://cn.vibraenergia.com.br/premmia/revendedor/login"
PAGAMENTOS = "https://cn.vibraenergia.com.br/premmia/revendedor/conferencia/pagamentos"
TIMEOUT = 60_000
FUSO = ZoneInfo("America/Recife")
COLUNAS = (
    "CPF",
    "Nome",
    "Produto",
    "Valor líquido",
    "Data/Hora da transação",
    "Código Transação",
    "Forma de Pagamento",
    "Status",
)


class Parada(RuntimeError):
    pass


def normalizar(texto: str) -> str:
    return "".join(
        letra
        for letra in unicodedata.normalize("NFKD", texto.casefold())
        if not unicodedata.combining(letra)
    )


def verificar_pagina(page) -> None:
    texto = normalizar(page.locator("body").inner_text(timeout=TIMEOUT))
    desafios = (
        r"captcha|nao sou (?:um )?robo|verifi(?:que|car|cacao).*human",
        r"autenticacao (?:de|em) (?:dois|2)|duplo fator|dois fatores|2fa",
        r"codigo (?:de verificacao|de seguranca|enviado|por sms)|token de acesso",
        r"confirmar.*(?:sms|e-mail)|verificacao.*(?:sms|e-mail)|verifique.*(?:identidade|navegador)",
    )
    captcha = page.locator(
        "iframe[src*='recaptcha'], iframe[src*='hcaptcha'], "
        "iframe[src*='challenges.cloudflare'], .g-recaptcha, .h-captcha"
    )
    if any(re.search(padrao, texto) for padrao in desafios) or any(
        captcha.nth(indice).is_visible() for indice in range(captcha.count())
    ):
        raise Parada("bloqueado por CAPTCHA, 2FA ou verificacao humana")
    if re.search(
        r"aceit(?:ar|e).*termos|concordo.*termos|alter(?:ar|e|acao).*senha|"
        r"troca de senha|atualiz(?:ar|e).*cadastro|confirm(?:ar|e).*cadastro|"
        r"senha expirad|primeiro acesso|redefin(?:ir|a).*senha",
        texto,
    ):
        raise Parada("outro: aceite de termos ou alteracao de conta exigida")
    if re.search(
        r"senha.*(?:incorret|invalid)|(?:usuario|login).*invalid|"
        r"credencia(?:l|is).*invalid|acesso negado|conta bloqueada|"
        r"usuario ou senha|login ou senha",
        texto,
    ):
        raise Parada("senha recusada ou conta bloqueada; nenhuma nova tentativa")


def unico(locator, descricao: str):
    visiveis = [
        locator.nth(indice) for indice in range(locator.count()) if locator.nth(indice).is_visible()
    ]
    if len(visiveis) != 1:
        raise Parada(f"outro: seletor nao confirmado para {descricao}")
    return visiveis[0]


def aguardar(page, condicao, etapa: str) -> None:
    limite = time.monotonic() + TIMEOUT / 1000
    while time.monotonic() < limite:
        verificar_pagina(page)
        if condicao():
            return
        page.wait_for_timeout(300)
    raise Parada(f"outro: timeout na etapa {etapa}; sem repetir login")


def aguardar_exportacao(page) -> None:
    limite = time.monotonic() + 180
    while time.monotonic() < limite:
        verificar_pagina(page)
        try:
            page.wait_for_load_state("networkidle", timeout=1000)
        except PlaywrightTimeoutError:
            continue
        botao = page.get_by_role("button", name="Exportar", exact=True)
        try:
            botao.wait_for(state="visible", timeout=1000)
        except PlaywrightTimeoutError:
            continue
        if botao.is_enabled():
            verificar_pagina(page)
            return
    raise PlaywrightTimeoutError("Exportar nao ficou pronto em 180 segundos")


def url_periodo(url: str, inicio: date, fim: date, pagina: int = 1) -> str:
    partes = urlsplit(url)
    if partes.scheme != "https" or partes.netloc != "cn.vibraenergia.com.br":
        raise Parada("outro: endereco de pagamentos inesperado")
    parametros = partes.query.split("&")
    produtos = [item for item in parametros if unquote_plus(item.partition("=")[0]) == "products"]
    if not produtos or not all(item.partition("=")[2] for item in produtos):
        raise Parada("outro: products nao retornado pelo portal; nenhum valor inventado")
    preservados = [
        item
        for item in parametros
        if unquote_plus(item.partition("=")[0]) not in {"start", "end", "page"}
    ]
    novos = urlencode(
        {
            "page": pagina,
            "start": inicio.isoformat() + " 00:00",
            "end": fim.isoformat() + " 23:59",
        }
    )
    return urlunsplit(
        (
            partes.scheme,
            partes.netloc,
            partes.path,
            "&".join(preservados + [novos]),
            partes.fragment,
        )
    )


def salvar_resposta_pagamentos(resposta, pasta: Path, execucao: str, erros: list[str]) -> None:
    partes = urlsplit(resposta.url)
    if partes.netloc != "cn.vibraenergia.com.br" or not re.search(
        r"pagamento|payment|transaction|transacao", partes.path, re.IGNORECASE
    ):
        return
    if "json" not in resposta.headers.get("content-type", "").casefold():
        return
    try:
        payload = resposta.json()
    except (PlaywrightError, ValueError) as exc:
        erros.append(type(exc).__name__)
        return

    def tem_lista(valor) -> bool:
        if isinstance(valor, list):
            return bool(valor) and all(isinstance(item, dict) for item in valor)
        if isinstance(valor, dict):
            return any(tem_lista(item) for item in valor.values())
        return False

    def sem_segredos(valor):
        if isinstance(valor, dict):
            return {
                chave: sem_segredos(item)
                for chave, item in valor.items()
                if not re.search(
                    r"password|senha|token|secret|authorization|cookie|credential",
                    str(chave),
                    re.IGNORECASE,
                )
            }
        if isinstance(valor, list):
            return [sem_segredos(item) for item in valor]
        return valor

    if not tem_lista(payload):
        return
    caminho = pasta / f"reserva_{execucao}_{time.time_ns()}.json"
    try:
        conteudo = json.dumps(sem_segredos(payload), ensure_ascii=False).encode("utf-8")
        with caminho.open("xb") as arquivo:
            arquivo.write(conteudo)
    except (OSError, ValueError) as exc:
        erros.append(type(exc).__name__)
        return
    print(
        f"Reserva JSON: caminho={partes.path} status={resposta.status} bytes={len(conteudo)}",
        flush=True,
    )


def dinheiro(valor) -> Decimal:
    texto = str(valor).strip().replace("R$", "").replace(" ", "")
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        resultado = Decimal(texto)
    except InvalidOperation:
        raise Parada("outro: valor monetario desconhecido no relatorio") from None
    if not resultado.is_finite():
        raise Parada("outro: valor monetario nao finito")
    return resultado


def data_hora(valor) -> datetime:
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=FUSO) if valor.tzinfo is None else valor.astimezone(FUSO)
    try:
        momento = datetime.fromisoformat(str(valor).strip().replace("Z", "+00:00"))
    except ValueError:
        pass
    else:
        return momento.replace(tzinfo=FUSO) if momento.tzinfo is None else momento.astimezone(FUSO)
    for formato in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%d/%m/%Y",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(str(valor).strip(), formato).replace(tzinfo=FUSO)
        except ValueError:
            continue
    raise Parada("outro: formato de data desconhecido no relatorio")


def modalidade(texto: str) -> str:
    texto = normalizar(texto)
    for marca, rotulo in (
        ("pix", "PIX"),
        ("vale", "VALE"),
        ("desconto", "DESCONTO"),
        ("credito", "CREDITO"),
        ("cartao", "CREDITO"),
        ("pos", "CREDITO_POS"),
        ("carteira", "CARTEIRA_DIGITAL"),
    ):
        if marca in texto:
            return rotulo
    return "NAO_IDENTIFICADA"


def ler_relatorio(caminho: Path, inicio: date, fim: date) -> tuple[list[dict], dict]:
    if caminho.suffix.lower() == ".xlsx":
        workbook = load_workbook(caminho, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = (
                workbook["Conferência"]
                if "Conferência" in workbook.sheetnames
                else workbook.worksheets[0]
            )
            tabela = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    elif caminho.suffix.lower() == ".csv":
        payload = caminho.read_bytes()
        try:
            texto = payload.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = payload.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(texto[:8192], delimiters=";,")
        except csv.Error:
            raise Parada("outro: delimitador CSV nao confirmado") from None
        tabela = list(csv.reader(io.StringIO(texto), dialect))
    else:
        raise Parada("outro: formato exportado diferente de XLSX/CSV")
    cabecalho = None
    for indice, linha in enumerate(tabela[:30]):
        rotulos = [normalizar(str(celula or "").strip()) for celula in linha]
        if all(normalizar(coluna) in rotulos for coluna in COLUNAS):
            cabecalho = indice
            break
    if cabecalho is None:
        raise Parada("outro: layout do relatorio diferente da referencia; original preservado")
    indices = {coluna: rotulos.index(normalizar(coluna)) for coluna in COLUNAS}
    registros = []
    formas, statuses, horas = set(), set(), []
    totais = defaultdict(Decimal)
    for linha in tabela[cabecalho + 1 :]:
        if not any(celula is not None and str(celula).strip() for celula in linha):
            continue
        if len(linha) <= max(indices.values()):
            raise Parada("outro: linha incompleta no relatorio")
        momento = data_hora(linha[indices["Data/Hora da transação"]])
        if not inicio <= momento.date() <= fim:
            raise Parada("outro: exportacao fora do periodo filtrado")
        forma = str(linha[indices["Forma de Pagamento"]] or "").strip()
        status = str(linha[indices["Status"]] or "").strip()
        if not forma or not status or re.search(r"\d{5,}|@", forma + status):
            raise Parada("outro: forma/status inesperado; valores nao exibidos")
        valor = dinheiro(linha[indices["Valor líquido"]])
        registros.append(
            {
                "data_hora": momento.isoformat(),
                "forma": forma,
                "valor_liquido": str(valor),
                "codigo": str(linha[indices["Código Transação"]] or ""),
                "servico": str(linha[indices["Produto"]] or ""),
            }
        )
        formas.add(forma)
        statuses.add(status)
        horas.append(momento.strftime("%H:%M:%S"))
        totais[forma] += valor
    return registros, {
        "arquivo": caminho.name,
        "tipo": caminho.suffix.lower(),
        "colunas": list(COLUNAS),
        "linhas": len(registros),
        "formas": sorted(formas),
        "status": sorted(statuses),
        "horarios": [min(horas), max(horas)] if horas else [],
        "totais_por_forma": {forma: str(valor) for forma, valor in totais.items()},
    }


def buscar_erp(inicio: date, fim: date) -> dict:
    sys.path.insert(0, str(ROOT))
    from src.modules.cash_reconciliation.config import carregar_unidades

    unidade = carregar_unidades()[11495]
    chave = os.getenv(unidade.chave_env, "").strip()
    if not chave:
        raise Parada("outro: credencial do ERP ausente; portal salvo, comparacao indisponivel")
    base = os.getenv("WEBPOSTO_BASE_URL") or "https://web.qualityautomacao.com.br"
    resultado = defaultdict(lambda: {"quantidade": 0, "total": Decimal(0)})
    with httpx.Client(base_url=base, timeout=60) as cliente:

        def paginar(rota):
            linhas, cursor, vistos = [], None, set()
            for _ in range(100):
                params = {
                    "CHAVE": chave,
                    "empresaCodigo": 11495,
                    "limite": 1000,
                    "dataInicial": inicio.isoformat(),
                    "dataFinal": fim.isoformat(),
                }
                if cursor is not None:
                    params["ultimoCodigo"] = cursor
                resposta = cliente.get("/INTEGRACAO/V1/" + rota, params=params)
                if resposta.status_code != 200:
                    raise Parada(f"outro: ERP HTTP {resposta.status_code}; resposta nao exibida")
                payload = resposta.json()
                lote = payload.get("resultados")
                if not isinstance(lote, list):
                    raise Parada("outro: estrutura inesperada do ERP")
                linhas.extend(lote)
                novo = payload.get("ultimoCodigo")
                if not lote or novo is None:
                    return linhas
                if novo in vistos:
                    raise Parada("outro: cursor repetido no ERP")
                vistos.add(novo)
                cursor = novo
            raise Parada("outro: limite de paginacao do ERP atingido")

        cartoes = paginar("CARTOES")
        formas = paginar("VENDAS_FORMA_PAGAMENTO")
        vendas = paginar("VENDAS")
        dias_vendas = {
            venda["vendaCodigo"]: data_hora(venda["dataHora"]).date()
            for venda in vendas
            if venda.get("empresaCodigo") == 11495
        }
        for fonte, linhas in (("CARTOES", cartoes), ("VENDAS_FORMA_PAGAMENTO", formas)):
            for linha in linhas:
                if linha.get("empresaCodigo") != 11495:
                    continue
                nome = (
                    str(
                        linha.get("adiministradoraDescricao")
                        or linha.get("administradoraDescricao")
                        or ""
                    )
                    if fonte == "CARTOES"
                    else str(linha.get("nomeFormaPagamento") or "")
                )
                if "premmia" not in normalizar(nome):
                    continue
                dia = (
                    data_hora(linha["dataMovimento"]).date()
                    if fonte == "CARTOES"
                    else dias_vendas.get(linha.get("vendaCodigo"))
                )
                if dia is None or not inicio <= dia <= fim:
                    raise Parada("outro: recebimento Premmia do ERP sem data valida")
                item = resultado[(dia.isoformat(), fonte, modalidade(nome))]
                item["quantidade"] += 1
                item["total"] += dinheiro(
                    linha["valor" if fonte == "CARTOES" else "valorPagamento"]
                )
    return resultado


def executar(inicio: date, fim: date, dias: list[date]) -> None:
    usuario = os.getenv("PREMMIA_usuario_POSTO_VIP", "")
    senha = os.getenv("PREMMIA_SENHA_POSTO_VIP", "")
    if not usuario or not senha:
        raise Parada("outro: variaveis de credenciais Premmia ausentes")
    pasta = DESTINO / dias[0].isoformat() / "11495"
    pasta.mkdir(parents=True, exist_ok=True)
    execucao = datetime.now(FUSO).strftime("%Y%m%d_%H%M%S")
    erros_reserva: list[str] = []
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False, slow_mo=300)
        context = browser.new_context(
            accept_downloads=True, locale="pt-BR", timezone_id="America/Recife"
        )
        page = context.new_page()
        page.set_default_timeout(TIMEOUT)
        page.set_default_navigation_timeout(TIMEOUT)
        etapa = "abrir login"
        try:
            resposta = page.goto(LOGIN, wait_until="commit")
            if resposta is None or resposta.status >= 400:
                raise Parada("outro: portal indisponivel na pagina de login")
            etapa = "carregar formulario de login"
            aguardar(
                page, lambda: page.locator("input[data-testid='Login']").is_visible(), "formulario"
            )
            verificar_pagina(page)
            unico(page.locator("input[data-testid='Login']"), "usuario").fill(usuario)
            unico(page.locator("input[data-testid='Senha']"), "senha").fill(senha)
            verificar_pagina(page)
            unico(page.get_by_role("button", name="Entrar", exact=True), "Entrar").click()
            etapa = "confirmar login"
            aguardar(page, lambda: "/login" not in page.url, "login")
            print("Resultado do login: passou", flush=True)
            page.on(
                "response",
                lambda resposta: salvar_resposta_pagamentos(
                    resposta, pasta, execucao, erros_reserva
                ),
            )
            etapa = "abrir Conferencia/Pagamentos"
            print("Etapa: abrir Conferencia/Pagamentos.", flush=True)
            resposta = page.goto(PAGAMENTOS, wait_until="commit")
            verificar_pagina(page)
            if resposta is None or resposta.status >= 400 or "/login" in page.url:
                raise Parada("outro: Conferencia/Pagamentos indisponivel")
            etapa = "aguardar products da URL inicial"
            aguardar(
                page,
                lambda: bool(parse_qs(urlsplit(page.url).query).get("products")),
                "products da URL inicial",
            )
            url_original = page.url
            periodos = [(min(dias), max(dias))]
            if min(dias) != max(dias):
                periodos.append((min(dias), min(dias)))
            for tentativa, (inicio_exportado, fim_exportado) in enumerate(periodos):
                try:
                    etapa = "navegar para URL filtrada"
                    print(
                        f"Etapa: filtro pela URL {inicio_exportado:%d/%m/%Y} a "
                        f"{fim_exportado:%d/%m/%Y}; products preservado, sem campos de data.",
                        flush=True,
                    )
                    url_filtrada = url_periodo(url_original, inicio_exportado, fim_exportado)
                    resposta = page.goto(url_filtrada, wait_until="commit")
                    if resposta is None or resposta.status >= 400:
                        raise Parada("outro: filtro por URL indisponivel")
                    etapa = "aguardar networkidle e Exportar visivel/habilitado"
                    print(f"Etapa: {etapa} (limite 180 s; tabela nao exigida).", flush=True)
                    aguardar_exportacao(page)
                    atuais = parse_qs(urlsplit(page.url).query)
                    desejados = parse_qs(urlsplit(url_filtrada).query)
                    if any(
                        atuais.get(chave) != desejados.get(chave)
                        for chave in ("start", "end", "page", "products")
                    ):
                        raise Parada("outro: portal nao preservou os parametros do filtro")
                    verificar_pagina(page)
                    etapa = "clicar Exportar e aguardar download XLSX"
                    print(f"Etapa: {etapa} (limite 180 s).", flush=True)
                    with page.expect_download(timeout=180_000) as download_info:
                        unico(
                            page.get_by_role("button", name="Exportar", exact=True), "Exportar"
                        ).click()
                    verificar_pagina(page)
                    download = download_info.value
                    if Path(download.suggested_filename).suffix.lower() != ".xlsx":
                        raise Parada("outro: exportacao diferente de XLSX")
                    etapa = "salvar XLSX original"
                    original = pasta / (
                        f"original_{inicio_exportado}_{fim_exportado}_{execucao}.xlsx"
                    )
                    if original.exists():
                        raise Parada("outro: arquivo original ja existe; sem sobrescrever")
                    download.save_as(original)
                    print("XLSX original salvo localmente.", flush=True)
                    break
                except PlaywrightTimeoutError:
                    verificar_pagina(page)
                    if tentativa + 1 == len(periodos):
                        raise
                    print(
                        f"Timeout na etapa {etapa}; fallback autorizado para somente "
                        f"{min(dias):%d/%m/%Y}, na mesma sessao, uma vez.",
                        flush=True,
                    )
            dias = [dia for dia in dias if inicio_exportado <= dia <= fim_exportado]
            inicio, fim = inicio_exportado, fim_exportado
            if erros_reserva:
                raise Parada(
                    f"outro: falha ao salvar reserva JSON ({erros_reserva[0]}); XLSX preservado"
                )
        except (Parada, PlaywrightError) as exc:
            # Oculta todo o conteudo para nao persistir credenciais ou dados pessoais.
            try:
                captura = pasta / f"parada_{datetime.now(FUSO):%Y%m%d_%H%M%S}.png"
                page.screenshot(path=str(captura), mask=[page.locator("body")])
            except PlaywrightError:
                print("Captura indisponivel; navegador sera fechado.", flush=True)
            if isinstance(exc, PlaywrightError):
                raise Parada(f"outro: {type(exc).__name__} na etapa {etapa}") from None
            raise
        finally:
            context.close()
            browser.close()
    try:
        registros, estrutura = ler_relatorio(original, min(dias), max(dias))
    except (Parada, OSError, ValueError, KeyError) as exc:
        raise Parada(
            f"outro: etapa ler/validar layout do XLSX ({type(exc).__name__}); original preservado"
        ) from None
    for dia in dias:
        destino = DESTINO / dia.isoformat() / "11495"
        destino.mkdir(parents=True, exist_ok=True)
        if destino != pasta:
            shutil.copyfile(original, destino / original.name)
        diario = [item for item in registros if data_hora(item["data_hora"]).date() == dia]
        totais_dia = defaultdict(Decimal)
        for item in diario:
            totais_dia[item["forma"]] += dinheiro(item["valor_liquido"])
        with (destino / "normalizado.json").open("x", encoding="utf-8") as arquivo:
            json.dump(
                {
                    "empresa_codigo": 11495,
                    "dia": dia.isoformat(),
                    "registros": diario,
                    "periodo_original": [min(dias).isoformat(), max(dias).isoformat()],
                },
                arquivo,
                ensure_ascii=False,
                indent=2,
            )
        print(
            json.dumps(
                {
                    "dia": dia.strftime("%d/%m/%Y"),
                    "linhas": len(diario),
                    "totais_por_forma": {forma: str(valor) for forma, valor in totais_dia.items()},
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
    print(json.dumps(estrutura, ensure_ascii=False, indent=2), flush=True)
    erp = buscar_erp(inicio, fim)
    portal = defaultdict(lambda: {"quantidade": 0, "total": Decimal(0)})
    for dia in (inicio + timedelta(days=indice) for indice in range((fim - inicio).days + 1)):
        caminho = DESTINO / dia.isoformat() / "11495" / "normalizado.json"
        if not caminho.is_file():
            raise Parada("outro: dia sem arquivo normalizado para comparacao")
        for linha in json.loads(caminho.read_text(encoding="utf-8"))["registros"]:
            agregado = portal[(dia.isoformat(), modalidade(linha["forma"]))]
            agregado["quantidade"] += 1
            agregado["total"] += dinheiro(linha["valor_liquido"])
        formas = {forma for data, _, forma in erp if data == dia.isoformat()} | {
            forma for data, forma in portal if data == dia.isoformat()
        }
        for forma in sorted(formas):
            agregado_portal = portal[(dia.isoformat(), forma)]
            for fonte in ("CARTOES", "VENDAS_FORMA_PAGAMENTO"):
                agregado_erp = erp[(dia.isoformat(), fonte, forma)]
                print(
                    json.dumps(
                        {
                            "dia": dia.strftime("%d/%m/%Y"),
                            "forma": forma,
                            "fonte_erp": fonte,
                            "portal_qtd": agregado_portal["quantidade"],
                            "portal_total": str(agregado_portal["total"]),
                            "erp_qtd": agregado_erp["quantidade"],
                            "erp_total": str(agregado_erp["total"]),
                            "diferenca": str(agregado_portal["total"] - agregado_erp["total"]),
                            "observacao": "fontes ERP separadas, nao somadas para evitar dupla contagem",
                        },
                        ensure_ascii=False,
                    ),
                    flush=True,
                )


def ler_data(texto: str) -> date:
    try:
        dia = datetime.strptime(texto, "%d/%m/%Y").date()
    except ValueError:
        raise argparse.ArgumentTypeError("Use dd/mm/aaaa.") from None
    if dia.strftime("%d/%m/%Y") != texto:
        raise argparse.ArgumentTypeError("Use dd/mm/aaaa.")
    return dia


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--posto", type=int, choices=[11495], required=True)
    parser.add_argument("--inicio", type=ler_data)
    parser.add_argument("--fim", type=ler_data)
    args = parser.parse_args()
    if bool(args.inicio) != bool(args.fim):
        parser.error("--inicio e --fim devem ser informados juntos.")
    hoje = datetime.now(FUSO).date()
    inicio = args.inicio or hoje.replace(day=1)
    fim = args.fim or hoje - timedelta(days=1)
    if fim >= hoje or inicio > fim or (fim - inicio).days >= 31:
        parser.error("Periodo deve ter no maximo 31 dias e terminar ate ontem.")
    dias = [
        inicio + timedelta(days=indice)
        for indice in range((fim - inicio).days + 1)
        if not (
            DESTINO / (inicio + timedelta(days=indice)).isoformat() / "11495" / "normalizado.json"
        ).is_file()
    ]
    if not dias:
        print("Todos os dias ja possuem dados salvos; nenhum login realizado.")
        return 0
    if (max(dias) - min(dias)).days + 1 != len(dias):
        print("Parada: dias faltantes descontiguos; nao exportar dias ja salvos.")
        return 1
    load_dotenv(ROOT / ".env")
    if os.getenv("WEBPOSTO_WRITES", "0").strip() != "0":
        print("Parada: WEBPOSTO_WRITES nao esta desativado; nenhum login realizado.")
        return 1
    for nome in ("httpx", "httpcore"):
        logging.getLogger(nome).setLevel(logging.WARNING)
    print(
        f"Teste unico: VIP 11495, {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}, {len(dias)} dias faltantes.",
        flush=True,
    )
    try:
        executar(inicio, fim, dias)
    except Parada as exc:
        print(f"Parada controlada: {exc}", flush=True)
        return 1
    except (PlaywrightError, httpx.HTTPError, OSError, ValueError, KeyError, csv.Error) as exc:
        print(
            f"Parada controlada: outro ({type(exc).__name__}); detalhes omitidos por privacidade.",
            flush=True,
        )
        return 1
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
