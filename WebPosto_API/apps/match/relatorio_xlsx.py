"""Planilha do dia: um .xlsx por posto/dia, montado a partir do banco local."""
from __future__ import annotations

import glob
import hashlib
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.match import caminhos, coletas
from src.modules.cash_reconciliation.domain.tempo import FUSO

FORMATO_REAIS = "R$ #,##0.00"
_ROTULO_FONTE = {"premmia": "Premmia", "pagbank_edi": "PagBank"}
_ROTULO_SITUACAO = {"ok": "OK", "sem_movimento": "Sem movimento", "erro": "Erro",
                    "verificacao_humana": "Precisa de verificação humana"}
_MSG_SITUACAO = {"ok": "Coleta concluída", "sem_movimento": "Sem vendas registradas no dia"}
_SEM_COMPARACAO = ". Sem comparação com webPosto (cadastre a chave do webPosto em Fontes do posto)."


def _valor(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(Decimal(str(v)))
    except InvalidOperation:
        return None


def _local(iso: str) -> datetime:
    dt = datetime.fromisoformat(iso)
    return dt.astimezone(FUSO) if dt.tzinfo else dt.replace(tzinfo=FUSO)


def _cabecalho(ws, titulos: list[str]) -> None:
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _ajustar(ws) -> None:
    for i, col in enumerate(ws.columns, 1):
        largura = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)


def _moeda(ws, coluna: int) -> None:
    for row in ws.iter_rows(min_row=2, min_col=coluna, max_col=coluna):
        for c in row:
            if c.value is not None:
                c.number_format = FORMATO_REAIS


def _resumo(con, wb, posto_id: int, dia: date) -> None:
    ws = wb.active
    ws.title = "Resumo"
    _cabecalho(ws, ["Fonte", "Situação", "Quantidade", "Total", "Mensagem"])
    rows = con.execute(
        "SELECT f.tipo, c.situacao, c.qtd, c.total, c.erro_texto FROM coleta c"
        " JOIN fonte f ON f.id=c.fonte_id WHERE f.posto_id=? AND c.dia=? ORDER BY f.id",
        (posto_id, str(dia))).fetchall()
    posto = con.execute("SELECT * FROM posto WHERE id=?", (posto_id,)).fetchone()
    compara = posto is not None and coletas.tem_chave_webposto(con, posto)
    for r in rows:
        msg = r["erro_texto"] or _MSG_SITUACAO.get(r["situacao"], "")
        if r["situacao"] == "ok" and not compara:
            msg += _SEM_COMPARACAO
        total = None if r["situacao"] in ("erro", "verificacao_humana") else _valor(r["total"])
        ws.append([_ROTULO_FONTE.get(r["tipo"], r["tipo"]),
                   _ROTULO_SITUACAO.get(r["situacao"], r["situacao"]), r["qtd"], total, msg])
    _moeda(ws, 4)
    _ajustar(ws)


def _transacoes(con, wb, nome: str, fonte: str, posto_id: int, dia: date) -> None:
    ws = wb.create_sheet(nome)
    _cabecalho(ws, ["Data/hora", "NSU", "Autorização", "Valor", "Bandeira/forma", "Situação"])
    rows = con.execute(
        "SELECT * FROM transacao WHERE posto_id=? AND fonte=? ORDER BY momento, id",
        (posto_id, fonte)).fetchall()
    for r in rows:
        local = _local(r["momento"])
        if local.date() != dia:
            continue
        ws.append([local.strftime("%d/%m/%Y %H:%M"), r["nsu"], r["autorizacao"],
                   _valor(r["valor"]), r["bandeira"] or r["forma"], r["situacao"]])
    _moeda(ws, 4)
    _ajustar(ws)


def _divergencias(con, wb, posto_id: int, dia: date) -> None:
    ws = wb.create_sheet("Divergências")
    _cabecalho(ws, ["Tipo", "Valor esperado", "Valor encontrado", "NSU", "Caixa",
                    "Responsável", "O que houve", "O que fazer"])
    for r in con.execute("SELECT * FROM divergencia WHERE posto_id=? AND dia=? ORDER BY id",
                         (posto_id, str(dia))).fetchall():
        ws.append([r["tipo"], _valor(r["valor_esperado"]), _valor(r["valor_encontrado"]),
                   r["nsu"], r["caixa_codigo"], r["responsavel"], r["explicacao_pt"],
                   r["correcao_pt"]])
    _moeda(ws, 2)
    _moeda(ws, 3)
    _ajustar(ws)


def _nome_seguro(apelido: str | None, posto_id: int) -> str:
    limpo = re.sub(r"[^\w -]", "_", apelido or "")
    limpo = re.sub(r"\s+", "_", limpo.strip()).strip("_")
    return limpo or f"P{posto_id}"


PADRAO_ARQUIVO = re.compile(r"^(Logos_.+_\d{2}-\d{2}-\d{4})(?:_v(\d+))?\.xlsx$")


def _base(apelido: str | None, posto_id: int, dia: date) -> str:
    return f"Logos_{_nome_seguro(apelido, posto_id)}_{dia.strftime('%d-%m-%Y')}"


def _numero_versao(base: str, arq: Path) -> int | None:
    m = PADRAO_ARQUIVO.match(arq.name)
    if not m or m.group(1) != base:
        return None
    return int(m.group(2)) if m.group(2) else 1


def versoes(posto_id: int, dia: date, apelido: str | None) -> list[Path]:
    """Versões do dia, da mais antiga para a mais nova (a base é a versão 1)."""
    base = _base(apelido, posto_id, dia)
    achadas = []
    for arq in caminhos.planilhas().glob(f"{glob.escape(base)}*.xlsx"):
        n = _numero_versao(base, arq)
        if n is not None:
            achadas.append((n, arq))
    return [a for _, a in sorted(achadas)]


def _apelido(con, posto_id: int) -> str | None:
    posto = con.execute("SELECT apelido FROM posto WHERE id=?", (posto_id,)).fetchone()
    return posto["apelido"] if posto else None


def mais_recente(con, posto_id: int, dia: date) -> Path | None:
    achadas = versoes(posto_id, dia, _apelido(con, posto_id))
    return achadas[-1] if achadas else None


def _impressao(wb) -> str:
    h = hashlib.sha256()
    for ws in wb.worksheets:
        h.update(repr(ws.title).encode())
        for linha in ws.iter_rows(values_only=True):
            h.update(repr(linha).encode())
    return h.hexdigest()


def _mesmos_dados(wb, arq: Path) -> bool:
    try:
        antigo = load_workbook(arq, read_only=True)
        try:
            return _impressao(antigo) == _impressao(wb)
        finally:
            antigo.close()
    except Exception:  # noqa: BLE001 - arquivo ilegível: gera nova versão
        return False


def gerar(con, posto_id: int, dia: date, destino: Path | None = None) -> Path:
    apelido = _apelido(con, posto_id)
    wb = Workbook()
    _resumo(con, wb, posto_id, dia)
    _transacoes(con, wb, "Premmia", "premmia", posto_id, dia)
    _transacoes(con, wb, "PagBank", "pagbank_edi", posto_id, dia)
    _divergencias(con, wb, posto_id, dia)
    if destino is not None:
        wb.save(destino)
        return destino
    existentes = versoes(posto_id, dia, apelido)
    if existentes and _mesmos_dados(wb, existentes[-1]):
        return existentes[-1]
    base = _base(apelido, posto_id, dia)
    proxima = _numero_versao(base, existentes[-1]) + 1 if existentes else 1
    destino = caminhos.planilhas() / (f"{base}.xlsx" if proxima == 1 else f"{base}_v{proxima}.xlsx")
    wb.save(destino)
    return destino
