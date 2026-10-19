from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace

import openpyxl
import pytest

from apps.match import banco, credenciais, relatorio_xlsx
from src.modules.cash_reconciliation.domain.tempo import FUSO

DIA = date(2026, 10, 17)


@pytest.fixture
def con(tmp_path, monkeypatch):
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    c = banco.conectar(tmp_path / "t.db")
    yield c
    c.close()


def _t(ident, nsu, valor, momento=datetime(2026, 10, 17, 14, 5, tzinfo=FUSO)):
    return SimpleNamespace(identificador=ident, nsu=nsu, autorizacao="A1", valor=Decimal(valor),
                           momento=momento, bandeira="VISA",
                           terminal="T1", situacao="aprovada")


@pytest.fixture
def posto(con):
    pid = banco.inserir_posto(con, "Posto Doze", "Doze", None, 74014)
    fp = banco.salvar_fonte(con, pid, "premmia", True, {})
    fb = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    cp = banco.registrar_coleta(con, fp, DIA, "ok", 1, Decimal("100.50"), None)
    banco.registrar_coleta(con, fb, DIA, "erro", 0, Decimal("0"), "Falha ao entrar")
    banco.gravar_transacoes(con, cp, pid, "premmia", [_t("p1", "111", "100.50")])
    banco.gravar_divergencias(con, pid, DIA, "premmia", [{
        "tipo": "nsu_ausente", "valor_esperado": Decimal("100.50"), "valor_encontrado": None,
        "nsu": "111", "caixa_codigo": 7, "responsavel": "JOÃO (42)",
        "explicacao_pt": "Venda sem par", "correcao_pt": "Conferir no portal"}])
    return pid


def test_abas_e_nome_do_arquivo(con, posto, tmp_path):
    p = relatorio_xlsx.gerar(con, posto, DIA)
    assert p.name == "Logos_Doze_17-10-2026.xlsx"
    assert p.parent == tmp_path / "dados" / "planilhas"
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Resumo", "Premmia", "PagBank", "Divergências"]
    resumo = [[c.value for c in r] for r in wb["Resumo"].iter_rows()]
    assert ["Premmia", "OK", 1, 100.5] in [r[:4] for r in resumo]
    assert resumo[1][4].startswith("Coleta concluída")  # sem chave: + aviso de sem comparação
    assert ["PagBank", "Erro", 0, None, "Falha ao entrar"] in resumo


def test_data_no_formato_brasil(con, posto):
    wb = openpyxl.load_workbook(relatorio_xlsx.gerar(con, posto, DIA))
    assert wb["Premmia"]["A2"].value == "17/10/2026 14:05"
    assert wb["Premmia"]["B2"].value == "111"


def test_valor_numerico_formatado(con, posto):
    wb = openpyxl.load_workbook(relatorio_xlsx.gerar(con, posto, DIA))
    c = wb["Premmia"]["D2"]
    assert c.value == 100.5 and c.number_format == "R$ #,##0.00"
    d = wb["Divergências"]["B2"]
    assert d.value == 100.5 and d.number_format == "R$ #,##0.00"
    assert wb["Divergências"]["F2"].value == "JOÃO (42)"


def test_sem_segredo_na_planilha(con, posto, tmp_path):
    credenciais.salvar("premmia-1", "S3nh@-Ultra-Secreta")
    wb = openpyxl.load_workbook(relatorio_xlsx.gerar(con, posto, DIA, tmp_path / "x.xlsx"))
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert "S3nh@-Ultra-Secreta" not in str(c.value)


def test_apelido_com_caracteres_invalidos(con, tmp_path):
    pid = banco.inserir_posto(con, "Real", "Real S/A: 24h?", None, 1)
    p = relatorio_xlsx.gerar(con, pid, DIA)
    assert p.name == "Logos_Real_S_A__24h_17-10-2026.xlsx"
    assert not set("/:?") & set(p.name)
    assert p.parent == tmp_path / "dados" / "planilhas" and p.exists()


def test_apelido_ausente_usa_id(con):
    pid = banco.inserir_posto(con, "Sem", None, None, 1)
    assert relatorio_xlsx.gerar(con, pid, DIA).name == f"Logos_P{pid}_17-10-2026.xlsx"


def test_so_transacoes_do_dia_em_recife(con):
    pid = banco.inserir_posto(con, "X", "X", None, 1)
    f = banco.salvar_fonte(con, pid, "premmia", True, {})
    c = banco.registrar_coleta(con, f, DIA, "ok", 3, Decimal("3"), None)
    banco.gravar_transacoes(con, c, pid, "premmia", [
        _t("a", "1", "1", datetime(2026, 10, 16, 23, 59, tzinfo=FUSO)),
        _t("b", "2", "1", datetime(2026, 10, 17, 23, 59, tzinfo=FUSO)),
        _t("c", "3", "1", datetime(2026, 10, 18, 0, 1, tzinfo=FUSO))])
    ws = openpyxl.load_workbook(relatorio_xlsx.gerar(con, pid, DIA))["Premmia"]
    assert [r[0].value for r in ws.iter_rows(min_row=2)] == ["17/10/2026 23:59"]


def test_sem_movimento(con):
    pid = banco.inserir_posto(con, "X", "X", None, 1)
    f = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    banco.registrar_coleta(con, f, DIA, "sem_movimento", 0, Decimal("0"), None)
    ws = openpyxl.load_workbook(relatorio_xlsx.gerar(con, pid, DIA))["Resumo"]
    assert [c.value for c in ws[2]] == ["PagBank", "Sem movimento", 0, 0, "Sem vendas registradas no dia"]


def test_resumo_avisa_sem_comparacao_quando_sem_chave(con, posto):
    wb = openpyxl.load_workbook(relatorio_xlsx.gerar(con, posto, DIA))
    msgs = [r[4].value for r in wb["Resumo"].iter_rows(min_row=2)]
    assert "Sem comparação com webPosto" in msgs[0]
    banco.salvar_fonte(con, posto, "webposto", True, {})
    credenciais.salvar(f"{posto}/webposto/chave", "K")
    wb = openpyxl.load_workbook(relatorio_xlsx.gerar(con, posto, DIA))
    msgs = [r[4].value for r in wb["Resumo"].iter_rows(min_row=2)]
    assert msgs[0] == "Coleta concluída"


def _muda(con, posto, valor):
    banco.gravar_divergencias(con, posto, DIA, "premmia", [{
        "tipo": "valor_diferente", "valor_esperado": Decimal(valor), "valor_encontrado": None,
        "nsu": "222", "caixa_codigo": 7, "responsavel": "JOÃO (42)",
        "explicacao_pt": "Outra", "correcao_pt": "Conferir"}])


def test_primeira_versao_sem_sufixo(con, posto):
    p = relatorio_xlsx.gerar(con, posto, DIA)
    assert p.name == "Logos_Doze_17-10-2026.xlsx"
    assert relatorio_xlsx.versoes(posto, DIA, "Doze") == [p]
    assert relatorio_xlsx.mais_recente(con, posto, DIA) == p


def test_dados_iguais_reaproveita_ultima(con, posto):
    p1 = relatorio_xlsx.gerar(con, posto, DIA)
    p2 = relatorio_xlsx.gerar(con, posto, DIA)
    assert p1 == p2
    assert len(list(p1.parent.glob("*.xlsx"))) == 1


def test_dados_mudaram_cria_v2_e_v3(con, posto):
    p1 = relatorio_xlsx.gerar(con, posto, DIA)
    antes = p1.read_bytes()
    _muda(con, posto, "10")
    p2 = relatorio_xlsx.gerar(con, posto, DIA)
    assert p2.name == "Logos_Doze_17-10-2026_v2.xlsx"
    _muda(con, posto, "20")
    p3 = relatorio_xlsx.gerar(con, posto, DIA)
    assert p3.name == "Logos_Doze_17-10-2026_v3.xlsx"
    assert p1.read_bytes() == antes
    assert relatorio_xlsx.versoes(posto, DIA, "Doze") == [p1, p2, p3]
    assert relatorio_xlsx.mais_recente(con, posto, DIA) == p3


def test_versao_aberta_no_excel_nao_impede_nova(con, posto):
    p1 = relatorio_xlsx.gerar(con, posto, DIA)
    with open(p1, "rb"):
        _muda(con, posto, "10")
        p2 = relatorio_xlsx.gerar(con, posto, DIA)
    assert p2.name.endswith("_v2.xlsx")


def test_destino_explicito_nao_versiona(con, posto, tmp_path):
    alvo = tmp_path / "x.xlsx"
    assert relatorio_xlsx.gerar(con, posto, DIA, alvo) == alvo and alvo.exists()
    assert relatorio_xlsx.mais_recente(con, posto, DIA) is None


def test_valores_inteiros_e_zero_reaproveitam(con):
    pid = banco.inserir_posto(con, "Posto Doze", "Doze", None, 74014)
    fp = banco.salvar_fonte(con, pid, "premmia", True, {})
    banco.registrar_coleta(con, fp, DIA, "ok", 1, Decimal("100"), None)
    fb = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    banco.registrar_coleta(con, fb, DIA, "ok", 0, Decimal("0"), None)
    banco.gravar_divergencias(con, pid, DIA, "premmia", [{
        "tipo": "x", "valor_esperado": Decimal("100"), "valor_encontrado": Decimal("0")}])
    p1 = relatorio_xlsx.gerar(con, pid, DIA)
    assert relatorio_xlsx.gerar(con, pid, DIA) == p1
    assert len(list(p1.parent.glob("*.xlsx"))) == 1


def test_nome_em_uso_por_apelido_equivalente(con):
    a = banco.inserir_posto(con, "A", "Doze Filial", None, 1)
    assert relatorio_xlsx.nome_em_uso(con, "Doze_Filial", None)
    assert relatorio_xlsx.nome_em_uso(con, "doze filial", None)
    assert not relatorio_xlsx.nome_em_uso(con, "Doze Filial", a)  # o próprio posto
    assert not relatorio_xlsx.nome_em_uso(con, "Outro", None)


def test_nome_em_uso_fallback_pid(con):
    banco.inserir_posto(con, "A", None, None, 1)  # id 1 -> P1
    assert relatorio_xlsx.nome_em_uso(con, "P1", None)
    assert relatorio_xlsx.nome_em_uso(con, None, 1) is False
    b = banco.inserir_posto(con, "B", "P1", None, 2)  # já colide, banco aceita
    assert relatorio_xlsx.conflito(con, b) and relatorio_xlsx.conflito(con, 1)
    assert relatorio_xlsx.mais_recente(con, b, DIA) is None
