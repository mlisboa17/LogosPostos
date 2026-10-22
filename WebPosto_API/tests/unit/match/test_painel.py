import json
import re
import threading

import pytest
from fastapi.testclient import TestClient

from apps.match import acesso, banco, coletas, credenciais, rotina
from apps.match.acesso import Perfil
from apps.match.painel import criar_app, rotas
from src.modules.cash_reconciliation import config as cfg
from src.modules.cash_reconciliation.domain.tempo import ontem
from src.shared import segredos

SENHA = "senha-forte-1"
SEGREDO = "SEGREDO-NAO-MOSTRAR-123"


@pytest.fixture(autouse=True)
def _limpa():
    yield
    segredos.definir_resolvedor(None)
    cfg.definir_provedor_unidades(None)


@pytest.fixture
def fab(tmp_path):
    caminho = tmp_path / "m.db"
    banco.conectar(caminho).close()
    return lambda: banco.conectar(caminho)


@pytest.fixture
def con(fab):
    c = fab()
    yield c
    c.close()


@pytest.fixture
def cli(fab):
    return TestClient(criar_app(fab), base_url="http://127.0.0.1", follow_redirects=False)


@pytest.fixture
def admin(con):
    return acesso.criar_usuario(con, "Admin", "admin", SENHA, Perfil.ADMINISTRADOR, [])


@pytest.fixture
def executar_falso(monkeypatch):
    chamadas, feito = [], threading.Event()

    async def falso(con, **kw):
        chamadas.append(kw)
        feito.set()
        return []

    monkeypatch.setattr(rotina, "executar", falso)
    return chamadas, feito


def _csrf(html: str) -> str:
    return re.search(r'name="csrf" value="([^"]+)"', html).group(1)


def _entrar(cli, login="admin", senha=SENHA):
    tok = _csrf(cli.get("/entrar").text)
    return cli.post("/entrar", data={"login": login, "senha": senha, "csrf": tok})


def _logado(cli, login="admin") -> str:
    r = _entrar(cli, login)
    assert r.status_code == 303, r.text
    return _csrf(cli.get("/").text)


def _posto(con, nome="Doze", codigo=74014) -> int:
    return banco.inserir_posto(con, nome, nome, None, codigo)


def test_sem_login_redireciona(cli, admin):
    for rota in ("/", "/postos", "/usuarios", "/agenda", "/historico", "/coleta", "/config/exportar"):
        r = cli.get(rota)
        assert r.status_code == 303 and r.headers["location"] == "/entrar", rota
    r = cli.post("/coletar?posto_id=1", data={"csrf": "x"})
    assert r.status_code == 303 and r.headers["location"] == "/entrar"


def test_primeiro_uso_cria_admin(cli, con):
    r = cli.get("/entrar")
    assert r.status_code == 303 and r.headers["location"] == "/primeiro-uso"
    tok = _csrf(cli.get("/primeiro-uso").text)
    r = cli.post("/primeiro-uso", data={"nome": "Dono", "login": "dono", "senha": SENHA,
                                        "senha2": SENHA, "csrf": tok})
    assert r.status_code == 303 and r.headers["location"] == "/entrar"
    assert not acesso.precisa_primeiro_uso(con)
    assert cli.get("/primeiro-uso").headers["location"] == "/entrar"
    r = cli.post("/primeiro-uso", data={"nome": "Intruso", "login": "intruso", "senha": SENHA,
                                        "senha2": SENHA, "csrf": tok})
    assert r.status_code == 403
    assert con.execute("SELECT COUNT(*) FROM usuario").fetchone()[0] == 1
    assert _entrar(cli, "dono").status_code == 303


def test_primeiro_uso_sem_csrf_recusa(cli, con):
    r = cli.post("/primeiro-uso", data={"nome": "X", "login": "x", "senha": SENHA, "senha2": SENHA})
    assert r.status_code == 403
    assert acesso.precisa_primeiro_uso(con)


def test_login_errado_nao_entra(cli, admin):
    r = _entrar(cli, senha="errada-errada")
    assert r.status_code == 401
    assert "lm_sessao" not in r.cookies
    assert cli.get("/").headers["location"] == "/entrar"


def test_login_cookie_seguro_e_sair(cli, admin):
    r = _entrar(cli)
    cookie = r.headers["set-cookie"]
    assert "lm_sessao=" in cookie and "HttpOnly" in cookie and "SameSite=strict" in cookie.replace("Strict", "strict")
    tok = _csrf(cli.get("/").text)
    assert cli.post("/sair", data={"csrf": tok}).status_code == 303
    assert cli.get("/").headers["location"] == "/entrar"


def test_sessao_expira(cli, admin, monkeypatch):
    _logado(cli)
    agora = rotas._relogio()
    monkeypatch.setattr(rotas, "_relogio", lambda: agora + 12 * 3600 + 1)
    assert cli.get("/").headers["location"] == "/entrar"


def test_host_externo_recusado(fab, admin):
    c = TestClient(criar_app(fab), base_url="http://ataque.exemplo", follow_redirects=False)
    assert c.get("/entrar").status_code == 400


def test_post_sem_csrf_recusa(cli, admin, con):
    tok = _logado(cli)
    assert cli.post("/postos", data={"acao": "novo", "nome": "X"}).status_code == 403
    assert cli.post("/postos", data={"acao": "novo", "nome": "X", "csrf": "outro"}).status_code == 403
    assert banco.listar_postos(con, so_ativos=False) == []
    assert cli.post("/postos", data={"acao": "novo", "nome": "X", "csrf": tok}).status_code == 303
    assert [p["nome"] for p in banco.listar_postos(con)] == ["X"]


def test_gerente_posto_alheio_403(cli, admin, con, executar_falso):
    ids = [_posto(con, f"P{i}", 1000 + i) for i in range(7)]
    assert ids[-1] == 7
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [1])
    tok = _logado(cli, "gerente")
    r = cli.get("/postos/7/fontes")
    assert r.status_code == 403 and "Você não tem acesso a isto." in r.text
    r = cli.post("/coletar?posto_id=7&visivel=0", data={"csrf": tok})
    assert r.status_code == 403
    assert executar_falso[0] == []
    assert cli.get("/postos/1/fontes").status_code == 200
    for rota in ("/usuarios", "/canais", "/agenda", "/config/exportar", "/usuarios/log"):
        assert cli.get(rota).status_code == 403, rota


def test_conferente_nao_troca_credencial(cli, admin, con):
    pid = _posto(con)
    acesso.criar_usuario(con, "Cê", "conf", SENHA, Perfil.CONFERENTE, [pid])
    tok = _logado(cli, "conf")
    r = cli.post(f"/postos/{pid}/fontes/premmia", data={"csrf": tok, "ativo": "1",
                                                        "senha": SEGREDO, "usuario": "u"})
    assert r.status_code == 403
    assert not credenciais.tem(f"{pid}/premmia/senha")
    assert cli.get(f"/postos/{pid}/fontes").status_code == 403


def test_hoje_lista_so_postos_visiveis(cli, admin, con):
    a, _b = _posto(con, "Posto Alfa", 1), _posto(con, "Posto Beta", 2)
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [a])
    _logado(cli, "gerente")
    html = cli.get("/").text
    assert "Posto Alfa" in html and "Posto Beta" not in html
    menu = re.search(r"<nav.*?</nav>", html, re.S).group(0)
    assert "/usuarios" not in menu and "/canais" not in menu and "/agenda" not in menu


def test_hoje_mostra_responsavel(cli, admin, con):
    pid = _posto(con)
    banco.gravar_divergencias(con, pid, ontem(), "pagbank_edi", [{
        "tipo": "falta_no_webposto", "valor_esperado": "10", "valor_encontrado": "0",
        "responsavel": "JOÃO (42)", "explicacao_pt": "Venda <b>sem</b> lançamento.",
        "correcao_pt": "Lançar o cartão."}])
    _logado(cli)
    html = cli.get("/").text
    assert "JOÃO (42)" in html and "Lançar o cartão." in html
    assert "Venda &lt;b&gt;sem&lt;/b&gt; lançamento." in html  # autoescape


def test_log_registra_troca_de_credencial(cli, admin, con):
    pid = _posto(con)
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/fontes/premmia", data={"csrf": tok, "ativo": "1",
                                                        "usuario": "user", "senha": SEGREDO})
    assert r.status_code == 303
    assert credenciais.ler(f"{pid}/premmia/senha") == SEGREDO
    logs = con.execute("SELECT * FROM log_acesso WHERE acao='trocar_credencial'").fetchall()
    assert len(logs) == 1 and logs[0]["usuario_id"] == admin
    assert SEGREDO not in logs[0]["detalhe"] and "senha" in logs[0]["detalhe"]
    assert all(SEGREDO not in (r["detalhe"] or "") for r in con.execute("SELECT detalhe FROM log_acesso"))


def test_senha_em_branco_mantem_anterior(cli, admin, con):
    pid = _posto(con)
    credenciais.salvar(f"{pid}/pagbank/token", SEGREDO)
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/fontes/pagbank_edi", data={"csrf": tok, "ativo": "1",
                                                            "usuario": "novo", "token": ""})
    assert r.status_code == 303
    assert credenciais.ler(f"{pid}/pagbank/token") == SEGREDO
    assert credenciais.ler(f"{pid}/pagbank/usuario") == "novo"
    fonte = {f["tipo"]: f for f in banco.fontes_do_posto(con, pid)}["pagbank_edi"]
    assert fonte["ativo"] == 1


def test_salvar_fonte_preserva_config(cli, admin, con):
    pid = _posto(con)
    banco.salvar_fonte(con, pid, "premmia", False, {"administradoras": {"PREMMIA": [5]}})
    tok = _logado(cli)
    cli.post(f"/postos/{pid}/fontes/premmia", data={"csrf": tok, "ativo": "1"})
    fonte = banco.fontes_do_posto(con, pid)[0]
    assert json.loads(fonte["config_json"]) == {"administradoras": {"PREMMIA": [5]}}
    assert fonte["ativo"] == 1


def test_tela_nao_mostra_segredo(cli, admin, con):
    pid = _posto(con)
    for chave in ("premmia/usuario", "premmia/senha", "pagbank/usuario", "pagbank/token", "webposto/chave"):
        credenciais.salvar(f"{pid}/{chave}", f"{SEGREDO}-{chave}")
    _logado(cli)
    html = cli.get(f"/postos/{pid}/fontes").text
    assert SEGREDO not in html
    assert "••••" in html


def test_testar_pagbank_usa_erro_pt(cli, admin, con, monkeypatch):
    from src.modules.cash_reconciliation.adapters import pagbank_edi
    pid = _posto(con)
    vistos = []

    async def falso(emp, dia, **kw):
        vistos.append((emp, dia))
        raise pagbank_edi.PagBankCredencialInvalida(f"token {SEGREDO}")

    monkeypatch.setattr(pagbank_edi, "buscar_transacoes", falso)
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/fontes/pagbank_edi/testar", data={"csrf": tok})
    assert r.status_code == 200
    assert "PagBank recusou o usuário/token" in r.text and SEGREDO not in r.text
    assert vistos == [(74014, ontem())]


def test_testar_webposto_ok(cli, admin, con, monkeypatch):
    from src.modules.cash_reconciliation.adapters import webposto_caixas
    pid = _posto(con)
    vistos = []

    async def falso(unidade, inicio, fim, **kw):
        vistos.append((unidade.empresa_codigo, inicio, fim))
        return [object(), object()]

    monkeypatch.setattr(webposto_caixas, "buscar_caixas", falso)
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/fontes/webposto/testar", data={"csrf": tok})
    assert "2 caixa(s)" in r.text
    assert vistos == [(74014, ontem(), ontem())]


def test_testar_premmia_so_confere_existencia(cli, admin, con):
    pid = _posto(con)
    tok = _logado(cli)
    assert "Falta" in cli.post(f"/postos/{pid}/fontes/premmia/testar", data={"csrf": tok}).text
    credenciais.salvar(f"{pid}/premmia/usuario", "u")
    credenciais.salvar(f"{pid}/premmia/senha", "s")
    assert "estão salvos" in cli.post(f"/postos/{pid}/fontes/premmia/testar", data={"csrf": tok}).text


def test_coletar_inicia_em_thread(cli, admin, con, executar_falso):
    pid = _posto(con)
    tok = _logado(cli)
    r = cli.post(f"/coletar?posto_id={pid}&visivel=1", data={"csrf": tok})
    assert r.status_code == 303 and r.headers["location"].startswith("/coleta")
    chamadas, feito = executar_falso
    assert feito.wait(5)
    assert len(chamadas) == 1
    assert {k: chamadas[0][k] for k in ("posto_id", "visivel", "origem")} == {
        "posto_id": pid, "visivel": True, "origem": "painel"}
    assert chamadas[0]["usuario_id"] == admin
    assert con.execute("SELECT COUNT(*) FROM log_acesso WHERE acao='coletar'").fetchone()[0] == 1


def test_coletar_em_andamento(cli, admin, con, executar_falso):
    pid = _posto(con)
    tok = _logado(cli)
    assert rotina._TRAVA.acquire(blocking=False)
    try:
        r = cli.post(f"/coletar?posto_id={pid}&visivel=0", data={"csrf": tok})
        html = cli.get(r.headers["location"]).text
    finally:
        rotina._TRAVA.release()
    assert "Já existe uma coleta em andamento. Aguarde terminar." in html
    assert executar_falso[0] == []


def test_reenviar_usa_dados_salvos(cli, admin, con, monkeypatch):
    pid = _posto(con)
    fid = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    banco.registrar_coleta(con, fid, ontem(), "ok", 3, "30.00", None)
    enviados, feito = [], threading.Event()

    def falso(c, resultados, dia):
        enviados.append((resultados, dia))
        feito.set()

    monkeypatch.setattr(rotas.envio, "enviar_dia", falso)
    tok = _logado(cli)
    r = cli.post(f"/reenviar?posto_id={pid}", data={"csrf": tok})
    assert "Reenvio iniciado" in r.text
    assert feito.wait(5)
    resultados, dia = enviados[0]
    assert dia == ontem() and [(x.fonte, x.qtd) for x in resultados] == [("pagbank_edi", 3)]


def test_resultados_salvos_sem_coleta(con):
    pid = _posto(con)
    banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    assert coletas.resultados_salvos(con, pid, ontem()) == []


def test_exportar_sem_segredo(cli, admin, con):
    pid = _posto(con)
    banco.salvar_fonte(con, pid, "webposto", True, {"chave": SEGREDO})
    credenciais.salvar(f"{pid}/webposto/chave", SEGREDO)
    _logado(cli)
    r = cli.get("/config/exportar")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert SEGREDO not in r.text
    assert r.json()["postos"][0]["nome"] == "Doze"


def test_importar_config(cli, admin, con):
    tok = _logado(cli)
    dados = json.dumps({"postos": [{"nome": "Novo", "empresa_codigo_webposto": 9}]})
    r = cli.post("/config/importar", data={"csrf": tok},
                 files={"arquivo": ("c.json", dados, "application/json")})
    assert r.status_code == 200 and "1 posto(s)" in r.text
    assert [p["nome"] for p in banco.listar_postos(con)] == ["Novo"]
    r = cli.post("/config/importar", data={"csrf": tok},
                 files={"arquivo": ("c.json", "não é json", "application/json")})
    assert r.status_code == 400


def test_agenda_hora_invalida_recusa(cli, admin, con):
    tok = _logado(cli)
    r = cli.post("/agenda", data={"csrf": tok, "hora_coleta": "25:00", "hora_repescagem": "10:00"})
    assert r.status_code == 400 and "Horário inválido" in r.text
    assert rotina.horarios(con) == ("06:00", "10:00")
    r = cli.post("/agenda", data={"csrf": tok, "hora_coleta": "05:30", "hora_repescagem": "11:15"})
    assert r.status_code == 303
    assert rotina.horarios(con) == ("05:30", "11:15")


def test_historico_datas_brasil(cli, admin, con):
    pid = _posto(con)
    fid = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    banco.registrar_coleta(con, fid, ontem(), "erro", 0, "0", "PagBank recusou o usuário/token.")
    _logado(cli)
    html = cli.get("/historico").text
    assert f"{ontem():%d/%m/%Y}" in html and "PagBank recusou" in html
    assert re.search(r"\d{2}/\d{2}/\d{4} \d{2}:\d{2}", html)


def test_canais_cria_e_mantem_token(cli, admin, con):
    tok = _logado(cli)
    r = cli.post("/canais", data={"csrf": tok, "acao": "novo", "tipo": "telegram",
                                  "token": SEGREDO, "chat_ids": "11, 22", "ativo": "1"})
    assert r.status_code == 303
    canal = con.execute("SELECT * FROM canal").fetchone()
    assert json.loads(canal["config_json"])["chat_ids"] == ["11", "22"]
    assert SEGREDO not in canal["config_json"]
    assert credenciais.ler(f"canal/{canal['id']}/token") == SEGREDO
    cli.post("/canais", data={"csrf": tok, "acao": "editar", "id": canal["id"], "tipo": "telegram",
                              "token": "", "chat_ids": "33", "ativo": "1"})
    assert credenciais.ler(f"canal/{canal['id']}/token") == SEGREDO
    assert SEGREDO not in cli.get("/canais").text


def test_usuarios_criar_e_nao_desativar_ultimo_admin(cli, admin, con):
    pid = _posto(con)
    tok = _logado(cli)
    r = cli.post("/usuarios", data={"csrf": tok, "acao": "criar", "nome": "Gê", "login": "ge",
                                    "senha": SENHA, "perfil": "gerente", "postos": [str(pid)]})
    assert r.status_code == 303
    ge = con.execute("SELECT id FROM usuario WHERE login='ge'").fetchone()["id"]
    assert acesso.postos_visiveis_agora(con, ge) == [pid]
    r = cli.post("/usuarios", data={"csrf": tok, "acao": "desativar", "id": admin})
    assert r.status_code == 400
    assert acesso.carregar_usuario(con, admin)["ativo"] == 1
    r = cli.post("/usuarios", data={"csrf": tok, "acao": "editar", "id": admin, "nome": "Admin",
                                    "perfil": "gerente"})
    assert r.status_code == 400
    assert not acesso.precisa_primeiro_uso(con)
    assert "criar_usuario" in cli.get("/usuarios/log").text


def test_todas_as_telas_abrem_para_admin(cli, admin, con):
    pid = _posto(con)
    banco.inserir_canal(con, "telegram", {"chat_ids": ["1"], "postos": [pid]})
    banco.inserir_canal(con, "email", {"servidor": "s", "porta": 587, "usuario": "a@b.c",
                                       "destinatarios": ["x@y.z"]})
    banco.inserir_canal(con, "toast", {}, ativo=False)
    banco.registrar_envio(con, 1, ontem(), "erro", "Não foi possível falar com o Telegram.")
    _logado(cli)
    for rota in ("/", "/postos", f"/postos/{pid}/fontes", "/canais", "/agenda", "/historico", "/coleta",
                 "/usuarios", "/usuarios/log", "/static/estilo.css"):
        r = cli.get(rota)
        assert r.status_code == 200, rota
        assert r.headers["x-frame-options"] == "DENY"
    assert "Cadastre cada posto em uma só máquina para evitar coleta duplicada." in cli.get("/postos").text
    assert "Não foi possível falar com o Telegram." in cli.get("/historico").text


def test_usuario_desativado_perde_sessao(fab, admin, con):
    acesso.criar_usuario(con, "Gê", "ge", SENHA, Perfil.GERENTE, [])
    c1 = TestClient(criar_app(fab), base_url="http://127.0.0.1", follow_redirects=False)
    _logado(c1, "ge")
    con.execute("UPDATE usuario SET ativo=0 WHERE login='ge'")
    con.commit()
    assert c1.get("/").headers["location"] == "/entrar"


def test_gerente_posto_alheio_403_fontes_e_reenviar(cli, admin, con):
    _posto(con, "Meu", 1)
    alheio = _posto(con, "Alheio", 2)
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [1])
    tok = _logado(cli, "gerente")
    for rota in (f"/postos/{alheio}/fontes/premmia", f"/postos/{alheio}/fontes/pagbank_edi/testar",
                 f"/reenviar?posto_id={alheio}"):
        r = cli.post(rota, data={"csrf": tok, "ativo": "1", "senha": SEGREDO})
        assert r.status_code == 403, rota
    assert not credenciais.tem(f"{alheio}/premmia/senha")


def test_usuario_com_posto_invalido(cli, admin, con):
    tok = _logado(cli)
    r = cli.post("/usuarios", data={"csrf": tok, "acao": "criar", "nome": "Gê", "login": "ge",
                                    "senha": SENHA, "perfil": "gerente", "postos": ["999"]})
    assert r.status_code == 400 and "Posto inválido." in r.text and "já existe" not in r.text
    assert con.execute("SELECT COUNT(*) FROM usuario WHERE login='ge'").fetchone()[0] == 0


def test_editar_canal_preserva_chaves_desconhecidas(cli, admin, con):
    cid = banco.inserir_canal(con, "telegram", {"chat_ids": ["1"], "extra": "manter"})
    tok = _logado(cli)
    cli.post("/canais", data={"csrf": tok, "acao": "editar", "id": cid, "chat_ids": "2", "ativo": "1"})
    cfg = json.loads(con.execute("SELECT config_json FROM canal WHERE id=?", (cid,)).fetchone()[0])
    assert cfg == {"chat_ids": ["2"], "extra": "manter", "postos": []}


def test_falha_em_segundo_plano_nao_loga_mensagem(caplog, fab, monkeypatch):
    import logging
    from apps.match import agendador
    from apps.match.canais import toast

    def explode(*a, **k):
        raise RuntimeError("SEGREDO-XYZ")

    async def explode_async(*a, **k):
        raise RuntimeError("SEGREDO-XYZ")

    monkeypatch.setattr(rotina, "executar", explode_async)
    monkeypatch.setattr(rotina, "limpar_antigos", explode)
    monkeypatch.setattr(toast, "avisar", explode)
    caplog.set_level(logging.DEBUG)
    rotas._em_segundo_plano("teste", explode).join(5)
    agendador._rodar(fab, so_falhas=False, limpar=True)
    assert caplog.records
    for rec in caplog.records:
        assert "SEGREDO-XYZ" not in rec.getMessage()
        assert "SEGREDO-XYZ" not in (rec.exc_text or "")
        assert rec.exc_info is None


def test_postos_recusa_codigo_repetido_ou_invalido(cli, admin, con):
    pid = _posto(con)
    tok = _logado(cli)
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Clone",
                                  "empresa_codigo_webposto": "74014"})
    assert r.status_code == 400 and "Já existe um posto com esse código do webPosto." in r.text
    for ruim in ("0", "-1", f"-{pid + 1}"):
        r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Neg",
                                      "empresa_codigo_webposto": ruim})
        assert r.status_code == 400 and "Código do webPosto inválido." in r.text, ruim
    outro = _posto(con, "Outro", 5)
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(outro), "nome": "Outro",
                                  "empresa_codigo_webposto": "74014"})
    assert r.status_code == 400 and "Já existe um posto" in r.text
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(pid), "nome": "Doze 2",
                                  "empresa_codigo_webposto": "74014"})
    assert r.status_code == 303  # o próprio posto mantém o código
    assert len(banco.listar_postos(con)) == 2


def test_importar_config_codigo_repetido_mostra_motivo(cli, admin, con):
    _posto(con)
    tok = _logado(cli)
    dados = json.dumps({"postos": [{"nome": "Outro", "empresa_codigo_webposto": 74014}]})
    r = cli.post("/config/importar", data={"csrf": tok},
                 files={"arquivo": ("c.json", dados, "application/json")})
    assert r.status_code == 400 and "Já existe um posto com esse código do webPosto." in r.text
    assert len(banco.listar_postos(con)) == 1


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_baixar_planilha(cli, admin, con, tmp_path, monkeypatch):
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a, b = _posto(con, "Alfa", 1), _posto(con, "Beta", 2)
    acesso.criar_usuario(con, "Cê", "conf", SENHA, Perfil.CONFERENTE, [a])
    _logado(cli, "conf")
    r = cli.get(f"/planilha?posto_id={a}&dia=17-10-2026")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.content[:2] == b"PK"
    assert "Logos_Alfa_17-10-2026.xlsx" in r.headers["content-disposition"]
    assert cli.get(f"/planilha?posto_id={a}&dia=2026-10-17").status_code == 200
    assert cli.get(f"/planilha?posto_id={b}&dia=17-10-2026").status_code == 403
    assert cli.get(f"/planilha?posto_id={a}&dia=31-02-2026").status_code == 400
    assert cli.get(f"/planilha?posto_id={a}&dia=abc").status_code == 400
    acoes = [x["acao"] for x in con.execute("SELECT acao FROM log_acesso")]
    assert acoes.count("baixar_planilha") == 2


def test_hoje_e_historico_tem_link_da_planilha(cli, admin, con):
    pid = _posto(con)
    fid = banco.salvar_fonte(con, pid, "pagbank_edi", True, {})
    banco.registrar_coleta(con, fid, ontem(), "ok", 1, "10", None)
    _logado(cli)
    link = f"/planilha?posto_id={pid}&amp;dia={ontem():%d-%m-%Y}"
    assert link in cli.get("/").text
    assert f"/planilha?posto_id={pid}&amp;dia={ontem().isoformat()}" in cli.get("/historico").text


def test_baixar_planilha_usa_mais_recente(cli, admin, con, tmp_path, monkeypatch):
    from apps.match import relatorio_xlsx
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a = _posto(con, "Alfa", 1)
    _logado(cli)
    dia = ontem()
    v1 = relatorio_xlsx.gerar(con, a, dia)
    banco.gravar_divergencias(con, a, dia, "premmia", [{"tipo": "x", "explicacao_pt": "novo"}])
    v2 = relatorio_xlsx.gerar(con, a, dia)
    assert v2 != v1
    r = cli.get(f"/planilha?posto_id={a}&dia={dia:%d-%m-%Y}")
    assert r.status_code == 200 and r.content == v2.read_bytes()
    assert v2.name in r.headers["content-disposition"]


def test_postos_recusa_apelido_equivalente(cli, admin, con):
    _posto(con, "Doze Filial", 74014)
    tok = _logado(cli)
    msg = "Já existe um posto com esse apelido (ou um apelido equivalente)."
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "N", "apelido": "Doze_Filial"})
    assert r.status_code == 400 and msg in r.text
    sem = banco.inserir_posto(con, "Sem", None, None, 5)  # sem apelido: P<id>
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(sem), "nome": "Sem",
                                  "apelido": "Doze Filial"})
    assert r.status_code == 400 and msg in r.text
    outro = banco.inserir_posto(con, "O", "O", None, 6)
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(outro), "nome": "O",
                                  "apelido": f"P{sem}"})
    assert r.status_code == 400 and msg in r.text
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(sem), "nome": "Sem",
                                  "apelido": "Doze Filial 2"})
    assert r.status_code == 303


def test_planilha_com_conflito_existente_da_409(cli, admin, con, tmp_path, monkeypatch):
    from apps.match import relatorio_xlsx
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a = banco.inserir_posto(con, "A", "Doze Filial", None, 1)
    b = banco.inserir_posto(con, "B", "Doze_Filial", None, 2)
    outro = tmp_path / "dados" / "planilhas"
    outro.mkdir(parents=True, exist_ok=True)
    dia = ontem()
    (outro / f"Logos_Doze_Filial_{dia:%d-%m-%Y}.xlsx").write_bytes(b"PKalheio")
    _logado(cli)
    for pid in (a, b):
        r = cli.get(f"/planilha?posto_id={pid}&dia={dia:%d-%m-%Y}")
        assert r.status_code == 409 and b"alheio" not in r.content
        assert "Apelido em conflito com outro posto; ajuste o cadastro." in r.text


# ---------------------------------------------------------------- coleta ao vivo e histórico

REFRESH = '<meta http-equiv="refresh" content="3;url=/coleta">'


def _execucao(con, itens, situacao="em_andamento", origem="agendada", usuario_id=None, dia=None):
    eid = banco.abrir_execucao(con, origem, usuario_id, dia or ontem())
    for pid, fonte, sit, msg in itens:
        banco.item_execucao(con, eid, pid, fonte, sit, msg)
    if situacao != "em_andamento":
        banco.fechar_execucao(con, eid, situacao)
    return eid


def test_andamento_tem_refresh_quando_em_andamento(cli, admin, con, monkeypatch):
    monkeypatch.setattr(rotina, "em_andamento", lambda: True)
    pid = _posto(con, "Alfa", 1)
    _execucao(con, [(pid, "premmia", "coletando", None), (pid, "pagbank_edi", "aguardando", None)])
    _logado(cli)
    html = cli.get("/coleta").text
    assert REFRESH in html
    assert "Coletando…" in html and "Aguardando" in html and "Agendada" in html
    assert "Premmia" in html and "PagBank" in html


def test_andamento_sem_refresh_quando_concluida(cli, admin, con):
    pid = _posto(con, "Alfa", 1)
    _execucao(con, [(pid, "premmia", "ok", "3 vendas."), (pid, "pagbank_edi", "coletando", None)],
              situacao="concluida")
    _logado(cli)
    html = cli.get("/coleta").text
    assert 'http-equiv="refresh"' not in html
    assert re.search(r"Concluída às \d{2}/\d{2}/\d{4} \d{2}:\d{2}", html)
    assert "OK" in html and "3 vendas." in html
    assert "Não concluído" in html and "Coletando…" not in html


def test_andamento_flash_some_depois_de_concluida(cli, admin, con, monkeypatch):
    pid = _posto(con, "Alfa", 1)
    eid = _execucao(con, [(pid, "premmia", "coletando", None)])
    _logado(cli)
    monkeypatch.setattr(rotina, "em_andamento", lambda: True)
    html = cli.get("/coleta?ja=1").text
    assert "Já existe uma coleta em andamento." in html and REFRESH in html
    assert "Coleta iniciada." in cli.get("/coleta?iniciada=1").text
    banco.fechar_execucao(con, eid, "concluida")
    monkeypatch.setattr(rotina, "em_andamento", lambda: False)
    for url in ("/coleta?ja=1", "/coleta?iniciada=1"):
        html = cli.get(url).text
        assert "Concluída às" in html
        assert "Já existe" not in html and "Coleta iniciada." not in html


def test_andamento_nao_concluido_sem_tempo(cli, admin, con):
    pid = _posto(con, "Alfa", 1)
    _execucao(con, [(pid, "premmia", "coletando", None)], situacao="concluida")
    _logado(cli)
    html = cli.get("/coleta").text
    assert re.search(r"Não concluído</td><td></td><td>—</td>", html)


def test_andamento_gerente_ultima_execucao_dos_seus_postos(cli, admin, con):
    a, b = _posto(con, "Posto Alfa", 1), _posto(con, "Posto Beta", 2)
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [a])
    _execucao(con, [(b, "premmia", "ok", "Só Beta.")], situacao="concluida", origem="bandeja")
    _logado(cli, "gerente")
    html = cli.get("/coleta").text
    assert "Nenhuma coleta registrada ainda." in html
    assert "Bandeja" not in html and "Só Beta." not in html and "Concluída às" not in html
    _execucao(con, [(a, "premmia", "ok", "Alfa ok.")], situacao="concluida", origem="repescagem")
    _execucao(con, [(b, "premmia", "ok", "Beta de novo.")], situacao="concluida", origem="bandeja")
    html = cli.get("/coleta").text
    assert "Alfa ok." in html and "Repescagem" in html
    assert "Bandeja" not in html and "Beta de novo." not in html


def test_andamento_sem_execucao(cli, admin, con):
    pid = _posto(con, "Alfa", 1)
    _logado(cli)
    html = cli.get("/coleta").text
    assert "Nenhuma coleta registrada ainda." in html and REFRESH not in html
    assert f"/coletar?posto_id={pid}" in html


def test_andamento_so_postos_visiveis(cli, admin, con):
    a, b = _posto(con, "Posto Alfa", 1), _posto(con, "Posto Beta", 2)
    _execucao(con, [(a, "premmia", "erro", "Falha Alfa."), (b, "premmia", "erro", "Falha Beta.")])
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [a])
    _logado(cli, "gerente")
    html = cli.get("/coleta").text
    assert "Posto Alfa" in html and "Falha Alfa." in html
    assert "Posto Beta" not in html and "Falha Beta." not in html


def test_coletar_redireciona_para_andamento(cli, admin, con, executar_falso):
    pid = _posto(con)
    tok = _logado(cli)
    r = cli.post(f"/coletar?posto_id={pid}&visivel=0", data={"csrf": tok})
    assert r.status_code == 303 and r.headers["location"].startswith("/coleta")
    assert executar_falso[1].wait(5)
    assert cli.get(r.headers["location"]).status_code == 200


def test_coletar_duplo_mostra_em_andamento(cli, admin, con, executar_falso):
    pid = _posto(con, "Alfa", 1)
    _execucao(con, [(pid, "premmia", "coletando", None)], origem="painel", usuario_id=admin)
    tok = _logado(cli)
    assert rotina._TRAVA.acquire(blocking=False)
    try:
        r = cli.post(f"/coletar?posto_id={pid}&visivel=0", data={"csrf": tok})
        assert r.status_code == 303 and r.headers["location"].startswith("/coleta")
        html = cli.get(r.headers["location"]).text
    finally:
        rotina._TRAVA.release()
    assert executar_falso[0] == []
    assert "Já existe uma coleta em andamento." in html
    assert "Coletando…" in html and REFRESH in html and "por Admin" in html
    assert "Nenhuma coleta registrada ainda." not in html


def test_historico_mostra_execucao_duracao_origem(cli, admin, con):
    pid = _posto(con, "Alfa", 1)
    eid = _execucao(con, [(pid, "premmia", "ok", None), (pid, "pagbank_edi", "erro", "PagBank fora.")],
                    situacao="concluida", origem="repescagem")
    con.execute("UPDATE execucao SET iniciada_em='2026-10-18T09:00:00+00:00',"
                " terminada_em='2026-10-18T09:02:05+00:00' WHERE id=?", (eid,))
    _execucao(con, [], situacao="falhou", origem="painel", usuario_id=admin)
    canal = banco.inserir_canal(con, "telegram", {"chat_ids": ["1"]})
    banco.registrar_envio(con, canal, ontem(), "erro", "Telegram recusou.", eid)
    con.commit()
    _logado(cli)
    html = cli.get("/historico").text
    assert "18/10/2026 06:00" in html and "18/10/2026 06:02" in html and "2m 5s" in html
    assert "Repescagem" in html and "Concluída" in html and "Falhou" in html
    assert "PagBank fora." in html and "Telegram recusou." in html
    assert "por Admin" in html and "por Repescagem" in html


def test_historico_versoes_da_planilha(cli, admin, con, tmp_path, monkeypatch):
    from apps.match import relatorio_xlsx
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    pid = _posto(con, "Alfa", 1)
    dia = ontem()
    relatorio_xlsx.gerar(con, pid, dia)
    banco.gravar_divergencias(con, pid, dia, "premmia", [{"tipo": "x", "explicacao_pt": "novo"}])
    relatorio_xlsx.gerar(con, pid, dia)
    _execucao(con, [(pid, "premmia", "ok", None)], situacao="concluida", dia=dia)
    _logado(cli)
    html = cli.get("/historico").text
    for n in (1, 2):
        assert f"/planilha?posto_id={pid}&amp;dia={dia:%d-%m-%Y}&amp;versao={n}" in html
    assert "versao=3" not in html


def test_planilha_versao_especifica_e_404(cli, admin, con, tmp_path, monkeypatch):
    from apps.match import relatorio_xlsx
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a, b = _posto(con, "Alfa", 1), _posto(con, "Beta", 2)
    dia = ontem()
    v1 = relatorio_xlsx.gerar(con, a, dia)
    banco.gravar_divergencias(con, a, dia, "premmia", [{"tipo": "x", "explicacao_pt": "novo"}])
    v2 = relatorio_xlsx.gerar(con, a, dia)
    acesso.criar_usuario(con, "Cê", "conf", SENHA, Perfil.CONFERENTE, [a])
    _logado(cli, "conf")
    r = cli.get(f"/planilha?posto_id={a}&dia={dia:%d-%m-%Y}&versao=1")
    assert r.status_code == 200 and r.content == v1.read_bytes() and v1.name in r.headers["content-disposition"]
    r = cli.get(f"/planilha?posto_id={a}&dia={dia:%d-%m-%Y}&versao=2")
    assert r.status_code == 200 and r.content == v2.read_bytes()
    r = cli.get(f"/planilha?posto_id={a}&dia={dia:%d-%m-%Y}&versao=3")
    assert r.status_code == 404 and "Versão não encontrada." in r.text
    assert cli.get(f"/planilha?posto_id={b}&dia={dia:%d-%m-%Y}&versao=1").status_code == 403


def test_historico_gerente_so_seus_postos(cli, admin, con, tmp_path, monkeypatch):
    from apps.match import relatorio_xlsx
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a, b = _posto(con, "Posto Alfa", 1), _posto(con, "Posto Beta", 2)
    relatorio_xlsx.gerar(con, a, ontem())
    relatorio_xlsx.gerar(con, b, ontem())
    _execucao(con, [(a, "premmia", "ok", "Alfa ok."), (b, "premmia", "erro", "Beta caiu.")],
              situacao="concluida")
    _execucao(con, [(b, "pagbank_edi", "erro", "Só Beta.")], situacao="concluida", origem="bandeja")
    canal = banco.inserir_canal(con, "telegram", {"chat_ids": ["1"]})
    banco.registrar_envio(con, canal, ontem(), "erro", "Envio secreto.", 1)
    acesso.criar_usuario(con, "Gê", "gerente", SENHA, Perfil.GERENTE, [a])
    _logado(cli, "gerente")
    html = cli.get("/historico").text
    assert "Alfa ok." in html
    assert "Beta caiu." not in html and "Só Beta." not in html and "Bandeja" not in html
    assert "Posto Beta" not in html and "Envio secreto." not in html
    assert f"posto_id={a}&amp;" in html and f"posto_id={b}" not in html


# ---------------------------------------------------------------- pasta da rede (v0.2)

import html as _html  # noqa: E402
import socket  # noqa: E402

from apps.match import registro_rede as rr  # noqa: E402

MSG_CHAVE = ("Informe o código do webPosto ou o CNPJ para proteger este posto contra cadastro"
             " em duas máquinas.")


def _rede(con, tmp_path):
    p = tmp_path / "rede"
    p.mkdir(exist_ok=True)
    banco.salvar_pasta_rede(con, str(p))
    return p


def _texto(r) -> str:
    return _html.unescape(r.text)


def test_pasta_rede_salvar_registra_ativos(cli, admin, con, tmp_path):
    a = _posto(con, "Doze", 74014)
    b = _posto(con, "Vip", 5)
    banco.inserir_posto(con, "Sem", "Sem", None, None)
    rede = tmp_path / "rede"
    rede.mkdir()
    rr.registrar(rede, "5", "OUTRA", "Vip")
    tok = _logado(cli)
    r = cli.post("/pasta-rede", data={"csrf": tok, "acao": "salvar", "caminho": str(rede)})
    assert r.status_code == 200, r.text
    assert banco.pasta_rede(con) == rede
    assert rr.pasta_postos(rede).is_dir()
    assert rr.dono(rede, "74014").id_instalacao == banco.id_instalacao(con)
    assert rr.dono(rede, "5").id_instalacao == "OUTRA"
    assert banco.obter_posto(con, a)["ativo"] and not banco.obter_posto(con, b)["ativo"]
    assert rr.mensagem_bloqueio(rr.dono(rede, "5")) in _texto(r)
    log = con.execute("SELECT acao FROM log_acesso WHERE acao='pasta_rede'").fetchall()
    assert len(log) == 1
    r = cli.post("/pasta-rede", data={"csrf": tok, "acao": "salvar", "caminho": str(tmp_path / "nao")})
    assert r.status_code == 400 and banco.pasta_rede(con) == rede
    r = cli.post("/pasta-rede", data={"csrf": tok, "acao": "desligar"})
    assert r.status_code == 303 and banco.pasta_rede(con) is None
    assert cli.get("/agenda").status_code == 200


def test_pasta_rede_so_admin(cli, admin, con, tmp_path):
    pid = _posto(con)
    acesso.criar_usuario(con, "G", "ger", SENHA, Perfil.GERENTE, [pid])
    tok = _logado(cli, "ger")
    r = cli.post("/pasta-rede", data={"csrf": tok, "acao": "salvar", "caminho": str(tmp_path)})
    assert r.status_code == 403 and banco.pasta_rede(con) is None


def test_cadastro_bloqueado_outra_maquina_409(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    rr.registrar(rede, "74014", "OUTRA", "Doze")
    rr.registrar(rede, "12345678000190", "OUTRA", "Cnpj")
    bloqueio = rr.mensagem_bloqueio(rr.dono(rede, "74014"))
    tok = _logado(cli)
    pid = _posto(con)
    banco.definir_posto_ativo(con, pid, False)
    r = cli.post("/postos", data={"csrf": tok, "acao": "ativar", "id": str(pid)})
    assert r.status_code == 409 and bloqueio in _texto(r)
    assert not banco.obter_posto(con, pid)["ativo"]
    outro = _posto(con, "Outro", 7)
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(outro), "nome": "Outro",
                                  "cnpj": "12.345.678/0001-90", "empresa_codigo_webposto": ""})
    assert r.status_code == 409
    assert banco.obter_posto(con, outro)["empresa_codigo_webposto"] == 7


def test_cadastro_pasta_indisponivel_503(cli, admin, con, tmp_path):
    banco.salvar_pasta_rede(con, str(tmp_path / "pasta-offline"))
    pid = _posto(con, "Doze", 74014)
    banco.definir_posto_ativo(con, pid, False)
    ativo = _posto(con, "Vip", 5)
    tok = _logado(cli)
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Novo",
                                  "empresa_codigo_webposto": "9"})
    assert r.status_code == 503 and rr.MSG_INDISPONIVEL in _texto(r)
    r = cli.post("/postos", data={"csrf": tok, "acao": "ativar", "id": str(pid)})
    assert r.status_code == 503 and not banco.obter_posto(con, pid)["ativo"]
    r = cli.post("/postos", data={"csrf": tok, "acao": "editar", "id": str(ativo), "nome": "Vip 2",
                                  "empresa_codigo_webposto": "6"})
    assert r.status_code == 503 and rr.MSG_INDISPONIVEL in _texto(r)
    assert banco.obter_posto(con, ativo)["nome"] == "Vip"
    dados = json.dumps({"postos": [{"nome": "Imp", "empresa_codigo_webposto": 11}]})
    r = cli.post("/config/importar", data={"csrf": tok},
                 files={"arquivo": ("c.json", dados, "application/json")})
    assert r.status_code == 503 and rr.MSG_INDISPONIVEL in _texto(r)
    assert len(banco.listar_postos(con, so_ativos=False)) == 2
    assert not (tmp_path / "pasta-offline").exists()


def test_cadastro_exige_codigo_ou_cnpj_com_pasta(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    tok = _logado(cli)
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Sem chave"})
    assert r.status_code == 400 and MSG_CHAVE in _texto(r)
    assert banco.listar_postos(con, so_ativos=False) == []
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Com CNPJ",
                                  "cnpj": "12.345.678/0001-90"})
    assert r.status_code == 303
    assert rr.dono(rede, "12345678000190").posto_nome == "Com CNPJ"


def test_assumir_posto_admin_e_log(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    pid = _posto(con)
    rr.registrar(rede, "74014", "OUTRA", "Doze")
    acesso.criar_usuario(con, "G", "ger", SENHA, Perfil.GERENTE, [pid])
    tok = _logado(cli, "ger")
    r = cli.post(f"/postos/{pid}/assumir", data={"csrf": tok, "confirmar": "1"})
    assert r.status_code == 403 and rr.dono(rede, "74014").id_instalacao == "OUTRA"
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/assumir", data={"csrf": tok})
    assert r.status_code == 400 and rr.dono(rede, "74014").id_instalacao == "OUTRA"
    r = cli.post(f"/postos/{pid}/assumir", data={"csrf": tok, "confirmar": "1"})
    assert r.status_code == 303
    assert rr.dono(rede, "74014").id_instalacao == banco.id_instalacao(con)
    log = con.execute("SELECT usuario_id, detalhe FROM log_acesso WHERE acao='assumir_posto'").fetchall()
    assert [(x["usuario_id"], x["detalhe"]) for x in log] == [(admin, f"posto={pid};antes={socket.gethostname()}")]


def test_reinstalado_mostra_mesma_maquina_e_assume(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    pid = _posto(con)
    rr.registrar(rede, "74014", "ID-DA-INSTALACAO-ANTIGA", "Doze")  # mesmo PC, antes de reinstalar
    tok = _logado(cli)
    r = cli.get("/postos")
    assert f"Pertence à máquina {socket.gethostname()}" in _texto(r)
    assert f'action="/postos/{pid}/assumir"' in r.text
    assert "Confirmo que a outra máquina não coleta mais este posto" in _texto(r)
    r = cli.post(f"/postos/{pid}/assumir", data={"csrf": tok, "confirmar": "1"})
    assert r.status_code == 303
    r = cli.get("/postos")
    assert "Protegido (esta máquina)" in _texto(r) and "Pertence à máquina" not in _texto(r)


def test_desativar_libera(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    tok = _logado(cli)
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Doze",
                                  "empresa_codigo_webposto": "74014"})
    assert r.status_code == 303
    pid = banco.listar_postos(con)[0]["id"]
    assert rr.dono(rede, "74014").id_instalacao == banco.id_instalacao(con)
    r = cli.post("/postos", data={"csrf": tok, "acao": "desativar", "id": str(pid)})
    assert r.status_code == 303 and rr.dono(rede, "74014") is None
    assert not banco.obter_posto(con, pid)["ativo"]
    outro = _posto(con, "Vip", 5)
    banco.salvar_pasta_rede(con, str(tmp_path / "pasta-offline"))
    r = cli.post("/postos", data={"csrf": tok, "acao": "desativar", "id": str(outro)})
    assert r.status_code == 303 and not banco.obter_posto(con, outro)["ativo"]


def test_assumir_mostra_aviso_e_loga_maquina_anterior(cli, admin, con, tmp_path):
    rede = _rede(con, tmp_path)
    pid = _posto(con)
    banco.definir_posto_ativo(con, pid, False)
    rr.registrar(rede, "74014", "OUTRA", "Doze")
    tok = _logado(cli)
    r = cli.post(f"/postos/{pid}/assumir", data={"csrf": tok, "confirmar": "1"})
    assert r.status_code == 303
    r = cli.get(r.headers["location"])
    assert "Posto assumido. Clique em Ativar para voltar a coletar." in _texto(r)
    log = con.execute("SELECT detalhe FROM log_acesso WHERE acao='assumir_posto'").fetchone()
    assert log["detalhe"] == f"posto={pid};antes={socket.gethostname()}"


# ---------------------------------------------------------------- revisão final v0.2

def _desativado_aqui(dono) -> str:
    return (f"Este posto já está cadastrado na máquina {dono.maquina} desde "
            f"{rr.formatar_data_hora(rr.ler_data_hora(dono.registrado_em))}. Ele foi salvo desativado"
            " aqui; use 'Assumir posto' se a outra máquina não coleta mais.")


def test_baixar_planilha_regera_apos_nova_coleta(cli, admin, con, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    monkeypatch.setenv("LOGOS_MATCH_DIR", str(tmp_path / "dados"))
    a = _posto(con, "Alfa", 1)
    fid = banco.salvar_fonte(con, a, "premmia", True, {})
    dia = ontem()
    banco.registrar_coleta(con, fid, dia, "erro", 0, "0", "Premmia fora.")
    _logado(cli)
    url = f"/planilha?posto_id={a}&dia={dia:%d-%m-%Y}"
    r1 = cli.get(url)
    assert r1.status_code == 200 and f"Logos_Alfa_{dia:%d-%m-%Y}.xlsx" in r1.headers["content-disposition"]
    banco.registrar_coleta(con, fid, dia, "ok", 3, "30", None)  # nova coleta, sem canal ativo
    r2 = cli.get(url)
    assert r2.status_code == 200 and f"Logos_Alfa_{dia:%d-%m-%Y}_v2.xlsx" in r2.headers["content-disposition"]
    arq = tmp_path / "dados" / "planilhas" / f"Logos_Alfa_{dia:%d-%m-%Y}_v2.xlsx"
    assert r2.content == arq.read_bytes()
    wb = load_workbook(arq, read_only=True)
    try:
        situacoes = [linha[1] for linha in wb["Resumo"].iter_rows(min_row=2, values_only=True)]
    finally:
        wb.close()
    assert situacoes == ["OK"]
    r3 = cli.get(url)  # sem mudança: mesmo arquivo, nenhuma versão nova
    assert r3.content == r2.content and "_v2.xlsx" in r3.headers["content-disposition"]
    assert not (arq.parent / f"Logos_Alfa_{dia:%d-%m-%Y}_v3.xlsx").exists()
    r = cli.get(url + "&versao=1")
    assert r.status_code == 200 and r.content == r1.content


def _assumir_ativar_coletar(cli, tok, con, rede, pid, chave, executar_falso):
    r = cli.get("/postos")
    assert f'action="/postos/{pid}/assumir"' in r.text and "Desativado" in r.text
    assert cli.post(f"/postos/{pid}/assumir", data={"csrf": tok, "confirmar": "1"}).status_code == 303
    assert rr.dono(rede, chave).id_instalacao == banco.id_instalacao(con)
    r = cli.post("/postos", data={"csrf": tok, "acao": "ativar", "id": str(pid)})
    assert r.status_code == 303 and banco.obter_posto(con, pid)["ativo"]
    r = cli.post(f"/coletar?posto_id={pid}&visivel=0", data={"csrf": tok})
    assert r.status_code == 303 and executar_falso[1].wait(5)
    assert executar_falso[0][0]["posto_id"] == pid


def test_maquina_substituta_cadastra_desativado_e_assume(cli, admin, con, tmp_path, executar_falso):
    rede = _rede(con, tmp_path)
    rr.registrar(rede, "74014", "OUTRA", "Doze")
    antes = rr.dono(rede, "74014")
    tok = _logado(cli)
    r = cli.post("/postos", data={"csrf": tok, "acao": "novo", "nome": "Doze",
                                  "empresa_codigo_webposto": "74014"})
    assert r.status_code == 200 and _desativado_aqui(antes) in _texto(r)
    [posto] = banco.listar_postos(con, so_ativos=False)
    assert not posto["ativo"] and posto["empresa_codigo_webposto"] == 74014
    assert rr.dono(rede, "74014") == antes  # pasta da rede intocada
    _assumir_ativar_coletar(cli, tok, con, rede, posto["id"], "74014", executar_falso)


def test_maquina_substituta_importa_desativado_e_assume(cli, admin, con, tmp_path, executar_falso):
    rede = _rede(con, tmp_path)
    rr.registrar(rede, "74014", "OUTRA", "Doze")
    antes = rr.dono(rede, "74014")
    tok = _logado(cli)
    dados = json.dumps({"postos": [{"nome": "Doze", "empresa_codigo_webposto": 74014},
                                   {"nome": "Vip", "empresa_codigo_webposto": 5}]})
    r = cli.post("/config/importar", data={"csrf": tok},
                 files={"arquivo": ("c.json", dados, "application/json")})
    assert r.status_code == 200
    texto = _texto(r)
    assert "Salvos desativados" in texto and f"Doze: {_desativado_aqui(antes)}" in texto
    assert "1 posto(s) salvo(s) desativado(s)" in texto
    postos = {p["nome"]: p for p in banco.listar_postos(con, so_ativos=False)}
    assert not postos["Doze"]["ativo"] and postos["Vip"]["ativo"]
    assert rr.dono(rede, "74014") == antes
    assert rr.dono(rede, "5").id_instalacao == banco.id_instalacao(con)
    _assumir_ativar_coletar(cli, tok, con, rede, postos["Doze"]["id"], "74014", executar_falso)


def test_coleta_orfa_sem_refresh(cli, admin, con, monkeypatch):
    pid = _posto(con, "Alfa", 1)
    _execucao(con, [(pid, "premmia", "ok", "3 vendas."), (pid, "pagbank_edi", "coletando", None)])
    _logado(cli)
    monkeypatch.setattr(rotina, "em_andamento", lambda: False)
    html = cli.get("/coleta").text
    assert 'http-equiv="refresh"' not in html
    assert "Não concluído" in html and "Coletando…" not in html
    assert "Esta página se atualiza sozinha" not in html and "Concluída às" not in html
    assert "A coleta foi interrompida" in html
    monkeypatch.setattr(rotina, "em_andamento", lambda: True)
    html = cli.get("/coleta").text
    assert REFRESH in html and "Coletando…" in html


def test_pasta_rede_indisponivel_no_meio_nao_salva(cli, admin, con, tmp_path, monkeypatch):
    _posto(con, "Doze", 74014)
    b = _posto(con, "Vip", 5)
    _posto(con, "Real", 6)
    rede = tmp_path / "rede"
    rede.mkdir()
    rr.registrar(rede, "5", "OUTRA", "Vip")
    real = rr.registrar
    chamadas = []

    def falha_no_terceiro(pasta, chave, *a):
        chamadas.append(chave)
        if len(chamadas) == 3:
            raise rr.PastaIndisponivel(rr.MSG_INDISPONIVEL)
        return real(pasta, chave, *a)

    monkeypatch.setattr(rr, "registrar", falha_no_terceiro)
    tok = _logado(cli)
    r = cli.post("/pasta-rede", data={"csrf": tok, "acao": "salvar", "caminho": str(rede)})
    assert r.status_code == 503 and rr.MSG_INDISPONIVEL in _texto(r)
    assert len(chamadas) == 3
    assert banco.pasta_rede(con) is None
    assert banco.obter_posto(con, b)["ativo"]  # nada desativado
