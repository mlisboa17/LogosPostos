import asyncio
import ast
from types import SimpleNamespace
from unittest.mock import MagicMock
from datetime import date, datetime, timedelta
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import ValidationError

from src.jobs import noturno as composicao_noturna
from src.modules.cash_reconciliation.adapters import persistencia
from src.modules.cash_reconciliation.adapters import saude_robo
from src.modules.cash_reconciliation.domain.execucao import ResultadoDiario
from src.modules.cash_reconciliation.domain.integridade_execucao import ExecucaoNoturna, FonteExecucao
from src.modules.cash_reconciliation.domain.tempo import agora
from src.modules.cash_reconciliation.domain.models import (
    AdquirenteConfigurada,
    Proveniencia,
    Unidade,
)
from src.modules.cash_reconciliation.domain.recebimentos import (
    RecebimentoAdquirente,
    ResultadoRecebimentos,
    ResumoCasados,
)
from src.modules.cash_reconciliation.interfaces import http
from src.modules.cash_reconciliation.jobs import noturno
from src.infrastructure.security.jwt_utils import create_access_token
from src.shared.notificacao.telegram import (
    ResultadoDestinatario,
    enviar_mensagem as telegram_enviar_mensagem,
)

from .test_http import resultado_sintetico
from src.modules.cash_reconciliation.adapters import premmia_portal
from src.modules.cash_reconciliation.application.resumo_diario import gerar_mensagem_curta
from src.modules.cash_reconciliation.domain.tempo import FUSO

DIA = date(2026, 10, 1)
EMPRESA = 321


@pytest.fixture(autouse=True)
def sem_postos_comerciais_reais(monkeypatch, tmp_path):
    monkeypatch.setattr(saude_robo, "DIRETORIO_RESUMOS", tmp_path / "resumos")
    monkeypatch.setattr(saude_robo, "CAMINHO_ESTADO", tmp_path / "resumos" / "estado_robo.json")
    monkeypatch.setenv("LOGOS_CARTAO_AUTOMATICO", "0")
    from src.modules.cash_reconciliation.application import fontes_falta

    async def sem_fontes_locais(registro, unidade):
        return registro

    monkeypatch.setattr(fontes_falta, "anexar_fontes_falta", sem_fontes_locais)


def unidade(empresa=EMPRESA):
    return Unidade(
        empresa_codigo=empresa,
        nome="Sintetica",
        chave_env="CHAVE_SINTETICA",
        destinos=(),
        adquirentes=(
            AdquirenteConfigurada(nome="PAGBANK", situacao="ativo", modalidades=("pix",)),
            AdquirenteConfigurada(
                nome="MAIS_PAGAMENTOS", situacao="pendente", modalidades=("pix",)
            ),
        ),
    )


def diario(empresa=EMPRESA, dia=DIA, erro=None):
    fechamento = resultado_sintetico().model_copy(
        update={
            "empresa_codigo": empresa,
            "inicio": dia,
            "fim": dia,
            "caixas": resultado_sintetico().caixas if empresa == EMPRESA and dia == DIA else (),
        }
    )
    return ResultadoDiario(
        empresa_codigo=empresa,
        dia=dia,
        fechamento=fechamento if erro is None else None,
        erro_fechamento=erro,
        recebimentos=ResultadoRecebimentos(
            empresa_codigo=empresa,
            dia=dia,
            adquirentes=(
                RecebimentoAdquirente(
                    adquirente="PAGBANK",
                    situacao="ok",
                    casados=ResumoCasados(quantidade=1, total="10.15"),
                    proveniencia=Proveniencia(
                        execucao_id="recebimento-sintetico",
                        executado_em=datetime(2026, 10, 2, 3),
                        versao_regra="CARTOES_V1",
                        fonte_sangrias="",
                        extratos=(),
                        outras_fontes=("sintetica",),
                    ),
                ),
            ),
        ),
        proveniencia=Proveniencia(
            execucao_id="execucao-sintetica",
            executado_em=datetime(2026, 10, 2, 3),
            versao_regra=noturno.VERSAO_JOB,
            fonte_sangrias="sintetica",
            extratos=(),
            outras_fontes=("sintetica",),
        ),
        versoes_regras=("FECHAMENTO_V3", "CARTOES_V1"),
    )


def coleta_corrente(registro, run_id):
    instante = agora()
    fontes = [FonteExecucao(
        nome="fechamento", unidade=registro.empresa_codigo, dia=registro.dia, run_id=run_id,
        coletado_em=instante, cobertura_inicio=registro.dia, cobertura_fim=registro.dia,
        estado="FAILED" if registro.erro_fechamento else "SUCCESS",
    )]
    fontes.extend(FonteExecucao(
        nome=a.adquirente, unidade=registro.empresa_codigo, dia=registro.dia, run_id=run_id,
        coletado_em=instante, cobertura_inicio=registro.dia, cobertura_fim=registro.dia,
        estado="SUCCESS" if a.situacao == "ok" else "BLOCKED",
    ) for a in registro.recebimentos.adquirentes)
    return registro.model_copy(update={
        "proveniencia": registro.proveniencia.model_copy(update={
            "execucao_id": run_id, "executado_em": instante,
        }), "fontes_execucao": tuple(fontes),
    })


@pytest.fixture
def contexto_envio():
    execucao = ExecucaoNoturna(dia=DIA)
    registro = coleta_corrente(diario(), execucao.run_id)
    execucao.fontes = list(registro.fontes_execucao)
    execucao.finalizar()
    return {"execucao": execucao, "resultados": [registro]}


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(persistencia, "DIRETORIO", tmp_path)
    monkeypatch.setattr(http, "carregar_unidades", lambda: {EMPRESA: unidade()})
    app = FastAPI()
    app.include_router(http.router)
    with TestClient(app) as client:
        client.cookies.set(
            "access_token",
            create_access_token(
                "director@example.invalid", extra={"role": "diretor", "token_type": "access"}
            ),
        )
        yield client


def test_persistencia_idempotente_e_sem_temporarios(tmp_path):
    primeiro = diario()
    caminho = persistencia.salvar_dia(primeiro, diretorio=tmp_path)
    assert caminho == tmp_path / DIA.isoformat() / f"{EMPRESA}.json"
    assert persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path) == primeiro
    segundo = diario(erro="falha ao consultar fechamento")
    persistencia.salvar_dia(segundo, diretorio=tmp_path)
    assert persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path) == segundo
    assert list(caminho.parent.iterdir()) == [caminho]
    assert persistencia.carregar_dia(999, DIA, diretorio=tmp_path) is None


def test_falha_atomica_preserva_resultado_anterior(monkeypatch, tmp_path):
    anterior = diario()
    persistencia.salvar_dia(anterior, diretorio=tmp_path)

    def falhar(*args):
        raise OSError("segredo")

    monkeypatch.setattr(persistencia.os, "replace", falhar)
    with pytest.raises(persistencia.PersistenciaErro, match="persistir"):
        persistencia.salvar_dia(diario(erro="falha"), diretorio=tmp_path)
    assert persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path) == anterior
    assert len(list((tmp_path / str(DIA)).iterdir())) == 1


@pytest.mark.parametrize("conteudo", ["corrompido segredo", '{"empresa_codigo":999}'])
def test_corrompido_nao_expoe_conteudo(tmp_path, caplog, conteudo):
    caminho = persistencia.caminho_dia(EMPRESA, DIA, tmp_path)
    caminho.parent.mkdir()
    caminho.write_text(conteudo, encoding="utf-8")
    with pytest.raises(persistencia.PersistenciaErro) as exc:
        persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path)
    assert "segredo" not in str(exc.value) + caplog.text


def test_arquivo_de_outro_tenant_rejeitado(tmp_path):
    caminho = persistencia.caminho_dia(EMPRESA, DIA, tmp_path)
    caminho.parent.mkdir()
    caminho.write_text(diario(999).model_dump_json(), encoding="utf-8")
    with pytest.raises(persistencia.PersistenciaErro):
        persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path)


def test_modelo_rejeita_escopo_interno_e_resultado_com_erro():
    dados = diario().model_dump()
    dados["recebimentos"]["empresa_codigo"] = 999
    with pytest.raises(ValidationError):
        ResultadoDiario.model_validate(dados)
    dados = diario().model_dump()
    dados["erro_fechamento"] = "erro"
    with pytest.raises(ValidationError):
        ResultadoDiario.model_validate(dados)


def test_rotas_usam_persistido_sem_rede(client, monkeypatch, tmp_path):
    persistencia.salvar_dia(diario(), diretorio=tmp_path)

    async def proibido(*args):
        pytest.fail("Registro persistido nao pode consultar rede.")

    monkeypatch.setattr(http, "auditar_unidade", proibido)
    monkeypatch.setattr(http, "conciliar_recebimentos", proibido)
    receb = client.get("/api/v1/cash-audit/recebimentos", params={"unidade": EMPRESA, "dia": DIA})
    assert receb.status_code == 200
    assert receb.json()["adquirentes"][0]["casados"]["quantidade"] == 1
    fechamento = client.get(
        "/api/v1/cash-audit/fechamento", params={"unidade": EMPRESA, "inicio": DIA, "fim": DIA}
    )
    assert fechamento.status_code == 200
    assert fechamento.json()["quebra_total"] == "2.50"
    assert fechamento.json()["caixas"][0]["severidade"] == "vermelho"
    assert fechamento.json()["execucoes"][0]["execucao_id"] == "execucao-sintetica"


def test_periodo_misto_consulta_so_dia_ausente(client, monkeypatch, tmp_path):
    persistencia.salvar_dia(diario(), diretorio=tmp_path)
    chamadas = []
    depois = DIA + timedelta(days=1)

    async def buscar(empresa, inicio, fim):
        chamadas.append((empresa, inicio, fim))
        return diario(dia=inicio).fechamento

    monkeypatch.setattr(http, "auditar_unidade", buscar)
    response = client.get(
        "/api/v1/cash-audit/fechamento",
        params={"unidade": EMPRESA, "inicio": DIA, "fim": depois},
    )
    assert response.status_code == 200
    assert chamadas == [(EMPRESA, depois, depois)]
    assert len(response.json()["execucoes"]) == 2


def test_erro_persistido_nao_consulta_ao_vivo(client, tmp_path):
    persistencia.salvar_dia(diario(erro="falha ao consultar fechamento"), diretorio=tmp_path)
    response = client.get(
        "/api/v1/cash-audit/fechamento", params={"unidade": EMPRESA, "inicio": DIA, "fim": DIA}
    )
    assert response.status_code == 502


@pytest.mark.parametrize("rota", ["recebimentos", "fechamento"])
def test_rota_cache_corrompido_retorna_500(client, tmp_path, rota):
    caminho = persistencia.caminho_dia(EMPRESA, DIA, tmp_path)
    caminho.parent.mkdir()
    caminho.write_text("segredo arquivo corrompido", encoding="utf-8")
    params = (
        {"unidade": EMPRESA, "dia": DIA}
        if rota == "recebimentos"
        else {"unidade": EMPRESA, "inicio": DIA, "fim": DIA}
    )
    response = client.get(f"/api/v1/cash-audit/{rota}", params=params)
    assert response.status_code == 500 and "segredo" not in response.text


async def test_job_isola_fechamento_e_persiste_demais_etapas(monkeypatch, tmp_path, caplog):
    caplog.set_level("INFO", logger=noturno.logger.name)
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(noturno, "carregar_unidades", lambda: {EMPRESA: unidade()})

    async def falhar(*args):
        raise RuntimeError("segredo CPF URL")

    async def recebimento(empresa, dia, conf, *, timeout):
        return RecebimentoAdquirente(adquirente=conf.nome, situacao="pendente")

    monkeypatch.setattr(noturno, "auditar_unidade", falhar)
    monkeypatch.setattr(noturno, "conciliar_adquirente", recebimento)
    assert await noturno.executar(DIA, diretorio=tmp_path) == 1
    registro = persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path)
    assert registro.erro_fechamento == "falha ao consultar fechamento"
    assert len(registro.recebimentos.adquirentes) == 2
    estado = saude_robo.carregar_estado()
    assert estado.unidades_com_falha == (EMPRESA,)
    assert "unidades_com_falha=1 unidades_com_aviso=0" in caplog.text
    assert "segredo" not in registro.model_dump_json() + caplog.text


async def test_credencial_invalida_bloqueia_publicacao_normal(monkeypatch, tmp_path, caplog):
    caplog.set_level("INFO", logger=noturno.logger.name)
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(noturno, "carregar_unidades", lambda: {EMPRESA: unidade()})
    resultado = diario()
    recebimentos = resultado.recebimentos.model_copy(
        update={
            "adquirentes": (
                RecebimentoAdquirente(
                    adquirente="PAGBANK",
                    situacao="credencial inválida",
                    erro="credencial inválida",
                ),
            ),
        }
    )

    async def coletar(*args, **kwargs):
        return coleta_corrente(resultado.model_copy(update={"recebimentos": recebimentos}), args[2])

    monkeypatch.setattr(noturno, "coletar_unidade", coletar)
    assert await noturno.executar(DIA, diretorio=tmp_path) == 1
    estado = saude_robo.carregar_estado()
    assert estado.unidades_com_falha == (EMPRESA,)
    assert estado.unidades_com_aviso == (EMPRESA,)
    assert "unidades_com_falha=1 unidades_com_aviso=1" in caplog.text


def test_estado_antigo_sem_unidades_com_aviso_permanece_compativel(tmp_path):
    import json

    estado = saude_robo.EstadoRobo(
        ultima_execucao=datetime(2026, 10, 1, 3),
        dia_processado=DIA,
        unidades_processadas=(EMPRESA,),
        unidades_com_falha=(),
        resumo_gerado=True,
    )
    dados_antigos = estado.model_dump(mode="json")
    dados_antigos.pop("unidades_com_aviso")
    caminho = tmp_path / "estado_robo.json"
    caminho.write_text(json.dumps(dados_antigos), encoding="utf-8")
    carregado = saude_robo.carregar_estado(caminho=caminho)
    assert carregado is not None
    assert carregado.unidades_com_aviso == ()


async def test_job_registra_alertas_sem_duplicar_pendencias(monkeypatch, tmp_path):
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(noturno, "carregar_unidades", lambda: {EMPRESA: unidade()})

    async def coletar(*args, **kwargs):
        return coleta_corrente(diario(), args[2])

    monkeypatch.setattr(noturno, "coletar_unidade", coletar)
    banco = tmp_path / "pendencias.sqlite3"

    assert (
        await noturno.executar(
            DIA,
            diretorio=tmp_path / "snapshots",
            diretorio_pendencias=banco,
        )
        == 0
    )
    assert (
        await noturno.executar(
            DIA,
            diretorio=tmp_path / "snapshots",
            diretorio_pendencias=banco,
        )
        == 0
    )

    from src.modules.cash_reconciliation.adapters import pendencias as repositorio_pendencias

    itens, total = repositorio_pendencias.listar(unidade=EMPRESA, banco=banco)
    assert total == 1
    assert itens[0].tipo == "fechamento_quebra"
    assert [evento.acao for evento in itens[0].historico] == ["criada"]
    estado = saude_robo.carregar_estado()
    assert estado.dia_processado == DIA
    assert estado.unidades_processadas == (EMPRESA,)
    assert estado.resumo_gerado is True
    resumo = saude_robo.caminho_resumo(DIA).read_text(encoding="utf-8")
    assert "LOGOS — Caixas" in resumo
    assert "sobra R$ 2,50" in resumo
    assert "R$ 2,50" in resumo


async def test_job_timeout_nao_impede_recebimentos(monkeypatch, tmp_path):
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(noturno, "carregar_unidades", lambda: {EMPRESA: unidade()})

    async def lento(*args):
        await asyncio.sleep(1)

    async def recebimento(empresa, dia, conf, *, timeout):
        return RecebimentoAdquirente(
            adquirente=conf.nome, situacao="credencial inválida", erro="credencial inválida"
        )

    monkeypatch.setattr(noturno, "auditar_unidade", lento)
    monkeypatch.setattr(noturno, "conciliar_adquirente", recebimento)
    assert await noturno.executar(DIA, diretorio=tmp_path, timeout=0.001) == 1
    registro = persistencia.carregar_dia(EMPRESA, DIA, diretorio=tmp_path)
    assert registro.erro_fechamento == "tempo limite excedido"
    assert len(registro.recebimentos.adquirentes) == 2


async def test_job_falha_de_unidade_nao_interrompe_outra(monkeypatch, tmp_path):
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(
        noturno, "carregar_unidades", lambda: {321: unidade(321), 999: unidade(999)}
    )

    async def coletar(unit, dia, execucao_id, *, timeout):
        if unit.empresa_codigo == 321:
            raise RuntimeError("segredo")
        return diario(999)

    monkeypatch.setattr(noturno, "coletar_unidade", coletar)
    assert await noturno.executar(DIA, diretorio=tmp_path) == 1
    assert persistencia.carregar_dia(999, DIA, diretorio=tmp_path) is not None


async def test_job_sem_somente_leitura_bloqueia(monkeypatch, tmp_path):
    monkeypatch.delenv("WEBPOSTO_WRITES", raising=False)
    assert await noturno.executar(DIA, diretorio=tmp_path) == 1
    assert not list(tmp_path.glob("*.json"))
    estado = saude_robo.carregar_estado()
    assert estado.execucao.audit_status == "AUDIT_FAILED"
    assert estado.execucao.fontes[0].estado == "BLOCKED"
    assert not estado.resumo_gerado


def test_dia_cli_formato_estrito():
    import argparse

    assert noturno._dia("2026-10-01") == DIA
    for valor in ("invalido", "20261001", "2026-02-30"):
        with pytest.raises(argparse.ArgumentTypeError):
            noturno._dia(valor)


def test_envio_telegram_nao_duplica_data(monkeypatch, tmp_path, contexto_envio):
    monkeypatch.setenv("TELEGRAM_ENVIO", "1")
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "token-sintetico")
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "123")
    monkeypatch.setattr(composicao_noturna, "_coletar_dados_telegram", lambda *args: {})
    monkeypatch.setattr(
        composicao_noturna,
        "_mensagens_telegram",
        lambda dia, dados=None: ["resumo sintético"],
    )
    monkeypatch.setattr(
        composicao_noturna,
        "_criar_planilha_telegram",
        lambda dia, dados: tmp_path / "resumo.xlsx",
    )
    chamadas = []

    def enviar(texto, *, token, chat_ids, parse_mode):
        chamadas.append((texto, token, chat_ids, parse_mode))
        return [ResultadoDestinatario(chat_id="123", sucesso=True, status_http=200)]

    monkeypatch.setattr(composicao_noturna, "enviar_mensagem", enviar)
    monkeypatch.setattr(
        composicao_noturna,
        "enviar_documento",
        lambda *args, **kwargs: [
            ResultadoDestinatario(chat_id="123", sucesso=True, status_http=200)
        ],
    )
    registro = tmp_path / "envios_telegram.json"
    for _ in range(2):
        composicao_noturna._tentar_envio_telegram(
            DIA,
            dia_manual=True,
            enviar_manual=True,
            reenviar=False,
            caminho_registro=registro,
            **contexto_envio,
        )

    assert len(chamadas) == 1
    assert "2026-10-01" in registro.read_text(encoding="utf-8")


def test_envio_telegram_desligado_nao_chama_api(monkeypatch, tmp_path, contexto_envio):
    monkeypatch.setenv("TELEGRAM_ENVIO", "0")
    monkeypatch.setattr(
        composicao_noturna,
        "enviar_mensagem",
        lambda *args, **kwargs: pytest.fail("Telegram não deve ser chamado."),
    )

    composicao_noturna._tentar_envio_telegram(
        DIA,
        dia_manual=True,
        enviar_manual=True,
        reenviar=False,
        caminho_registro=tmp_path / "envios.json",
        **contexto_envio,
    )


def test_envio_telegram_401_nao_registra_token_no_log(monkeypatch, tmp_path, caplog, contexto_envio):
    monkeypatch.setenv("TELEGRAM_ENVIO", "1")
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "token-sintetico-secreto")
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "123")
    monkeypatch.setattr(composicao_noturna, "_coletar_dados_telegram", lambda *args: {})
    monkeypatch.setattr(
        composicao_noturna,
        "_mensagens_telegram",
        lambda dia, dados=None: ["resumo sintético"],
    )
    monkeypatch.setattr(
        composicao_noturna,
        "_criar_planilha_telegram",
        lambda dia, dados: tmp_path / "resumo.xlsx",
    )
    with httpx.Client(
        transport=httpx.MockTransport(
            lambda request: httpx.Response(401, json={"ok": False, "description": "unauthorized"})
        )
    ) as cliente:
        monkeypatch.setattr(
            composicao_noturna,
            "enviar_mensagem",
            lambda texto, *, token, chat_ids, parse_mode: telegram_enviar_mensagem(
                texto,
                token=token,
                chat_ids=chat_ids,
                cliente=cliente,
                parse_mode=parse_mode,
            ),
        )
        monkeypatch.setattr(
            composicao_noturna,
            "enviar_documento",
            lambda *args, **kwargs: [
                ResultadoDestinatario(
                    chat_id="123",
                    sucesso=False,
                    status_http=401,
                    tipo_erro="HTTPStatusError",
                )
            ],
        )

        composicao_noturna._tentar_envio_telegram(
            DIA,
            dia_manual=True,
            enviar_manual=True,
            reenviar=False,
            caminho_registro=tmp_path / "envios.json",
            **contexto_envio,
        )

    assert "HTTP 401" in caplog.text
    assert "token-sintetico-secreto" not in caplog.text


def test_falha_na_planilha_nao_invalida_envio_das_mensagens(monkeypatch, tmp_path, caplog, contexto_envio):
    monkeypatch.setenv("TELEGRAM_ENVIO", "1")
    monkeypatch.setenv("TELEGRAM_BOT_TOKEN", "token-sintetico")
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "123")
    monkeypatch.setattr(composicao_noturna, "_coletar_dados_telegram", lambda *args: {})
    monkeypatch.setattr(
        composicao_noturna,
        "_mensagens_telegram",
        lambda dia, dados=None: ["resumo sintético"],
    )
    monkeypatch.setattr(
        composicao_noturna,
        "_criar_planilha_telegram",
        lambda dia, dados: tmp_path / "resumo.xlsx",
    )
    monkeypatch.setattr(
        composicao_noturna,
        "enviar_mensagem",
        lambda *args, **kwargs: [
            ResultadoDestinatario(chat_id="123", sucesso=True, status_http=200)
        ],
    )
    monkeypatch.setattr(
        composicao_noturna,
        "enviar_documento",
        lambda *args, **kwargs: [
            ResultadoDestinatario(
                chat_id="123",
                sucesso=False,
                status_http=500,
                tipo_erro="HTTPStatusError",
            )
        ],
    )
    registro = tmp_path / "envios.json"

    composicao_noturna._tentar_envio_telegram(
        DIA,
        dia_manual=True,
        enviar_manual=True,
        reenviar=False,
        caminho_registro=registro,
        **contexto_envio,
    )

    assert '"situacao": "ok"' in registro.read_text(encoding="utf-8")
    assert "Telegram planilha envio=falhou" in caplog.text


def test_reprocessamento_manual_nao_envia_sem_opcao(monkeypatch, tmp_path, contexto_envio):
    monkeypatch.setenv("TELEGRAM_ENVIO", "1")
    monkeypatch.setattr(
        composicao_noturna,
        "enviar_mensagem",
        lambda *args, **kwargs: pytest.fail("Reprocessamento manual não deve enviar."),
    )

    composicao_noturna._tentar_envio_telegram(
        DIA,
        dia_manual=True,
        enviar_manual=False,
        reenviar=False,
        caminho_registro=tmp_path / "envios.json",
        **contexto_envio,
    )


def test_modulos_nao_importam_uns_aos_outros():
    raiz = Path(__file__).resolve().parents[3] / "src" / "modules"
    dependencias_proibidas = (
        ("cash_reconciliation", "commercial_performance"),
        ("commercial_performance", "cash_reconciliation"),
    )
    violacoes = []
    for origem, destino in dependencias_proibidas:
        for caminho in (raiz / origem).rglob("*.py"):
            arvore = ast.parse(caminho.read_text(encoding="utf-8"))
            for no in ast.walk(arvore):
                if isinstance(no, ast.Import):
                    nomes = [alias.name for alias in no.names]
                elif isinstance(no, ast.ImportFrom):
                    nomes = [no.module or ""]
                else:
                    continue
                if any(
                    nome == f"src.modules.{destino}" or nome.startswith(f"src.modules.{destino}.")
                    for nome in nomes
                ):
                    violacoes.append(f"{caminho}:{no.lineno}")
    assert violacoes == []


@pytest.mark.parametrize("falha", [
    premmia_portal.PremmiaErro("Detalhe privado", etapa="abrir login"),
    premmia_portal.PlaywrightTimeoutError("URL privada"),
    OSError("Arquivo privado"),
    RuntimeError("Resposta privada"),
])
@pytest.mark.parametrize("historico_corrompido", [False, True])
async def test_premmia_qualquer_falha_e_aviso_com_data_sem_alterar_saida(monkeypatch, tmp_path, falha, historico_corrompido):
    u = unidade().model_copy(update={"premmia_ativo":True,
        "adquirentes":(*unidade().adquirentes, AdquirenteConfigurada(nome="PREMMIA",situacao="ativo",modalidades=("premmia",)))})
    monkeypatch.setenv("WEBPOSTO_WRITES","0")
    monkeypatch.setattr(noturno,"carregar_unidades",lambda:{EMPRESA:u})
    monkeypatch.setattr(premmia_portal,"DIRETORIO",tmp_path/"premmia")
    def dias_coletados(*a):
        if historico_corrompido:
            from zipfile import BadZipFile
            raise BadZipFile("Arquivo privado")
        return {date(2026,9,30)}
    monkeypatch.setattr(premmia_portal,"dias_coletados",dias_coletados)
    chamadas=[]
    async def coletar(unit, *, ate):
        chamadas.append((unit.empresa_codigo,ate))
        raise falha
    async def recebimento(empresa,dia,config,*,timeout,aviso_premmia=None):
        return RecebimentoAdquirente(adquirente=config.nome,situacao="pendente" if aviso_premmia else "ok",erro=aviso_premmia)
    async def auditar(*a): return diario().fechamento
    monkeypatch.setattr(premmia_portal,"coletar",coletar)
    monkeypatch.setattr(noturno,"auditar_unidade",auditar)
    monkeypatch.setattr(noturno,"conciliar_adquirente",recebimento)
    monkeypatch.setattr(noturno,"pendencias_do_dia",lambda *a:[])
    monkeypatch.setattr(noturno,"executar_reincidencia_mes",lambda *a,**k:[])
    assert await noturno.executar(DIA,diretorio=tmp_path/"dados") == 0
    assert chamadas == [(EMPRESA,DIA)]
    estado=saude_robo.carregar_estado()
    assert estado.unidades_com_falha == () and estado.unidades_com_aviso == (EMPRESA,)
    r=persistencia.carregar_dia(EMPRESA,DIA,diretorio=tmp_path/"dados")
    erro=r.recebimentos.adquirentes[-1].erro
    assert "coleta não realizada hoje" in erro
    assert ("dados até indisponíveis" if historico_corrompido else "dados até 30/09") in erro
    assert "privad" not in erro
    mensagem="\n".join(gerar_mensagem_curta(DIA,[r],[EMPRESA],set(),{EMPRESA:0}))
    assert f"Premmia: {erro}" in mensagem


def test_login_invisivel_novo_configuracao_e_fallback_unico_sem_credenciais(monkeypatch):
    pw=MagicMock()
    browsers=[MagicMock(),MagicMock()]
    for browser in browsers:
        browser.version="131.0.6778.33"
        browser.new_context.return_value.new_page.return_value.goto.return_value=SimpleNamespace(status=200)
    pw.chromium.launch.side_effect=browsers
    monkeypatch.setattr(premmia_portal,"aguardar",lambda *a,**k:None)
    monkeypatch.setattr(premmia_portal,"verificar_pagina",lambda *a:None)
    browser,page=premmia_portal.abrir_login(pw,unidade(),visivel=False)
    assert pw.chromium.launch.call_args.kwargs == {"headless":False,"args":["--headless=new"],"slow_mo":0}
    contexto=browser.new_context.call_args.kwargs
    assert "Headless" not in contexto["user_agent"]
    assert contexto["locale"]=="pt-BR" and contexto["timezone_id"]=="America/Recife"
    assert contexto["viewport"]=={"width":1366,"height":768}
    assert page.goto.call_args.kwargs["timeout"]==90_000
    page.locator.return_value.fill.assert_not_called()
    browser.close()
    pw.chromium.launch.reset_mock()
    pw.chromium.launch.side_effect=browsers
    browsers[0].new_context.return_value.new_page.return_value.goto.side_effect=premmia_portal.PlaywrightTimeoutError("sintético")
    browser,page=premmia_portal.abrir_login(pw,unidade(),visivel=False)
    assert pw.chromium.launch.call_count==2
    assert pw.chromium.launch.call_args.kwargs["args"]==[]
    browsers[0].close.assert_called()
    page.locator.return_value.fill.assert_not_called()
    browser.close()


def test_fallback_nao_repete_apos_segunda_falha_ou_desafio(monkeypatch):
    pw=MagicMock()
    browser=pw.chromium.launch.return_value
    page=browser.new_context.return_value.new_page.return_value
    page.goto.side_effect=premmia_portal.PlaywrightTimeoutError("sintético")
    with pytest.raises(premmia_portal.PlaywrightTimeoutError):
        premmia_portal.abrir_login(pw,unidade(),visivel=False)
    assert pw.chromium.launch.call_count==2
    pw.chromium.launch.reset_mock()
    page.goto.side_effect=None
    page.goto.return_value=SimpleNamespace(status=200)
    monkeypatch.setattr(premmia_portal,"aguardar",lambda *a,**k:None)
    monkeypatch.setattr(premmia_portal,"verificar_pagina",lambda *a: (_ for _ in ()).throw(premmia_portal.PremmiaVerificacaoHumana()))
    with pytest.raises(premmia_portal.PremmiaVerificacaoHumana):
        premmia_portal.abrir_login(pw,unidade(),visivel=False)
    assert pw.chromium.launch.call_count==1
    page.locator.return_value.fill.assert_not_called()


def test_login_ja_tentado_na_noite_nao_abre_navegador(monkeypatch,tmp_path):
    u=unidade().model_copy(update={"premmia_usuario_env":"USUARIO_SINTETICO","premmia_senha_env":"SENHA_SINTETICA"})
    monkeypatch.setenv("USUARIO_SINTETICO","sintético")
    monkeypatch.setenv("SENHA_SINTETICA","sintético")
    monkeypatch.setattr(premmia_portal,"agora",lambda:datetime(2026,10,11,4,tzinfo=FUSO))
    pasta=tmp_path/"logins"
    pasta.mkdir()
    (pasta/f"{EMPRESA}_2026-10-11.json").write_text("{}",encoding="utf-8")
    with pytest.raises(premmia_portal.PremmiaErro,match="já tentado"):
        premmia_portal.coletar_sync(u,[DIA],diretorio=tmp_path)


def test_planilha_telegram_todas_abas_atualizadas_e_valores_monetarios(monkeypatch,tmp_path):
    from decimal import Decimal
    from openpyxl import load_workbook
    from src.modules.cash_reconciliation.application.classificar_recebimentos import montar_conferencia
    from .test_conferencia_recebimentos import fonte,inv,registro
    r=registro(fonte("PREMMIA",maior=(inv(dinheiro=True),)),unidade=EMPRESA)
    conf=montar_conferencia(EMPRESA,r.dia,r.dia,[r])
    monkeypatch.setattr(composicao_noturna,"conferencia_para_registro",lambda *a:conf)
    dados={"resultados":[r],"unidades":{EMPRESA:unidade()},"nomes_funcionarios":{EMPRESA:{7:"Funcionário sintético"}},
           "abertas":{EMPRESA:0},"vencidas":{EMPRESA:0}}
    arquivo=composicao_noturna._criar_planilha_telegram(r.dia,dados,tmp_path)
    wb = load_workbook(arquivo)
    try:
        assert set(wb.sheetnames)=={"Resumo","Caixas","Despesas","Vales","Ações de recebimento","Correções prontas",
                                   "Detalhes por pagamento",f"Dia {EMPRESA}",f"Formas {EMPRESA}"}
        for aba,colunas in {"Ações de recebimento":("D",),"Correções prontas":("F",),
                            "Detalhes por pagamento":("E","N","R"),f"Dia {EMPRESA}":("B","C","D","E"),
                            f"Formas {EMPRESA}":("D","F","G")}.items():
            for coluna in colunas: assert wb[aba][f"{coluna}2"].number_format=='"R$" #,##0.00'
        assert Decimal(str(wb["Ações de recebimento"]["D2"].value))==Decimal("50")
        assert wb["Ações de recebimento"]["C2"].number_format=="dd/mm/yyyy hh:mm"
    finally:
        wb.close()


def test_sem_fallback_apos_enviar_login_e_sem_repeticao_na_noite(monkeypatch,tmp_path):
    u=unidade().model_copy(update={"premmia_usuario_env":"USUARIO_SINTETICO","premmia_senha_env":"SENHA_SINTETICA"})
    monkeypatch.setenv("USUARIO_SINTETICO","sintético")
    monkeypatch.setenv("SENHA_SINTETICA","sintético")
    pw=MagicMock()
    browser=MagicMock()
    page=MagicMock()
    monkeypatch.setattr(premmia_portal,"sync_playwright",lambda:pw)
    abrir=MagicMock(return_value=(browser,page))
    monkeypatch.setattr(premmia_portal,"abrir_login",abrir)
    monkeypatch.setattr(premmia_portal,"verificar_pagina",lambda *a:None)
    monkeypatch.setattr(premmia_portal,"unico",lambda loc:loc)
    monkeypatch.setattr(premmia_portal,"aguardar",lambda *a,**k: (_ for _ in ()).throw(premmia_portal.PlaywrightTimeoutError("sintético")))
    with pytest.raises(premmia_portal.PremmiaErro) as exc:
        premmia_portal.coletar_sync(u,[DIA],diretorio=tmp_path)
    assert exc.value.etapa=="confirmar login"
    assert abrir.call_count==1
    page.get_by_role.return_value.click.assert_called_once()
    browser.close.assert_called_once()
    with pytest.raises(premmia_portal.PremmiaErro,match="já tentado"):
        premmia_portal.coletar_sync(u,[DIA],diretorio=tmp_path)
    assert abrir.call_count==1
