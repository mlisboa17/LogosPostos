"""Testes unitários sintéticos para identificação de despesas avulsas sem lançamento (Parte A).

Regras da SPRINT DESPESAS_00:
- Dados sintéticos (nenhuma chamada externa a webPosto, Telegram ou Gemini);
- Cobertura de todos os casos: registrado (baixa, movimento, aberto, agrupado),
  ignorado (grupo, aplicação, tarifa), sem registro (PIX pessoa, aprendizado por título/movimento, FGTS)
  e deduplicação por FITID.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from src.modules.cash_reconciliation.adapters.extrato_reader import ExtratoBancario, LancamentoBancario
from src.modules.cash_reconciliation.application.despesas_sem_lancamento import (
    DespesaAvulsa,
    ResumoDespesasAvulsas,
    classificar_tipo_despesa,
    eh_debito_ignorado,
    extrair_favorecido,
    formatar_resumo_terminal,
    gerar_planilha_despesas_xlsx,
    identificar_despesas_sem_lancamento,
)
from src.modules.cash_reconciliation.domain.models import Unidade


@pytest.fixture
def unidade_sintetica() -> Unidade:
    return Unidade.model_validate({
        "empresa_codigo": 74014,
        "nome": "POSTO DOZE FILIAL II",
        "apelido": "Doze",
        "chave_env": "WEBPOSTO_API_KEY_POSTO_DOZE_FILIAL_II",
        "destinos": [],
        "contas_extrato": [
            {"banco": "BB", "conta_sufixo": "4068-X", "cd_conta": 86248},
            {"banco": "PAGBANK", "conta_sufixo": "9048-8", "cd_conta": 86188},
        ],
    })


def _criar_extrato(
    banco: str,
    lancamentos: list[LancamentoBancario],
    conta_sufixo: str = "9048-8",
) -> ExtratoBancario:
    return ExtratoBancario(
        banco=banco,
        agencia="0001",
        conta_sufixo=conta_sufixo,
        inicio=date(2026, 8, 1),
        fim=date(2026, 9, 30),
        sha256="sintetico_sha256",
        lancamentos=tuple(lancamentos),
    )


@pytest.mark.asyncio
async def test_debito_registrado_por_baixa_de_titulo(unidade_sintetica):
    """Débito que coincide com pagamento de título baixado é considerado registrado."""
    l_deb = LancamentoBancario(
        fitid="FITID_01",
        data=date(2026, 8, 15),
        momento=None,
        valor=Decimal("-500.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pagamento de conta - Fornecedor Alpha",
    )
    extrato = _criar_extrato("PAGBANK", [l_deb])

    # Título pago com pagamento de mesmo valor na mesma data
    titulos_mock = [{
        "tituloPagarCodigo": 101,
        "situacao": "Pago",
        "valor": 500.00,
        "vencimento": "2026-08-15",
        "nomeFornecedor": "FORNECEDOR ALPHA LTDA",
        "pagamento": [{
            "valor": 500.00,
            "dataPagamento": "2026-08-15",
        }],
    }]

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 8, 31),
        extratos_mock=[extrato],
        titulos_mock=titulos_mock,
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 0
    assert resumo.debitos_registrados_qtd == 1
    assert resumo.debitos_registrados_total == Decimal("500.00")


@pytest.mark.asyncio
async def test_debito_registrado_por_movimento_de_conta(unidade_sintetica):
    """Débito que já possui lançamento manual em V1/MOVIMENTOS_CONTA é considerado registrado."""
    l_deb = LancamentoBancario(
        fitid="FITID_02",
        data=date(2026, 8, 20),
        momento=None,
        valor=Decimal("-1850.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Jose Pedreiro",
    )
    extrato = _criar_extrato("PAGBANK", [l_deb])

    movimentos_mock = [{
        "codigo": 201,
        "tipo": "Débito",
        "tipoDocumentoOrigem": "MANUAL",
        "dataMovimento": "2026-08-20",
        "valor": 1850.00,
        "descricao": "PEDREIRO",
        "planoContaGerencialCodigo": 66430,
        "centroCustoCodigo": 7295,
    }]

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 8, 31),
        extratos_mock=[extrato],
        titulos_mock=[],
        movimentos_mock=movimentos_mock,
    )

    assert len(resumo.despesas) == 0
    assert resumo.debitos_registrados_qtd == 1
    assert resumo.debitos_registrados_total == Decimal("1850.00")


@pytest.mark.asyncio
async def test_debito_registrado_por_titulo_em_aberto(unidade_sintetica):
    """Débito que bate com título em aberto cujo fornecedor confere no extrato é registrado."""
    l_deb = LancamentoBancario(
        fitid="FITID_03",
        data=date(2026, 8, 25),
        momento=None,
        valor=Decimal("-1200.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pagamento de conta - Beta Energia S/A",
    )
    extrato = _criar_extrato("BB", [l_deb])

    titulos_mock = [{
        "tituloPagarCodigo": 102,
        "situacao": "Aberto",
        "valor": 1200.00,
        "vencimento": "2026-08-25",
        "nomeFornecedor": "BETA ENERGIA S/A",
        "pagamento": [],
    }]

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 8, 31),
        extratos_mock=[extrato],
        titulos_mock=titulos_mock,
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 0
    assert resumo.debitos_registrados_qtd == 1
    assert resumo.debitos_registrados_total == Decimal("1200.00")


@pytest.mark.asyncio
async def test_debito_registrado_por_pagamento_agrupado(unidade_sintetica):
    """Débito que soma múltiplos títulos em aberto do mesmo fornecedor (ex.: frete) é registrado."""
    l_deb = LancamentoBancario(
        fitid="FITID_04",
        data=date(2026, 9, 15),
        momento=None,
        valor=Decimal("-900.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Transportadora Gama",
    )
    extrato = _criar_extrato("PAGBANK", [l_deb])

    # 3 títulos de R$ 300,00 da Transportadora Gama
    titulos_mock = [
        {
            "tituloPagarCodigo": 103,
            "situacao": "Aberto",
            "valor": 300.00,
            "vencimento": "2026-09-10",
            "nomeFornecedor": "TRANSPORTADORA GAMA LTDA",
        },
        {
            "tituloPagarCodigo": 104,
            "situacao": "Aberto",
            "valor": 300.00,
            "vencimento": "2026-09-11",
            "nomeFornecedor": "TRANSPORTADORA GAMA LTDA",
        },
        {
            "tituloPagarCodigo": 105,
            "situacao": "Aberto",
            "valor": 300.00,
            "vencimento": "2026-09-12",
            "nomeFornecedor": "TRANSPORTADORA GAMA LTDA",
        },
    ]

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 9, 1),
        data_ate=date(2026, 9, 30),
        extratos_mock=[extrato],
        titulos_mock=titulos_mock,
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 0
    assert resumo.debitos_registrados_qtd == 1
    assert resumo.debitos_registrados_total == Decimal("900.00")


@pytest.mark.asyncio
async def test_debito_ignorado_transferencia_grupo_aplicacao_e_tarifa(unidade_sintetica):
    """Transferências para o grupo, aplicações automáticas e tarifas são ignoradas."""
    l_grupo = LancamentoBancario(
        fitid="FITID_05",
        data=date(2026, 8, 10),
        momento=None,
        valor=Decimal("-5000.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Loja De Conveniencia",
    )
    l_app = LancamentoBancario(
        fitid="FITID_06",
        data=date(2026, 8, 10),
        momento=None,
        valor=Decimal("-10000.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="BB RENDE FACIL - APLICACAO",
    )
    l_tarifa = LancamentoBancario(
        fitid="FITID_07",
        data=date(2026, 8, 10),
        momento=None,
        valor=Decimal("-9.90"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Cobrança Seguro Cartão Protegido",
    )
    extrato = _criar_extrato("PAGBANK", [l_grupo, l_app, l_tarifa])

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 8, 31),
        extratos_mock=[extrato],
        titulos_mock=[],
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 0
    assert resumo.debitos_ignorados_qtd == 3
    assert resumo.debitos_ignorados_total == Decimal("15009.90")


@pytest.mark.asyncio
async def test_despesa_avulsa_pix_pessoa_sem_registro(unidade_sintetica):
    """Débito PIX para pessoa física sem registro entra como despesa avulsa a classificar."""
    l_pix = LancamentoBancario(
        fitid="FITID_08",
        data=date(2026, 9, 17),
        momento=None,
        valor=Decimal("-650.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Leticia Malta Farias",
    )
    extrato = _criar_extrato("PAGBANK", [l_pix])

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 9, 1),
        data_ate=date(2026, 9, 30),
        extratos_mock=[extrato],
        titulos_mock=[],
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 1
    d = resumo.despesas[0]
    assert d.valor == Decimal("650.00")
    assert d.favorecido == "Leticia Malta Farias"
    assert d.tipo == "PIX pessoa"
    assert d.sugestao_plano_nome == "A CLASSIFICAR"
    assert d.origem_sugestao == "a classificar"


@pytest.mark.asyncio
async def test_despesa_avulsa_com_sugestao_aprendida_de_titulo(unidade_sintetica):
    """Débito avulso de fornecedor com título histórico herda plano de contas e centro de custo."""
    l_deb = LancamentoBancario(
        fitid="FITID_09",
        data=date(2026, 9, 18),
        momento=None,
        valor=Decimal("-25000.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Sy Display E Solutions",
    )
    extrato = _criar_extrato("PAGBANK", [l_deb])

    # Título anterior com outro valor (ex.: R$ 2.500)
    titulos_mock = [{
        "tituloPagarCodigo": 3213475,
        "situacao": "Aberto",
        "valor": 2500.00,
        "vencimento": "2026-10-06",
        "nomeFornecedor": "SY DISPLAY E SOLUTIONS LTDA",
        "fornecedorCodigo": 9991,
        "planoContaGerencialCodigo": 29063,
        "planoContaGerencialDescricao": "A19 - MARKETING DIGITAL E PUBLICIDADE",
        "centroCustoCodigo": 7295,
        "centroCustoDescricao": "PISTA",
    }]

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 9, 1),
        data_ate=date(2026, 9, 30),
        extratos_mock=[extrato],
        titulos_mock=titulos_mock,
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 1
    d = resumo.despesas[0]
    assert d.valor == Decimal("25000.00")
    assert d.sugestao_plano_codigo == 29063
    assert "MARKETING" in d.sugestao_plano_nome
    assert d.sugestao_centro_codigo == 7295
    assert "mesmo fornecedor em título" in d.origem_sugestao


@pytest.mark.asyncio
async def test_despesa_avulsa_imposto_fgts_conferir_folha(unidade_sintetica):
    """Guias de FGTS / CEF Matriz recebem flag 'conferir - folha'."""
    l_fgts = LancamentoBancario(
        fitid="FITID_10",
        data=date(2026, 9, 18),
        momento=None,
        valor=Decimal("-4167.71"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="QR Code Pix enviado - Cef Matriz",
    )
    extrato = _criar_extrato("PAGBANK", [l_fgts])

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 9, 1),
        data_ate=date(2026, 9, 30),
        extratos_mock=[extrato],
        titulos_mock=[],
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 1
    d = resumo.despesas[0]
    assert d.tipo == "imposto"
    assert d.observacao == "conferir - folha"


@pytest.mark.asyncio
async def test_deduplicacao_por_fitid(unidade_sintetica):
    """Transações duplicadas no mesmo ou em outro extrato com mesmo FITID contam só uma vez."""
    l_deb1 = LancamentoBancario(
        fitid="FITID_REPETIDO",
        data=date(2026, 8, 15),
        momento=None,
        valor=Decimal("-350.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Josie Jose da Silva",
    )
    l_deb2 = LancamentoBancario(
        fitid="FITID_REPETIDO",
        data=date(2026, 8, 15),
        momento=None,
        valor=Decimal("-350.00"),
        tipo="DEBITO",
        tipo_original="OUT",
        descricao="Pix enviado - Josie Jose da Silva",
    )
    extrato1 = _criar_extrato("PAGBANK", [l_deb1])
    extrato2 = _criar_extrato("PAGBANK", [l_deb2])

    resumo = await identificar_despesas_sem_lancamento(
        unidade=unidade_sintetica,
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 8, 31),
        extratos_mock=[extrato1, extrato2],
        titulos_mock=[],
        movimentos_mock=[],
    )

    assert len(resumo.despesas) == 1
    assert resumo.total_geral == Decimal("350.00")


def test_geracao_planilha_despesas_xlsx(tmp_path):
    """Valida que a planilha XLSX é gerada com as duas abas e formatação esperada."""
    d1 = DespesaAvulsa(
        data=date(2026, 9, 18),
        banco="PAGBANK",
        conta_sufixo="9048-8",
        valor=Decimal("25000.00"),
        fitid="FITID_X1",
        favorecido="Sy Display E Solutions",
        descricao_extrato="Pix enviado - Sy Display E Solutions",
        tipo="PIX pessoa",
        sugestao_plano_codigo=29063,
        sugestao_plano_nome="A19 - MARKETING DIGITAL",
        sugestao_centro_codigo=7295,
        sugestao_centro_nome="PISTA",
        sugestao_fornecedor_codigo=9991,
        sugestao_fornecedor_nome="SY DISPLAY",
        origem_sugestao="título 3213475",
        observacao="",
    )
    resumo = ResumoDespesasAvulsas(
        unidade_codigo=74014,
        unidade_nome="POSTO DOZE",
        data_de=date(2026, 8, 1),
        data_ate=date(2026, 9, 30),
        despesas=(d1,),
        total_geral=Decimal("25000.00"),
        totais_por_tipo={"PIX pessoa": (1, Decimal("25000.00"))},
        debitos_ignorados_qtd=5,
        debitos_ignorados_total=Decimal("15000.00"),
        debitos_registrados_qtd=10,
        debitos_registrados_total=Decimal("40000.00"),
    )

    caminho = tmp_path / "despesas_teste.xlsx"
    salvo = gerar_planilha_despesas_xlsx(resumo, caminho)

    assert salvo.is_file()
    wb = openpyxl.load_workbook(salvo)
    assert "Despesas Avulsas" in wb.sheetnames
    assert "Resumo por Tipo" in wb.sheetnames

    ws1 = wb["Despesas Avulsas"]
    assert ws1["D4"].value == 25000.0
    assert ws1["F4"].value == "Sy Display E Solutions"

    ws2 = wb["Resumo por Tipo"]
    assert ws2["A4"].value == "PIX pessoa"
    assert ws2["C4"].value == 25000.0
