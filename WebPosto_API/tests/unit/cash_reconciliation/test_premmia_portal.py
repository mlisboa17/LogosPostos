from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

from src.modules.cash_reconciliation.adapters import premmia_portal as portal
from src.modules.cash_reconciliation.application import (
    conciliar_cartoes,
    recebimentos,
    resumo_diario,
)
from src.modules.cash_reconciliation.domain.models import AdquirenteConfigurada, Unidade
from src.modules.cash_reconciliation.domain.recebimentos import (
    RecebimentoAdquirente,
    ResultadoRecebimentos,
)
from src.modules.cash_reconciliation.jobs import noturno

from .test_cartoes import C
from .test_pendencias import criar_diario

DIA = date(2026, 10, 5)


def unidade():
    return Unidade(
        empresa_codigo=321,
        nome="Sintética",
        apelido="Casa Caiada",
        destinos=(),
        chave_env="CHAVE_TESTE",
        premmia_ativo=True,
        premmia_usuario_env="PREMMIA_USER_TESTE",
        premmia_senha_env="PREMMIA_PASS_TESTE",
        adquirentes=(
            AdquirenteConfigurada(nome="PREMMIA", situacao="ativo", modalidades=("premmia",)),
        ),
    )


def registro(codigo="uuid-sintetico", dia=DIA, forma="Desconto"):
    return dict(
        zip(
            portal.COLUNAS,
            (
                "000.000.000-00",
                "Cliente Sintético",
                "Produto sintético",
                "30,00",
                f"{dia.isoformat()}T10:00:00-03:00",
                codigo,
                forma,
                "Processada",
            ),
        )
    )


def original(caminho, registros):
    wb = Workbook()
    wb.active.title = "Conferência"
    wb.active.append(portal.COLUNAS)
    for r in registros:
        wb.active.append([r[c] for c in portal.COLUNAS])
    wb.save(caminho)
    wb.close()


def test_mensal_dedup_dias_vazios_e_resultado_sem_dados_cliente(tmp_path):
    u = unidade()
    portal.acrescentar_coleta(u, [registro(), registro()], [DIA, date(2026, 10, 6)], tmp_path)
    caminho = portal.caminho_mensal(u, DIA, tmp_path)
    assert caminho.name == "postocasacaiada_10_2026.xlsx"
    antes = caminho.read_bytes()
    portal.acrescentar_coleta(u, [registro()], [DIA], tmp_path)
    assert caminho.read_bytes() == antes
    linhas, coletas = portal.ler_mensal(caminho)
    assert len(linhas) == 1
    assert [c["Quantidade"] for c in coletas] == [1, 0]
    assert portal.buscar_transacoes(u, date(2026, 10, 6), tmp_path) == []
    transacoes = portal.buscar_transacoes(u, DIA, tmp_path)
    assert len(transacoes) == 1
    assert transacoes[0].bandeira == "Desconto" and transacoes[0].valor == Decimal("30")
    assert transacoes[0].nsu == "uuid-sintetico"
    assert "CPF" not in transacoes[0].model_dump_json()
    assert "Cliente Sintético" not in transacoes[0].model_dump_json()
    with pytest.raises(portal.PremmiaDiaNaoColetado):
        portal.buscar_transacoes(u, date(2026, 10, 7), tmp_path)


def test_v33_mensal_preserva_tres_usos_e_deduplica_linha_identica(tmp_path):
    u = unidade()
    registros = []
    for dia, hora in [(date(2026, 10, 4), "14:28:35"), (date(2026, 10, 4), "14:55:41"), (date(2026, 10, 7), "14:17:33")]:
        r = registro("CUPOM777", dia)
        r["Data/Hora da transação"] = f"{dia}T{hora}-03:00"
        r["Valor líquido"] = "7,00"
        registros.append(r)
    dias = [date(2026, 10, 4), date(2026, 10, 7)]
    portal.acrescentar_coleta(u, registros + [registros[0]], dias, tmp_path)
    linhas, coletas = portal.ler_mensal(portal.caminho_mensal(u, dias[0], tmp_path))
    assert len(linhas) == 3
    assert [c["Quantidade"] for c in coletas] == [2, 1]
    ts = [t for d in dias for t in portal.buscar_transacoes(u, d, tmp_path)]
    assert len({t.identificador for t in ts}) == 3
    assert {t.nsu for t in ts} == {"CUPOM777"}
    antes = portal.caminho_mensal(u, dias[0], tmp_path).read_bytes()
    portal.acrescentar_coleta(u, registros, dias, tmp_path)
    assert portal.caminho_mensal(u, dias[0], tmp_path).read_bytes() == antes


def test_v33_integridade_detecta_linha_ausente_em_dia_ja_coletado(tmp_path, caplog):
    u = unidade()
    portal.acrescentar_coleta(u, [registro()], [DIA], tmp_path)
    r = registro()
    r["Data/Hora da transação"] = f"{DIA}T11:00:00-03:00"
    with pytest.raises(portal.PremmiaIntegridadeAviso, match="1 linhas do portal não entraram"):
        portal.acrescentar_coleta(u, [registro(), r], [DIA], tmp_path)
    assert "1 linhas do portal não entraram" in caplog.text


@pytest.mark.parametrize("coluna,valor", [
    ("Valor líquido", "8,00"), ("Forma de Pagamento", "Pix"),
])
def test_v33_mesmo_codigo_e_horario_com_valor_ou_forma_distintos_permanece(tmp_path, coluna, valor):
    primeiro = registro()
    segundo = {**primeiro, coluna: valor}
    portal.acrescentar_coleta(unidade(), [primeiro, segundo], [DIA], tmp_path)
    linhas, _ = portal.ler_mensal(portal.caminho_mensal(unidade(), DIA, tmp_path))
    assert len(linhas) == 2


def test_dias_faltantes_e_virada_mes(tmp_path):
    u = unidade()
    assert portal.dias_faltantes(u, date(2026, 10, 3), tmp_path) == [
        date(2026, 10, d) for d in (1, 2, 3)
    ]
    portal.acrescentar_coleta(u, [], [date(2026, 10, d) for d in range(1, 30)], tmp_path)
    assert portal.dias_faltantes(u, date(2026, 11, 1), tmp_path) == [
        date(2026, 10, 30),
        date(2026, 10, 31),
        date(2026, 11, 1),
    ]
    sem_planilha = tmp_path / "novo"
    assert len(portal.dias_faltantes(u, date(2026, 11, 1), sem_planilha)) == 32


def test_url_preserva_products_sem_campos_de_data():
    base = portal.PAGAMENTOS + "?page=2&products=1%2C2&start=2026-01-01&end=2026-01-02"
    filtrada = portal.url_periodo(base, DIA, DIA)
    query = parse_qs(urlsplit(filtrada).query)
    assert query == {
        "page": ["1"],
        "products": ["1,2"],
        "start": ["2026-10-05 00:00"],
        "end": ["2026-10-05 23:59"],
    }
    with pytest.raises(portal.PremmiaErro):
        portal.url_periodo(portal.PAGAMENTOS, DIA, DIA)


@pytest.mark.parametrize(
    "texto",
    [
        "CAPTCHA",
        "Código enviado por SMS",
        "Autenticação em dois fatores",
        "Aceitar termos",
        "Alterar senha",
        "Senha incorreta",
    ],
)
def test_paradas_obrigatorias_sem_dados_no_erro(texto):
    page = MagicMock()
    page.locator("body").inner_text.return_value = texto
    with pytest.raises(portal.PremmiaVerificacaoHumana, match="verificação humana"):
        portal.verificar_pagina(page)


def test_original_layout_periodo_e_formula_nao_executada(tmp_path):
    arquivo = tmp_path / "original.xlsx"
    r = registro()
    r["Nome"] = '=HYPERLINK("https://example.invalid")'
    original(arquivo, [registro()])
    registros = portal.ler_original(arquivo, DIA, DIA)
    assert len(registros) == 1
    with pytest.raises(portal.PremmiaErro, match="período"):
        portal.ler_original(arquivo, date(2026, 10, 6), date(2026, 10, 6))
    portal.acrescentar_coleta(unidade(), [r], [DIA], tmp_path)
    wb = load_workbook(portal.caminho_mensal(unidade(), DIA, tmp_path))
    try:
        assert wb["Transações"]["B2"].data_type == "s"
    finally:
        wb.close()


async def test_reprocessar_desconto_sem_portal_e_sem_par_vermelho(monkeypatch, tmp_path):
    u = unidade()
    portal.acrescentar_coleta(u, [registro()], [DIA], tmp_path)
    monkeypatch.setattr(portal, "DIRETORIO", tmp_path)
    monkeypatch.setattr(conciliar_cartoes, "carregar_unidades", lambda: {321: u})
    cartoes = [
        C(1, "30", "10:00", nsu="uuid-sintetico", adm="PREMMIA DESCONTO CUPOM"),
        C(2, "30", "10:05", nsu="outro-uuid", adm="PREMMIA PIX"),
    ]
    cartoes = [
        c.model_copy(
            update={
                "empresa_codigo": 321,
                "frentista": 7,
                "nome_frentista": "Funcionário Sintético",
            }
        )
        for c in cartoes
    ]

    async def contexto(*args):
        return cartoes, [], set()

    monkeypatch.setattr(conciliar_cartoes, "coletar_contexto_cartoes", contexto)
    monkeypatch.setattr(portal, "coletar_sync", lambda *a, **k: pytest.fail("Portal acessado"))
    resultado = await conciliar_cartoes.conciliar_premmia(321, DIA)
    assert len(resultado.casados) == 1 and resultado.casados[0].cartao.codigo == 1
    assert len(resultado.a_menor) == 1 and resultado.a_menor[0].codigo == 2
    assert not resultado.pares_provaveis
    assert "Cliente Sintético" not in resultado.model_dump_json()
    from src.modules.cash_reconciliation.application.gerar_pendencias import pendencias_do_dia

    diario = criar_diario(DIA).model_copy(
        update={
            "recebimentos": ResultadoRecebimentos(
                empresa_codigo=321,
                dia=DIA,
                adquirentes=(RecebimentoAdquirente.de_resultado(resultado),),
            ),
        }
    )
    pendencia = next(p for p in pendencias_do_dia(diario) if p.tipo == "recebimento_a_menor")
    assert pendencia.severidade == "vermelho" and "Funcionário Sintético (7)" in pendencia.mensagem
    mensagem = "\n".join(
        resumo_diario.gerar_mensagem_curta(
            DIA,
            [diario],
            [321],
            set(),
            {321: 0},
        )
    )
    assert "Premmia: portal R$ 30,00 · webPosto R$ 60,00 · ❌ diferença R$ -30,00 (1 casados)" in mensagem
    assert "A menor: R$ 30,00 · 10:05 — Funcionário Sintético (7)" in mensagem


async def test_coleta_uma_vez_e_reprocessamento_nao_logam(monkeypatch, tmp_path):
    chamadas = []

    def coletar(u, dias, **kwargs):
        chamadas.append(dias)
        portal.acrescentar_coleta(u, [], dias, tmp_path)

    monkeypatch.setattr(portal, "coletar_sync", coletar)
    await portal.coletar(unidade(), ate=date(2026, 10, 3), diretorio=tmp_path)
    await portal.coletar(unidade(), ate=date(2026, 10, 3), diretorio=tmp_path)
    assert len(chamadas) == 1 and len(chamadas[0]) == 3


def test_um_login_fallback_timeout_um_dia_mesma_sessao(monkeypatch, tmp_path):
    monkeypatch.setenv("PREMMIA_USER_TESTE", "usuario-sintetico")
    monkeypatch.setenv("PREMMIA_PASS_TESTE", "senha-sintetica")
    pw = MagicMock()
    monkeypatch.setattr(portal, "sync_playwright", lambda: pw)
    browser = pw.__enter__.return_value.chromium.launch.return_value
    page = browser.new_context.return_value.new_page.return_value
    locators = {}
    botoes = {}
    page.locator.side_effect = lambda seletor: locators.setdefault(seletor, MagicMock())
    page.get_by_role.side_effect = lambda role, name, **kwargs: botoes.setdefault(name, MagicMock())
    page.url = portal.LOGIN

    def goto(url, **kwargs):
        page.url = url + "?products=1%2C2" if url == portal.PAGAMENTOS else url
        return SimpleNamespace(status=200)

    page.goto.side_effect = goto
    monkeypatch.setattr(portal, "verificar_pagina", lambda *a: None)
    monkeypatch.setattr(portal, "aguardar", lambda *a, **k: None)
    monkeypatch.setattr(portal, "aguardar_exportacao", lambda *a: None)
    monkeypatch.setattr(portal, "unico", lambda loc: loc)
    download_context = MagicMock()
    download_context.__enter__.return_value = download_context
    download_context.value.suggested_filename = "relatorio.xlsx"
    download_context.value.save_as.side_effect = lambda path: original(path, [registro()])
    page.expect_download.side_effect = [
        portal.PlaywrightTimeoutError("sintético"),
        download_context,
    ]
    portal.coletar_sync(unidade(), [DIA, date(2026, 10, 6)], diretorio=tmp_path)
    assert pw.__enter__.return_value.chromium.launch.call_count == 1
    page.locator("input[data-testid='Login']").fill.assert_called_once()
    page.get_by_role("button", name="Entrar", exact=True).click.assert_called_once()
    assert page.expect_download.call_count == 2
    assert portal.dias_coletados(portal.caminho_mensal(unidade(), DIA, tmp_path)) == {DIA}
    assert len(list((tmp_path / "originais").glob("*.xlsx"))) == 1


@pytest.mark.parametrize("integridade", [False, True])
async def test_job_captcha_aviso_nao_falha_nem_repete(monkeypatch, tmp_path, integridade):
    u = unidade()
    monkeypatch.setenv("WEBPOSTO_WRITES", "0")
    monkeypatch.setattr(noturno, "carregar_unidades", lambda: {321: u})
    chamadas = []

    async def coletar(*args, **kwargs):
        chamadas.append(1)
        if integridade:
            raise portal.PremmiaIntegridadeAviso("Premmia: 2 linhas do portal não entraram na planilha")
        raise portal.PremmiaVerificacaoHumana()

    monkeypatch.setattr(portal, "coletar", coletar)

    async def auditar(*args):
        return criar_diario(DIA).fechamento.model_copy(update={"inicio": DIA, "fim": DIA})

    monkeypatch.setattr(noturno, "auditar_unidade", auditar)
    monkeypatch.setattr(noturno, "executar_reincidencia_mes", lambda *a, **k: [])
    monkeypatch.setattr(noturno, "pendencias_do_dia", lambda *a: [])
    from src.modules.cash_reconciliation.adapters import saude_robo, persistencia

    monkeypatch.setattr(saude_robo, "DIRETORIO_RESUMOS", tmp_path / "resumos")
    monkeypatch.setattr(saude_robo, "CAMINHO_ESTADO", tmp_path / "estado.json")
    codigo = await noturno.executar(
        DIA, diretorio=tmp_path / "dados", diretorio_pendencias=tmp_path / "db"
    )
    assert codigo == 0 and chamadas == [1]
    item = persistencia.carregar_dia(
        321, DIA, diretorio=tmp_path / "dados"
    ).recebimentos.adquirentes[0]
    assert item.situacao == "pendente"
    assert "coleta não realizada hoje" in item.erro
    assert (
        "Premmia: 2 linhas do portal não entraram na planilha" if integridade else portal.AVISO_HUMANO
    ) in item.erro
    assert saude_robo.carregar_estado().unidades_com_aviso == (321,)
    if integridade:
        from src.modules.cash_reconciliation.application import resumo_diario

        monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {321: u})
        mensagem = "\n".join(resumo_diario.gerar_mensagem_curta(
            DIA, [persistencia.carregar_dia(321, DIA, diretorio=tmp_path / "dados")],
            [321], set(), {321: 0},
        ))
        assert "Premmia: 2 linhas do portal não entraram na planilha" in mensagem
        assert "auditoria com falha" not in mensagem
        assert "conferência não rodou" not in mensagem


async def test_verificacao_humana_bloqueia_novo_login_no_mesmo_dia_alvo(monkeypatch, tmp_path):
    chamadas = []

    def bloquear(*args, **kwargs):
        chamadas.append(1)
        raise portal.PremmiaVerificacaoHumana()

    monkeypatch.setattr(portal, "coletar_sync", bloquear)
    for _ in range(2):
        with pytest.raises(portal.PremmiaVerificacaoHumana):
            await portal.coletar(unidade(), ate=DIA, diretorio=tmp_path)
    assert chamadas == [1]
    with pytest.raises(portal.PremmiaVerificacaoHumana):
        await portal.coletar(unidade(), ate=date(2026, 10, 6), diretorio=tmp_path)
    assert chamadas == [1, 1]


def test_liberar_nova_tentativa_so_do_posto(tmp_path):
    from src.modules.cash_reconciliation.config import carregar_unidades
    from src.modules.cash_reconciliation.domain.tempo import agora

    unidades = carregar_unidades()
    hoje = agora().date().isoformat()
    (tmp_path / "logins").mkdir(exist_ok=True)
    (tmp_path / "avisos").mkdir(exist_ok=True)
    for emp in (74014, 5555):
        (tmp_path / "logins" / f"{emp}_{hoje}.json").write_text("{}")
        (tmp_path / "avisos" / f"{emp}.json").write_text('{"dia_alvo": "2026-10-17"}')
    portal.liberar_nova_tentativa(unidades[74014], tmp_path)
    assert not (tmp_path / "logins" / f"74014_{hoje}.json").exists()
    assert not (tmp_path / "avisos" / "74014.json").exists()
    assert (tmp_path / "logins" / f"5555_{hoje}.json").exists()
    assert (tmp_path / "avisos" / "5555.json").exists()


@pytest.mark.parametrize("valor_env, esperado", [("msedge", "msedge"), (None, None), ("", None)])
def test_abrir_login_passa_canal_do_navegador(monkeypatch, valor_env, esperado):
    from src.modules.cash_reconciliation.config import carregar_unidades

    if valor_env is None:
        monkeypatch.delenv("LOGOS_NAVEGADOR_CANAL", raising=False)
    else:
        monkeypatch.setenv("LOGOS_NAVEGADOR_CANAL", valor_env)
    monkeypatch.setattr(portal, "verificar_pagina", lambda *a: None)
    monkeypatch.setattr(portal, "aguardar", lambda *a, **k: None)
    pw = MagicMock()
    pagina = pw.chromium.launch.return_value.new_context.return_value.new_page.return_value
    pagina.goto.return_value = SimpleNamespace(status=200)
    portal.abrir_login(pw, carregar_unidades()[74014], visivel=True)
    assert pw.chromium.launch.call_args.kwargs["channel"] == esperado
