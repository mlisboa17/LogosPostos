from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.infrastructure.security.jwt_utils import create_access_token
from src.modules.cash_reconciliation.application.relatorio_caixa_dia import gerar_relatorio_caixa_dia
from src.modules.cash_reconciliation.domain.fechamento import (
    DescontoFalta,
    MovimentoDespesa,
    ValeFuncionario,
)
from src.modules.cash_reconciliation.domain.models import Sangria, Unidade
from src.modules.cash_reconciliation.domain.recebimentos import (
    RecebimentoAdquirente,
    ResultadoRecebimentos,
    ResumoCasados,
    ResumoForma,
)
from src.modules.cash_reconciliation.interfaces import http

from .test_http import resultado_sintetico
from .test_resumo_diario import _registro, DIA, EMPRESA


def test_relatorio_diario_sintetico_mascara_dados_e_exibe_setores():
    unidade = Unidade(
        empresa_codigo=11495,
        nome="Posto VIP",
        apelido="VIP",
        chave_env="NAO_USADA",
        destinos=(),
    )
    base = resultado_sintetico()
    primeiro = base.caixas[0].model_copy(update={
        "caixa": base.caixas[0].caixa.model_copy(update={
            "empresa_codigo": 11495,
            "data": date(2026, 10, 10),
            "codigo": 100,
            "centro_custo": 7295,
            "funcionario_codigo": 30,
            "consolidado": False,
        }),
        "despesas": (MovimentoDespesa(
            codigo=1,
            caixa_codigo=100,
            tipo="D",
            valor=Decimal("25.00"),
            plano_conta_codigo=None,
            descricao="Pagamento para Ana CPF 123.456.789-00",
            centro_custo=7295,
        ),),
        "vales": (ValeFuncionario(
            codigo=1,
            empresa_codigo=11495,
            caixa_codigo=100,
            funcionario_codigo=30,
            origem="C",
            valor=Decimal("20.00"),
        ),),
        "vales_falta": (ValeFuncionario(
            codigo=2,
            empresa_codigo=11495,
            caixa_codigo=100,
            funcionario_codigo=30,
            origem="D",
            valor=Decimal("5.00"),
        ),),
        "sangrias": (Sangria(
            codigo=3,
            empresa_codigo=11495,
            caixa_codigo=100,
            conta_codigo=None,
            funcionario_codigo=30,
            valor=Decimal("15.00"),
            momento=datetime(2026, 10, 10, 9),
            alterada=True,
        ),),
        "desconto_falta": DescontoFalta(
            situacao="sem_desconto",
            falta=Decimal("5.00"),
            total_vale=Decimal(0),
            diferenca=Decimal("5.00"),
        ),
    })
    segundo = primeiro.model_copy(update={
        "caixa": primeiro.caixa.model_copy(update={
            "codigo": 101,
            "centro_custo": 24423,
            "funcionario_codigo": 31,
            "consolidado": True,
        }),
        "despesas": (),
        "vales": (),
        "vales_falta": (),
        "sangrias": (),
        "desconto_falta": None,
    })
    segundo = segundo.model_copy(update={
        "despesas": (MovimentoDespesa(
            codigo=2,
            caixa_codigo=101,
            tipo="D",
            valor=Decimal("80.00"),
            plano_conta_codigo="NF-1",
            descricao="Ajuste nota fiscal",
            centro_custo=18713,
        ),),
    })
    fechamento = base.model_copy(update={
        "empresa_codigo": 11495,
        "inicio": date(2026, 10, 10),
        "fim": date(2026, 10, 10),
        "caixas": (primeiro, segundo),
    })
    registro = _registro(alertas=()).model_copy(update={
        "empresa_codigo": 11495,
        "dia": date(2026, 10, 10),
        "fechamento": fechamento,
        "recebimentos": ResultadoRecebimentos(
            empresa_codigo=11495,
            dia=date(2026, 10, 10),
            adquirentes=(),
        ),
    })

    relatorio = gerar_relatorio_caixa_dia(
        registro,
        unidade,
        nomes_funcionarios={30: "Operador Sintético", 31: "Frentista Sintético"},
        situacao_robo="aviso",
    )

    assert "## Caixas" in relatorio
    assert "| Código |" not in relatorio and "Código interno" not in relatorio
    assert "| 100 |" not in relatorio and "| 101 |" not in relatorio
    exportacao = gerar_relatorio_caixa_dia(registro, unidade, incluir_codigo_interno=True)
    assert "Código interno |" in exportacao
    assert "| 100 |" in exportacao and "| 101 |" in exportacao
    assert "Operador Sintético (30)" in relatorio
    assert "Frentista Sintético (31)" in relatorio
    assert "PISTA" in relatorio and "FOOD" in relatorio
    assert "SEM PLANO" in relatorio
    assert "R$ 5,00 | provisória" in relatorio
    assert "vale comum" in relatorio and "falta de caixa" in relatorio
    assert "10/10/2026 09:00" in relatorio
    assert "Ajustes de NF (FECHAMENTO)" in relatorio
    assert "Ajuste nota fiscal" in relatorio
    despesas = relatorio.split("## Despesas", 1)[1].split("### Ajustes de NF", 1)[0]
    assert "Ajuste nota fiscal" not in despesas
    from src.modules.cash_reconciliation.application.relatorio_caixa_dia import _setor
    assert _setor(unidade, 22310) == "CONVENIÊNCIA"
    assert "CPF [omitido]" in relatorio
    assert "123.456.789-00" not in relatorio
    assert "não persistido" in relatorio
    assert "## Recebimentos" in relatorio and "## Pendências do dia" in relatorio


def test_relatorio_mascara_nsu_e_formata_data_hora():
    unidade = Unidade(
        empresa_codigo=EMPRESA,
        nome="Unidade sintética",
        chave_env="NAO_USADA",
        destinos=(),
    )
    registro = _registro(alertas=())
    item = SimpleNamespace(
        grupo=4,
        momento=datetime.fromisoformat("2026-10-09T13:45:00-03:00"),
        valor=Decimal("12.34"),
        de=None,
        para=None,
        responsavel="Nome Sintético (7)",
        caixa=10,
        nsu="NSU-987654321234",
    )
    troca = SimpleNamespace(
        grupo=2,
        momento=datetime.fromisoformat("2026-10-09T14:00:00-03:00"),
        valor=Decimal("12.34"),
        de="Dinheiro",
        para="Cartão",
        responsavel="Nome Sintético (7)",
        caixa=10,
        nsu=None,
    )
    relatorio = gerar_relatorio_caixa_dia(
        registro,
        unidade,
        conferencia=SimpleNamespace(itens=(item, troca)),
    )

    assert "09/10/2026 13:45" in relatorio
    assert "…1234" in relatorio
    assert "NSU-987654321234" not in relatorio
    assert "Dinheiro → Cartão" in relatorio
    assert all(f"### {grupo}:" in relatorio for grupo in ("Lançar", "Alterar", "Apurar", "Conferir à mão"))


def test_status_do_relatorio_reflete_casa_caiada_e_doze_10_outubro():
    from src.modules.cash_reconciliation.domain.cartoes import AlertaCartao, CartaoErp
    from src.modules.cash_reconciliation.domain.fechamento import LinhaModalidade

    def adquirente(nome, qtd_fonte, total_fonte, qtd_erp, total_erp, *, situacao="ok", **kwargs):
        return RecebimentoAdquirente(
            adquirente=nome,
            situacao=situacao,
            evidencias_completas=situacao == "ok",
            formas=(
                ResumoForma(
                    forma="Cartão",
                    portal=ResumoCasados(quantidade=qtd_fonte, total=Decimal(total_fonte)),
                    webposto=ResumoCasados(quantidade=qtd_erp, total=Decimal(total_erp)),
                    diferenca=Decimal(total_fonte) - Decimal(total_erp),
                ),
            ) if situacao == "ok" else (),
            **kwargs,
        )

    def registro_do_dia(codigo, modalidades, fontes):
        dia = date(2026, 10, 10)
        base = _registro(alertas=())
        audit = base.fechamento.caixas[0].model_copy(
            update={
                "caixa": base.fechamento.caixas[0].caixa.model_copy(
                    update={"empresa_codigo": codigo, "data": dia, "consolidado": False}
                ),
                "modalidades": tuple(modalidades),
            }
        )
        fechamento = base.fechamento.model_copy(
            update={"empresa_codigo": codigo, "inicio": dia, "fim": dia, "caixas": (audit,)}
        )
        return base.model_copy(
            update={
                "empresa_codigo": codigo,
                "dia": dia,
                "fechamento": fechamento,
                "recebimentos": ResultadoRecebimentos(
                    empresa_codigo=codigo,
                    dia=dia,
                    adquirentes=tuple(fontes),
                ),
            }
        )

    def linhas(apresentado, apurado, transferencia):
        return (
            LinhaModalidade(
                modalidade="dinheiro",
                rotulo="Dinheiro",
                apresentado=Decimal(apresentado),
                apurado=Decimal(apurado),
                diferenca=Decimal(apresentado) - Decimal(apurado),
            ),
            LinhaModalidade(
                modalidade="transfBanc",
                rotulo="PIX / Transferência",
                apresentado=Decimal(transferencia),
                apurado=Decimal(transferencia),
                diferenca=Decimal(0),
            ),
        )

    casa = registro_do_dia(
        5555,
        linhas("1000.00", "1011.71", "3702.54"),
        (
            adquirente("PAGBANK", 109, "8983.98", 109, "8983.98"),
            adquirente("PREMMIA", 14, "1032.57", 14, "1032.57"),
            adquirente("MAIS_PAGAMENTOS", 0, "0", 0, "0", situacao="pendente"),
        ),
    )
    relatorio_casa = gerar_relatorio_caixa_dia(
        casa,
        Unidade(
            empresa_codigo=5555,
            nome="AP CASA CAIADA",
            apelido="Casa Caiada",
            chave_env="NAO_USADA",
            destinos=(),
        ),
    )
    assert "- ✅ **PagBank (EDI):** 109 · R$ 8.983,98" in relatorio_casa
    assert "- ✅ **Premmia (portal):** 14 · R$ 1.032,57" in relatorio_casa
    assert "Mais Pagamentos × Transferência Bancária:** R$ 3.702,54" in relatorio_casa
    assert "Dinheiro:** apresentado R$ 1.000,00 × apurado R$ 1.011,71 → falta R$ 11,71" in relatorio_casa
    assert "1 diferença, R$ 11,71 e 1 fonte sem conferência" in relatorio_casa
    assert relatorio_casa.index("## Status da conferência") < relatorio_casa.index("## Resumo")

    pix = CartaoErp(
        codigo=800,
        empresa_codigo=74014,
        venda_codigo=None,
        valor=Decimal("800.00"),
        momento=datetime(2026, 10, 10, 16, 0),
        administradora="PIX CNPJ",
        nsu=None,
        nsu_tef=None,
        autorizacao=None,
        tipo_inclusao="Manualmente",
        caixa_codigo=10,
        frentista=299246,
    )
    doze = registro_do_dia(
        74014,
        linhas("21793.00", "21877.80", "23881.93"),
        (
            adquirente(
                "PAGBANK", 495, "36541.86", 494, "36291.79",
                alertas=(AlertaCartao(
                    codigo="CARTAO_COBRANCA_DUPLICADA",
                    severidade="vermelho",
                    mensagem="duplicidade Visa 19:12",
                    valor=Decimal("250.07"),
                ),),
            ).model_copy(update={"sem_fonte": (pix,)}),
            adquirente("PREMMIA", 32, "2570.17", 32, "2570.17"),
            adquirente("MAIS_PAGAMENTOS", 0, "0", 0, "0", situacao="pendente"),
        ),
    )
    relatorio_doze = gerar_relatorio_caixa_dia(
        doze,
        Unidade(
            empresa_codigo=74014,
            nome="POSTO DOZE FILIAL II",
            apelido="Doze",
            chave_env="NAO_USADA",
            destinos=(),
        ),
        mudou_desde_robo=("Novo lançamento PIX CNPJ R$ 800,00 manual 16:00.",),
    )
    assert "495 · R$ 36.541,86 × 494 · R$ 36.291,79 → diferença R$ 250,07" in relatorio_doze
    assert "verificar estorno ao cliente" in relatorio_doze
    assert "Dinheiro:** apresentado R$ 21.793,00 × apurado R$ 21.877,80 → falta R$ 84,80" in relatorio_doze
    assert "PIX CNPJ × extrato:** 1 · R$ 800,00" in relatorio_doze
    assert "Novo lançamento PIX CNPJ R$ 800,00 manual 16:00" in relatorio_doze


def _cliente_api(monkeypatch, *, role: str, company_id: int | None = None):
    unidade = Unidade(
        empresa_codigo=EMPRESA,
        nome="Unidade Sintética",
        chave_env="NAO_USADA",
        destinos=(),
    )
    outra = Unidade(
        empresa_codigo=654,
        nome="Outra Unidade",
        chave_env="NAO_USADA",
        destinos=(),
    )
    monkeypatch.setattr(http, "carregar_unidades", lambda: {EMPRESA: unidade, 654: outra})
    monkeypatch.setattr(http, "carregar_dia", lambda codigo, dia: _registro(alertas=()))
    monkeypatch.setattr(http.saude_robo, "carregar_estado", lambda: None)
    monkeypatch.setattr(http, "conferencia_para_registro", lambda registro: SimpleNamespace(itens=()))
    monkeypatch.setattr(http, "_listar_pendencias_com_prazo", lambda **kwargs: [])
    monkeypatch.setattr(http, "_carregar_prazos_pendencia", lambda: {"vermelho": 3, "laranja": 5})
    app = FastAPI()
    app.include_router(http.router)
    client = TestClient(app)
    client.cookies.set(
        "access_token",
        create_access_token(
            "relatorio@example.invalid",
            extra={"role": role, "company_id": company_id, "token_type": "access"},
        ),
    )
    return client


def test_api_relatorio_gravado_aplica_escopo_do_gerente(monkeypatch):
    client = _cliente_api(monkeypatch, role="gerente", company_id=EMPRESA)

    response = client.get("/api/v1/cash-audit/relatorio-caixa?dia=2026-10-09&gravado=true")
    proibido = client.get(
        "/api/v1/cash-audit/relatorio-caixa",
        params={"dia": "2026-10-09", "unidade": 654, "gravado": True},
    )

    assert response.status_code == 200
    assert len(response.json()["relatorios"]) == 1
    assert "Relatório do caixa" in response.json()["relatorios"][0]["markdown"]
    assert proibido.status_code == 403


def test_api_relatorio_baixa_xlsx_com_abas_por_bloco(monkeypatch):
    client = _cliente_api(monkeypatch, role="diretor")

    response = client.get(
        "/api/v1/cash-audit/relatorio-caixa.xlsx",
        params={"dia": "2026-10-09", "unidade": EMPRESA, "gravado": True},
    )

    workbook = load_workbook(BytesIO(response.content), read_only=True)
    assert response.status_code == 200
    assert {"Resumo", "Caixas", "Recebimentos", "Pendências do dia"}.issubset(workbook.sheetnames)
    workbook.close()


def test_api_relatorio_padrao_reconsulta_fontes(monkeypatch):
    client = _cliente_api(monkeypatch, role="diretor")

    async def consultar(unidade, dia):
        return _registro(alertas=()), (), (), {}, ()

    monkeypatch.setattr(http, "consultar_dados_atuais", consultar)
    response = client.get(
        "/api/v1/cash-audit/relatorio-caixa",
        params={"dia": "2026-10-09", "unidade": EMPRESA},
    )

    assert response.status_code == 200
    assert "Dados consultados em:" in response.json()["relatorios"][0]["markdown"]


def test_api_sincronizar_aplica_limite_de_dois_minutos(monkeypatch):
    client = _cliente_api(monkeypatch, role="diretor")
    monkeypatch.setattr(http, "_ultima_sincronizacao", {})

    async def consultar(unidade, dia):
        return _registro(alertas=()), (), (), {}, ()

    monkeypatch.setattr(http, "consultar_dados_atuais", consultar)
    params = {"dia": "2026-10-09", "unidade": EMPRESA}
    primeira = client.post("/api/v1/cash-audit/sincronizar", params=params)
    segunda = client.post("/api/v1/cash-audit/sincronizar", params=params)

    assert primeira.status_code == 200
    assert "Dados consultados em" in primeira.json()["markdown"]
    assert segunda.status_code == 429