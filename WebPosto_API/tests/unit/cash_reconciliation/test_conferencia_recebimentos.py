from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from src.modules.cash_reconciliation.adapters import pendencias, persistencia, propostas_lancamento
from src.modules.cash_reconciliation.application.classificar_recebimentos import (
    aplicar_reincidencia, classificar_recebimentos, montar_conferencia,
)
from src.modules.cash_reconciliation.application.gerar_pendencias import pendencias_do_dia
from src.modules.cash_reconciliation.application.resumo_diario import abas_conferencia_recebimentos, linhas_acoes_recebimentos
from src.modules.cash_reconciliation.domain.cartoes import Atribuicao, Casamento, Candidato, GrupoProvavel, Investigacao, ParProvavel
from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente, ResultadoRecebimentos, ResumoCasados, ResumoForma
from src.modules.cash_reconciliation.interfaces import http
from src.modules.cash_reconciliation.domain.tempo import FUSO
from src.infrastructure.security.jwt_utils import create_access_token
from .test_cartoes import A, C, T, DIA
from .test_http import client
from .test_pendencias import criar_diario


@pytest.fixture(autouse=True)
def relogio_conferencia(monkeypatch):
    monkeypatch.setattr(http, "agora", lambda: datetime(2026, 10, 11, 12, tzinfo=FUSO))


def inv(valor="50", ident="transacao", *, dinheiro=False):
    t = T(valor, "10:00", nsu=ident, ident=ident).model_copy(update={"adquirente": "PREMMIA", "bandeira": "Pix"})
    candidato = Candidato(abastecimento=A(1, valor, "09:59").model_copy(update={"nome_frentista": "Funcionário sintético"}),
                          pontos=90, motivos=("Dinheiro",), venda_em_dinheiro=True)
    return Investigacao(transacao=t, atribuicao=Atribuicao.ATRIBUIDO if dinheiro else Atribuicao.SEM_ABASTECIMENTO,
                        frentista=7 if dinheiro else None, candidatos=(candidato,) if dinheiro else (),
                        troca_de_forma=dinheiro, nome_frentista="Funcionário sintético" if dinheiro else None)


def fonte(nome, *, maior=(), menor=(), recebido="100", erp="100", **kwargs):
    return RecebimentoAdquirente(
        adquirente=nome, situacao="ok", a_maior=tuple(maior), a_menor=tuple(menor),
        evidencias_completas=True, casados=ResumoCasados(quantidade=1, total=Decimal("100")),
        formas=(ResumoForma(forma="Pix", portal=ResumoCasados(quantidade=1, total=Decimal(recebido)),
                            webposto=ResumoCasados(quantidade=1, total=Decimal(erp)), diferenca=Decimal(recebido)-Decimal(erp)),),
        **kwargs,
    )


def resultado(*fontes):
    return ResultadoRecebimentos(empresa_codigo=74014, dia=DIA, adquirentes=tuple(fontes))


def registro(*fontes, dia=DIA, unidade=74014):
    return criar_diario(dia, unidade).model_copy(update={
        "conferencia_recebimentos": True,
        "recebimentos": resultado(*fontes).model_copy(update={"empresa_codigo": unidade, "dia": dia}),
    })


def test_quatro_grupos_e_troca_dinheiro_sem_segundo_motor():
    c = C(1, "40", "12:00", adm="PREMMIA")
    fonte_p = fonte("PREMMIA", maior=(inv("160"), inv("100", "dinheiro", dinheiro=True)), menor=(c,),
                    sem_fonte=(C(2, "30", "11:00", adm="PIX CNPJ"),))
    itens = classificar_recebimentos(resultado(fonte_p))
    assert {i.grupo for i in itens} == {1, 2, 3, 4}
    dinheiro = next(i for i in itens if i.valor == 100)
    assert dinheiro.grupo == 2 and dinheiro.de == "Dinheiro"
    assert all(i.regra == "RN-CARTAO-010" for i in itens)
    assert "SUSPEITA" in next(i.motivo for i in itens if i.grupo == 3)


def test_legado_sem_comprovante_nao_acusa_ausencia_e_avisa():
    a = fonte("PREMMIA", menor=(C(1, "40", "12:00", adm="PREMMIA"),)).model_copy(update={"evidencias_completas": False})
    conf = montar_conferencia(74014, DIA, DIA, [registro(a)])
    assert conf.itens[0].grupo == 4 and conf.avisos


@pytest.mark.parametrize("dia,origem,destino,valor", [
    (2, "MAIS PAGAMENTOS", "PREMMIA", "50"),
    (4, "PREMMIA", "PAGBANK", "10"),
])
def test_b2_modelo_diretor_compensacao_e_proposta(dia, origem, destino, valor):
    data = date(2026, 10, dia)
    i = inv(valor).model_copy(update={"transacao": inv(valor).transacao.model_copy(update={"adquirente": destino, "momento": datetime(2026, 10, dia, 10, tzinfo=FUSO)})})
    c = C(1, valor, "10:15", adm=origem, dia=data.isoformat())
    if dia == 2:
        fontes = (fonte("MAIS PAGAMENTOS", menor=(c,), recebido="600", erp="650"),
                  fonte("PREMMIA", maior=(i,), recebido="1050", erp="1000"),
                  fonte("PAGBANK", recebido="800", erp="800"))
    else:
        fontes = (fonte("PREMMIA", menor=(c,), recebido="890", erp="900"),
                  fonte("PAGBANK", maior=(i,), recebido="760", erp="750"),
                  fonte("MAIS PAGAMENTOS", recebido="500", erp="500"))
    conf = montar_conferencia(74014, data, data, [registro(*fontes, dia=data)])
    assert conf.diario[0].diferenca == conf.diferenca_periodo == 0
    assert conf.diario[0].compensacao == "compensacao identificada"
    assert len(conf.itens) == 1 and conf.itens[0].grupo == 2
    assert conf.itens[0].de == origem and destino in conf.itens[0].para
    assert conf.itens[0].lancamento == 1
    assert len(conf.diario[0].adquirentes) == 3


def test_b2_compensacao_combinacao_unica_dois_itens():
    a = fonte("PAGBANK", menor=(C(1, "50", "10:15"),), recebido="100", erp="150")
    b = fonte("PREMMIA", maior=(inv("20", "a"), inv("30", "b")), recebido="150", erp="100")
    conf = montar_conferencia(74014, DIA, DIA, [registro(a, b)])
    assert len(conf.itens) == 1 and conf.itens[0].grupo == 2
    assert conf.itens[0].para == "2 pagamentos PREMMIA"
    assert conf.diario[0].diferenca == 0


def test_compensacao_sem_par_nao_forca_casamento():
    a = fonte("PAGBANK", menor=(C(1, "50", "16:00"),), recebido="100", erp="150")
    b = fonte("PREMMIA", maior=(inv("50"),), recebido="150", erp="100")
    conf = montar_conferencia(74014, DIA, DIA, [registro(a, b)])
    assert {i.grupo for i in conf.itens} == {1, 3}
    assert conf.diario[0].compensacao == "compensação a apurar"


def test_cruzada_ambigua_nao_reutiliza_transacao():
    a = fonte("PAGBANK", menor=(C(1, "50", "10:15"), C(2, "50", "10:16")))
    b = fonte("PREMMIA", maior=(inv(),))
    itens = classificar_recebimentos(resultado(a, b))
    assert not any(i.grupo == 2 for i in itens)


def test_par_grupo_e_manual_casado_forma_errada():
    p = ParProvavel(investigacao=inv(), cartao=C(1, "50", "01:00"))
    g = GrupoProvavel(investigacoes=(inv("20", "a"), inv("30", "b")), cartao=C(2, "50", "01:00"))
    t = inv("10").transacao.model_copy(update={"bandeira": "Desconto"})
    casado = Casamento(transacao=t, cartao=C(3, "10", "10:01", adm="PREMMIA CREDITO"), chave="cupom")
    itens = classificar_recebimentos(resultado(fonte("PREMMIA", pares_provaveis=(p,), grupos_provaveis=(g,), detalhes_casados=(casado,))))
    assert len(itens) == 3 and all(i.grupo == 2 for i in itens)
    assert next(i.para for i in itens if i.lancamento == 2) == "2 pagamentos"


def test_reincidencia_unidade_mes_e_funcionario_sem_contar_duas_vezes():
    item = classificar_recebimentos(resultado(fonte("PREMMIA", maior=(inv(dinheiro=True),))))[0]
    itens = [item.model_copy(update={"id": str(n), "valor": v}) for n, v in enumerate(map(Decimal, ("100", "13.90", "50")))]
    resultado_r = aplicar_reincidencia([*itens, itens[0]])
    assert all(i.reincidente and i.reincidencia_total == Decimal("163.90") for i in resultado_r)
    outro = itens[0].model_copy(update={"unidade": 5555})
    assert not aplicar_reincidencia([itens[0], outro])[0].reincidente


def test_manual_sqlite_nao_encontrado_cria_vermelha_e_historia(tmp_path, monkeypatch):
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    a = fonte("PREMMIA", sem_fonte=(C(2, "30", "11:00", adm="PIX CNPJ"),))
    r = registro(a)
    i = classificar_recebimentos(r.recebimentos)[0]
    pendencias.registrar_conferencia(i, "conferido", "Comprovante sintético", "gerente")
    pendencias.registrar_conferencia(i, "nao_encontrado", "Extrato sintético sem crédito", "diretor")
    pendencias.registrar_conferencia(i, "nao_encontrado", "Extrato sintético sem crédito", "diretor")
    eventos = pendencias.eventos_recebimentos(74014)
    conf = montar_conferencia(74014, DIA, DIA, [r], eventos)
    assert conf.itens[0].grupo == 3
    ps, total = pendencias.listar()
    assert total == 1 and ps[0].severidade == "vermelho" and ps[0].historico[0].usuario == "diretor"
    assert ps[0].valor == Decimal("30")


def test_proposta_fase_zero_arquivo_e_decisao(tmp_path, monkeypatch):
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    monkeypatch.setattr(propostas_lancamento, "DIRETORIO", tmp_path / "propostas")
    i = classificar_recebimentos(resultado(fonte("PREMMIA", maior=(inv(dinheiro=True),))))[0]
    path = propostas_lancamento.salvar_proposta(i)
    assert '"status": "proposta"' in path.read_text(encoding="utf-8")
    propostas_lancamento.decidir_proposta(i, "descartada", "Teste sintético", "gerente")
    propostas_lancamento.salvar_proposta(i)
    assert '"status": "descartada"' in path.read_text(encoding="utf-8")
    assert pendencias.eventos_recebimentos(74014)[i.id]["tipo"] == "proposta"
    with pytest.raises(pendencias.TransicaoPendenciaInvalida):
        propostas_lancamento.decidir_proposta(i, "executada pelo gerente", "Teste", "diretor")


def test_pendencias_idempotentes_sem_dobrar_alertas(tmp_path, monkeypatch):
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    monkeypatch.setattr(propostas_lancamento, "DIRETORIO", tmp_path / "propostas")
    r = registro(fonte("PREMMIA", maior=(inv(dinheiro=True),)))
    ps = pendencias_do_dia(r)
    assert pendencias.registrar(ps) > 0
    assert pendencias.registrar(pendencias_do_dia(r)) == 0
    grupos = [p for p in ps if p.tipo == "caixa_alertas_laranja"]
    assert sum((p.valor for p in grupos), Decimal(0)) == Decimal("50")


def test_migracao_historico_manual_sem_referencia_idempotente(tmp_path, monkeypatch):
    import json
    from src.modules.cash_reconciliation.domain.pendencias import NovaPendencia
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    monkeypatch.setattr(propostas_lancamento, "DIRETORIO", tmp_path / "propostas")
    pendencias.registrar([NovaPendencia(
        unidade=74014, dia=DIA, tipo="caixa_alertas_laranja", severidade="laranja", valor=Decimal("30"),
        referencia="caixa:sem caixa:laranja", identidade=f"74014|{DIA}|caixa:sem caixa|laranja",
        mensagem=json.dumps({"alertas": [{"codigo":"CARTAO_MANUAL_SEM_VENDA","valor":"30","mensagem":"Sintético"}]}),
    )])
    r = registro(fonte("PREMMIA", maior=(inv(dinheiro=True),)))
    pendencias.registrar(pendencias_do_dia(r))
    assert pendencias.registrar(pendencias_do_dia(r)) == 0


def test_get_e_post_http_escopo_sem_consultas_externas(client, tmp_path, monkeypatch):
    monkeypatch.setattr(persistencia, "DIRETORIO", tmp_path / "dados")
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    r = registro(fonte("PREMMIA", sem_fonte=(C(1, "30", "11:00", adm="PIX CNPJ").model_copy(update={"empresa_codigo": 321}),)), unidade=321)
    persistencia.salvar_dia(r)
    params = {"unidade": 321, "inicio": DIA.isoformat(), "fim": DIA.isoformat()}
    response = client.get("/api/v1/cash-audit/conferencia-recebimentos", params=params)
    assert response.status_code == 200
    i = response.json()["itens"][0]
    url = f"/api/v1/cash-audit/conferencia-recebimentos/{i['id']}"
    payload = {"unidade": 321, "dia": DIA.isoformat(), "decisao": "nao_encontrado", "observacao": "Teste sintético"}
    assert client.post(url, json=payload).status_code == 200
    assert client.get("/api/v1/cash-audit/conferencia-recebimentos", params=params).json()["itens"][0]["grupo"] == 3
    client.cookies.set("access_token", create_access_token("gerente@example.invalid", extra={"role": "gerente", "company_id": 74014}))
    assert client.get("/api/v1/cash-audit/conferencia-recebimentos", params=params).status_code == 403
    assert client.post(url, json=payload).status_code == 403


def test_planilha_modelo_b2_acoes_e_formula_segura():
    a = fonte("PREMMIA", maior=(inv(dinheiro=True),))
    conf = montar_conferencia(74014, DIA, DIA, [registro(a)])
    wb = Workbook()
    abas_conferencia_recebimentos(wb, [conf])
    assert wb["Dia 74014"].cell(1, 2).value == "Webposto PREMMIA"
    assert wb["Dia 74014"].cell(1, 5).value == "Diferença do dia"
    assert "Conferido (S/N)" in [c.value for c in wb["Ações de recebimento"][1]]
    assert "Nota de certeza" in [c.value for c in wb["Correções prontas"][1]]
    assert all(c.data_type != "f" for s in wb for row in s for c in row)
    assert "Recebimentos: 1 conferidos" in "\n".join(linhas_acoes_recebimentos(conf))
    wb.close()


def test_pix_cnpj_local_descarta_cliente_e_nao_importa_conferencia(tmp_path, monkeypatch):
    from src.modules.cash_reconciliation.adapters import webposto_cartoes
    monkeypatch.setattr(webposto_cartoes, "DIRETORIO_CONFERENCIA", tmp_path)
    wb = Workbook()
    wb.active.title = "PIX CNPJ Doze"
    wb.active.append(["Data", "Hora", "Valor", "Cliente", "CPF", "Código", "Venda", "Tipo", "Caixa", "Operador", "Conferido", "Observação"])
    wb.active.append(["05/10/2026", "01:00", 50, "Cliente proibido", "CPF proibido", 1, None, "Manualmente", 10, "Funcionário sintético (7)", "S", "Anotação sem auditoria"])
    wb.save(tmp_path / "pix_cnpj_doze_10_2026.xlsx")
    wb.close()
    cs = webposto_cartoes.ler_pix_cnpj_gravado(74014, DIA)
    assert len(cs) == 1 and cs[0].frentista == 7 and cs[0].administradora == "PIX CNPJ"
    assert "proibido" not in cs[0].model_dump_json()
    i = classificar_recebimentos(resultado(fonte("PAGBANK", sem_fonte=tuple(cs))))[0]
    assert i.grupo == 4 and i.conferencia is None


def test_grupo_quatro_pagamentos_atribuicao_sugestao_e_caixa():
    from .test_fechamento import caixa
    cx = caixa(caixaCodigo=4460052, empresaCodigo=74014, funcionarioCodigo=158924,
               abertura="2026-10-05T00:00:00-03:00", fechamento="2026-10-05T23:58:00-03:00")
    pagamentos = []
    for n, pontos in enumerate((96, 97, 77, None)):
        i = inv("25", str(n))
        if pontos is not None:
            c = Candidato(abastecimento=A(n+1, "25", "09:59").model_copy(update={"frentista":158924, "nome_frentista":"Funcionário sintético"}),
                          pontos=pontos, motivos=("valor igual", "1 min antes"), venda_em_dinheiro=False)
            i = i.model_copy(update={"candidatos":(c,), "frentista":158924 if pontos >= 80 else None,
                                    "nome_frentista":"Funcionário sintético" if pontos >= 80 else None})
        pagamentos.append(i)
    grupo = GrupoProvavel(cartao=C(1,"100","01:00"), investigacoes=tuple(pagamentos))
    itens = classificar_recebimentos(resultado(fonte("PAGBANK", grupos_provaveis=(grupo,))), caixas=(cx,),
                                      nomes={158924:"Funcionário sintético"})
    item = itens[0]
    assert len(item.pagamentos) == 4 and item.valor == 100
    assert [p.responsaveis[0].codigo for p in item.pagamentos] == [158924]*4
    assert [p.responsaveis[0].nota for p in item.pagamentos] == [96,97,77,0]
    assert "sugestão" in item.pagamentos[2].responsaveis[0].identificacao
    assert item.pagamentos[3].responsaveis[0].caixa == 4460052
    assert all(p.nsu is None or p.nsu.startswith("…") for p in item.pagamentos)
    conf = montar_conferencia(74014, DIA, DIA, [registro(fonte("PAGBANK", grupos_provaveis=(grupo,)))])
    wb = Workbook()
    abas_conferencia_recebimentos(wb,[conf])
    assert wb["Detalhes por pagamento"].max_row == 5
    assert wb["Dia 74014"].cell(2,1).value == DIA
    assert wb["Dia 74014"].cell(2,1).number_format == "dd/mm/yyyy"
    wb.close()


def test_empate_nao_escolhe_pessoa_e_multiplos_caixas_lista_todos():
    from .test_fechamento import caixa
    cs = tuple(Candidato(abastecimento=A(n,"50","09:59").model_copy(update={"frentista":n,"nome_frentista":f"Funcionário {n}"}),
                         pontos=77,motivos=("valor igual",),venda_em_dinheiro=False) for n in (7,8))
    i = inv().model_copy(update={"candidatos":cs})
    item = classificar_recebimentos(resultado(fonte("PREMMIA",maior=(i,))))[0]
    assert item.frentista is None and len(item.responsaveis) == 2
    assert all(r.identificacao == "empate" for r in item.responsaveis)
    assert item.pagamentos[0].abastecimento is None
    caixas = tuple(caixa(caixaCodigo=n,funcionarioCodigo=n) for n in (7,8))
    item = classificar_recebimentos(resultado(fonte("PREMMIA",maior=(inv(),))),caixas=caixas)[0]
    assert len(item.responsaveis) == 2 and {r.codigo for r in item.responsaveis} == {7,8}


@pytest.mark.parametrize("inicio,fim,esperado", [
    ("2026-10-05","2026-10-11",422), ("2026-10-04","2026-10-10",422),
    ("2026-10-02","2026-10-09",422), ("2026-10-03","2026-10-09",200),
])
def test_b3_backend_valida_periodo_antes_de_ler(client, tmp_path, monkeypatch, inicio, fim, esperado):
    monkeypatch.setattr(persistencia,"DIRETORIO",tmp_path/"dados")
    monkeypatch.setattr(pendencias,"BANCO_PENDENCIAS",tmp_path/"banco.sqlite")
    r = client.get("/api/v1/cash-audit/conferencia-recebimentos",params={"unidade":321,"inicio":inicio,"fim":fim})
    assert r.status_code == esperado
    assert not pendencias.BANCO_PENDENCIAS.exists()


def test_manual_legado_ja_contado_nao_duplica_conferidos():
    from src.modules.cash_reconciliation.domain.cartoes import AlertaCartao
    c = C(1,"100","01:00")
    a = fonte("PREMMIA",codigos_conferidos=(1,),alertas=(AlertaCartao(
        codigo="MANUAL",severidade="laranja",mensagem="sintético",valor=c.valor,cartoes=(c,)),))
    r = registro(a)
    item = classificar_recebimentos(r.recebimentos)[0]
    evento = {item.id:{"tipo":"manual","decisao":"conferido","observacao":"Sintético"}}
    conf = montar_conferencia(74014,DIA,DIA,[r],evento)
    assert conf.conferidos.quantidade == 1 and conf.conferidos.total == 100


@pytest.mark.parametrize("hora,esperado", [("10:30:00", 2), ("10:30:01", 3)])
def test_busca_cruzada_limite_exato_30_minutos(hora, esperado):
    a = fonte("PAGBANK", menor=(C(1, "50", hora),))
    b = fonte("PREMMIA", maior=(inv(),))
    itens = classificar_recebimentos(resultado(a, b))
    assert next(i.grupo for i in itens if i.lancamento == 1) == esperado


def test_adquirente_pendente_e_dia_ausente_nao_inventam_totais():
    a = RecebimentoAdquirente(adquirente="FUTURA", situacao="pendente")
    conf = montar_conferencia(74014, DIA, date(2026,10,6), [registro(a)])
    assert conf.diferenca_periodo is None and len(conf.diario) == 2
    assert all(d.diferenca is None for d in conf.diario)
    assert conf.totais_adquirentes[0].recebido is None


def test_http_exportacao_e_auditor_nao_marca(client, tmp_path, monkeypatch):
    monkeypatch.setattr(persistencia, "DIRETORIO", tmp_path / "dados")
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    r = registro(fonte("PREMMIA", sem_fonte=(C(1,"30","11:00").model_copy(update={"empresa_codigo":321}),)), unidade=321)
    persistencia.salvar_dia(r)
    params = {"unidade":321, "inicio":DIA.isoformat(), "fim":DIA.isoformat()}
    response = client.get("/api/v1/cash-audit/conferencia-recebimentos.xlsx", params=params)
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert "Ações de recebimento" in wb.sheetnames
    wb.close()
    assert not pendencias.BANCO_PENDENCIAS.exists()
    i = client.get("/api/v1/cash-audit/conferencia-recebimentos", params=params).json()["itens"][0]
    client.cookies.set("access_token",create_access_token("auditor@example.invalid",extra={"role":"auditor"}))
    assert client.post(f"/api/v1/cash-audit/conferencia-recebimentos/{i['id']}",
                       json={"unidade":321,"dia":DIA.isoformat(),"decisao":"conferido","observacao":"Sintético"}).status_code == 403
    assert client.get("/api/v1/cash-audit/conferencia-recebimentos",
                      params={**params,"fim":"2026-11-05"}).status_code == 422


def test_proposta_falha_espelho_preserva_auditoria_e_recupera(tmp_path, monkeypatch):
    monkeypatch.setattr(pendencias, "BANCO_PENDENCIAS", tmp_path / "pendencias.sqlite3")
    monkeypatch.setattr(propostas_lancamento, "DIRETORIO", tmp_path / "propostas")
    i = classificar_recebimentos(resultado(fonte("PREMMIA", maior=(inv(dinheiro=True),))))[0]
    path = propostas_lancamento.salvar_proposta(i)
    with monkeypatch.context() as ctx:
        def falhar(*args):
            raise OSError("Falha sintética no espelho")
        ctx.setattr(propostas_lancamento.os, "replace", falhar)
        with pytest.raises(pendencias.ErroPersistenciaPendencias):
            propostas_lancamento.decidir_proposta(i, "descartada", "Teste sintético", "gerente")
    assert pendencias.eventos_recebimentos(74014)[i.id]["decisao"] == "descartada"
    propostas_lancamento.salvar_proposta(i)
    assert '"status": "descartada"' in path.read_text(encoding="utf-8")


def test_manual_hora_ficticia_casada_ganha_correcao_hora():
    c = C(1,"50","01:00").model_copy(update={"tipo_inclusao":"Manualmente"})
    p = Casamento(transacao=inv().transacao.model_copy(update={"adquirente":"PAGBANK"}), cartao=c, chave="nsu")
    i = classificar_recebimentos(resultado(fonte("PAGBANK", detalhes_casados=(p,))))[0]
    assert i.grupo == 2 and "01:00:00" in i.de and "10:00:00" in i.para


def test_combinacao_nao_bloqueia_erp_sem_conflito_de_valor():
    a = fonte("PAGBANK", menor=(C(1,"50","10:15"),C(2,"999","10:15")))
    b = fonte("PREMMIA", maior=(inv("20","a"),inv("30","b")))
    itens = classificar_recebimentos(resultado(a,b))
    assert next(i.grupo for i in itens if i.lancamento==1) == 2
    assert next(i.grupo for i in itens if i.lancamento==2) == 3
