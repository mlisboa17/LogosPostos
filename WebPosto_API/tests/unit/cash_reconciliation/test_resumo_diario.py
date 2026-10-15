from datetime import date
from decimal import Decimal
import pytest

from openpyxl import load_workbook

from src.jobs import noturno
from src.modules.cash_reconciliation.application import resumo_diario
from src.modules.cash_reconciliation.domain.execucao import ResultadoDiario
from src.modules.cash_reconciliation.domain.fechamento import (
    Alerta,
    MovimentoDespesa,
    Severidade,
    ValeFuncionario,
)
from src.modules.cash_reconciliation.domain.models import Unidade
from src.modules.cash_reconciliation.domain.recebimentos import ResultadoRecebimentos

from .test_http import resultado_sintetico

DIA = date(2026, 10, 9)
EMPRESA = 321


def test_resumo_v32_quantidade_valor_e_forma():
    from src.modules.cash_reconciliation.application.conciliar_cartoes import conciliar
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import T, DIA as DIA_CARTOES
    from .test_pendencias import criar_diario

    t = T("12", "10:00", nsu="CUPOM001").model_copy(update={
        "adquirente": "PREMMIA", "bandeira": "Desconto",
    })
    resultado = conciliar(74014, "PREMMIA", DIA_CARTOES, [t], [], [], set())
    diario = criar_diario(DIA_CARTOES, 74014).model_copy(update={
        "recebimentos": ResultadoRecebimentos(
            empresa_codigo=74014, dia=DIA_CARTOES,
            adquirentes=(RecebimentoAdquirente.de_resultado(resultado),),
        ),
    })
    mensagem = "\n".join(resumo_diario.gerar_mensagem_curta(
        DIA_CARTOES, [diario], [74014], set(), {74014: 1},
    ))
    assert "Premmia: portal R$ 12,00 · webPosto sem lançamentos · ❌ diferença R$ 12,00 (0 casados)" in mensagem
    assert "A maior: R$ 12,00 · 10:00 · Desconto" in mensagem
    assert "Desconto: portal" not in mensagem
    assert resultado.a_maior[0].transacao.bandeira == "Desconto"
    assert "R$ 0,00" not in mensagem


def _registro(*, alertas=None, despesas=(), vales=(), adquirentes=()):
    base = resultado_sintetico()
    auditoria = base.caixas[0].model_copy(
        update={
            "caixa": base.caixas[0].caixa.model_copy(update={"data": DIA}),
            "alertas": tuple(alertas if alertas is not None else base.caixas[0].alertas),
            "despesas": tuple(despesas),
            "vales": tuple(vales),
        }
    )
    fechamento = base.model_copy(update={"inicio": DIA, "fim": DIA, "caixas": (auditoria,)})
    return ResultadoDiario(
        empresa_codigo=EMPRESA,
        dia=DIA,
        fechamento=fechamento,
        recebimentos=ResultadoRecebimentos(
            empresa_codigo=EMPRESA,
            dia=DIA,
            adquirentes=tuple(adquirentes),
        ),
        proveniencia=base.proveniencia,
        versoes_regras=("FECHAMENTO_V4",),
    )


def _unidade(apelido=None):
    return Unidade(
        empresa_codigo=EMPRESA,
        nome="Unidade sintética",
        apelido=apelido,
        chave_env="NAO_USADA",
        destinos=(),
    )


def test_unidade_sem_problema_exibe_check_e_nome_de_fallback(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade()})
    registro = _registro(alertas=())

    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [registro],
        [EMPRESA],
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
    )[0]

    assert "✅ EM ORDEM" in mensagem
    assert "Unidade sintética · cc 2" in mensagem
    assert "LOGOS — Caixas de sexta-feira, 09/10/2026" in mensagem
    assert "sobra R$ 2,50" in mensagem
    assert "Pendências abertas: 0" in mensagem


def test_falta_em_caixa_nao_consolidado_mostra_aviso(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("Doze")})
    alertas = (
        Alerta(
            codigo="FALTA_SEM_DESCONTO",
            severidade=Severidade.VERMELHO,
            mensagem="falta sintética",
            valor=Decimal("2263.18"),
            referencia=10,
        ),
        Alerta(
            codigo="NAO_CONSOLIDADO",
            severidade=Severidade.LARANJA,
            mensagem="não consolidado",
            referencia=10,
        ),
    )
    registro = _registro(alertas=alertas)
    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [registro],
        [EMPRESA],
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
        {EMPRESA: {3: "Operador Sintético"}},
    )[0]

    assert "⚠️ PRECISA DE VOCÊ" in mensagem
    assert "Doze · cc 2 — 09/10, turno sintético (Operador, 08:00)" in mensagem
    assert "falta R$ 2.263,18, sem desconto, não consolidado" in mensagem
    assert "caixa 10" not in mensagem


def test_despesa_sem_plano_e_limite_de_oito(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("Doze")})
    despesas = tuple(
        MovimentoDespesa(
            codigo=codigo,
            caixa_codigo=10,
            tipo="D",
            valor=Decimal(str(10 - codigo)),
            plano_conta_codigo=None,
            descricao="Lanche equipe",
        )
        for codigo in range(1, 10)
    )
    registro = _registro(alertas=(), despesas=despesas)

    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [registro],
        [EMPRESA],
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
    )[0]

    assert "SEM PLANO ⚠" in mensagem
    assert "+1 outras (R$ 1,00)" in mensagem
    assert mensagem.count("— Lanche equipe — SEM PLANO ⚠ — R$ ") == 8


def test_descricoes_nao_expoem_nome_pessoal(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("Doze")})
    descricoes = (
        "extra de Ronaldo <b>",
        "gasolina Henrique",
        "pagamento do Uber Daniel",
        "NSU 1234567 CPF 123.456.789-00 conta 45678",
    )
    despesas = tuple(
        MovimentoDespesa(
            codigo=codigo,
            caixa_codigo=10,
            tipo="D",
            valor=Decimal("40"),
            plano_conta_codigo="7",
            descricao=descricao,
        )
        for codigo, descricao in enumerate(descricoes, start=1)
    )

    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [_registro(alertas=(), despesas=despesas)],
        [EMPRESA],
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
    )[0]

    assert "extra [omitido]" in mensagem
    assert "gasolina [omitido]" in mensagem
    assert "Uber [omitido]" in mensagem
    assert all(nome not in mensagem for nome in ("Ronaldo", "Henrique", "Daniel"))
    assert all(
        identificador not in mensagem for identificador in ("1234567", "123.456.789-00", "45678")
    )
    assert "&lt;b&gt;" in mensagem
    assert "<b> <b>" not in mensagem


def test_vale_mostra_nome_de_funcionario(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    vale = ValeFuncionario.model_construct(
        codigo=1,
        empresa_codigo=EMPRESA,
        caixa_codigo=10,
        funcionario_codigo=302036,
        origem="D",
        valor=Decimal("300"),
    )
    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [_registro(alertas=(), vales=(vale,))],
        [EMPRESA],
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
        {EMPRESA: {302036: "Funcionário Sintético"}},
    )[0]

    assert "🧾 VALES" in mensagem
    assert "Funcionário Sintético (302036) R$ 300,00 (falta de caixa)" in mensagem


def test_mensagem_longa_e_dividida_abaixo_do_limite(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("Doze")})
    mensagens = resumo_diario.gerar_mensagem_curta(
        DIA,
        [_registro(alertas=())],
        [EMPRESA] * 100,
        set(),
        {EMPRESA: 0},
        {EMPRESA: 0},
    )

    assert len(mensagens) > 1
    assert all(len(mensagem) <= 4000 for mensagem in mensagens)
    assert all(len(mensagem.encode("utf-16-le")) // 2 <= 4000 for mensagem in mensagens)
    assert mensagens[0].startswith("🌙 LOGOS — Caixas de sexta-feira, 09/10/2026")
    assert "Doze · cc 2" in mensagens[0]
    assert "\n".join(mensagens).count("• Doze · cc 2") == 100
    assert all("LOGOS — Caixas" not in mensagem for mensagem in mensagens[1:])


def test_filiais_em_alerta_antes_das_sem_alerta(monkeypatch):
    segunda = 654
    monkeypatch.setattr(
        resumo_diario,
        "carregar_unidades",
        lambda: {EMPRESA: _unidade("Normal"), segunda: _unidade("Alerta")},
    )
    mensagem = resumo_diario.gerar_mensagem_curta(
        DIA,
        [_registro(alertas=())],
        [EMPRESA, segunda],
        set(),
        {EMPRESA: 0, segunda: 0},
        {EMPRESA: 0, segunda: 0},
    )[0]

    assert mensagem.index("Alerta — conferência não rodou") < mensagem.index("Normal · cc 2")


def test_planilha_com_quatro_abas_valores_numericos_e_nome(tmp_path):
    despesa = MovimentoDespesa(
        codigo=1,
        caixa_codigo=10,
        tipo="D",
        valor=Decimal("12.50"),
        plano_conta_codigo="7",
        descricao="Despesa sintética",
        plano_conta_nome="Plano sintético",
    )
    vale = ValeFuncionario(
        codigo=1,
        empresa_codigo=EMPRESA,
        caixa_codigo=10,
        funcionario_codigo=3,
        origem="D",
        valor=Decimal("5.25"),
    )
    registro = _registro(alertas=(), despesas=(despesa,), vales=(vale,))
    dados = {
        "unidades": {EMPRESA: _unidade("Filial Sintética")},
        "resultados": [registro],
        "nomes_funcionarios": {EMPRESA: {3: "Funcionário Planilha Sintético"}},
        "abertas": {EMPRESA: 2},
        "vencidas": {EMPRESA: 1},
    }
    caminho = noturno._criar_planilha_telegram(DIA, dados, tmp_path)
    workbook = load_workbook(caminho, data_only=True)

    assert workbook.sheetnames == ["Resumo", "Caixas", "Despesas", "Vales"]
    assert isinstance(workbook["Resumo"]["F2"].value, (int, float))
    assert "R$" in workbook["Resumo"]["F2"].number_format
    assert isinstance(workbook["Despesas"]["C2"].value, (int, float))
    assert workbook["Caixas"]["C2"].value == "Funcionário Planilha Sintético"
    assert workbook["Caixas"]["B2"].value.startswith("Filial Sintética · cc 2 — caixa de 09/10/2026")
    assert "(10)" not in workbook["Caixas"]["B2"].value
    assert workbook["Caixas"]["I1"].value == "Código interno"
    assert workbook["Caixas"]["I2"].value == 10
    assert workbook["Vales"]["C2"].value == "Funcionário Planilha Sintético"


def _caixa_com_quebra(registro, codigo, centro_custo, valor, *, fechado=True, consolidado=True):
    base = registro.fechamento.caixas[0]
    return base.model_copy(update={
        "caixa": base.caixa.model_copy(update={
            "codigo": codigo, "centro_custo": centro_custo, "turno": "1º turno",
            "fechado": fechado, "consolidado": consolidado,
        }),
        "modalidades": (base.modalidades[0].model_copy(update={"diferenca": Decimal(valor)}),),
        "alertas": (),
    })


def test_vip_dois_caixas_por_setor_aberto_fora_das_faltas_e_ordem(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    registro = _registro(alertas=())
    caixas = (
        _caixa_com_quebra(registro, 1, 7295, "-999", fechado=False),
        _caixa_com_quebra(registro, 2, 21650, "-1910.09", consolidado=False),
        _caixa_com_quebra(registro, 3, 24886, "-20"),
    )
    registro = registro.model_copy(update={"fechamento": registro.fechamento.model_copy(update={"caixas": caixas})})
    texto = "\n".join(resumo_diario.gerar_mensagem_curta(DIA, [registro], [EMPRESA], set(), {EMPRESA: 1}))
    precisa, abertos = texto.split("⏳ AINDA ABERTOS")
    assert "VIP · Loja — 09/10, 1º turno" in precisa
    assert "falta R$ 1.910,09, não consolidado" in precisa
    assert "caixa 1" not in precisa and "R$ 999,00" not in texto
    assert "VIP · Pista — 09/10, 1º turno" in abertos
    assert "aberto desde 08:00" in abertos
    assert precisa.index("R$ 1.910,09") < precisa.index("R$ 20,00")
    assert "cc 24886" in precisa
    assert str(EMPRESA) not in texto


@pytest.mark.parametrize("valor,sem_desconto,consolidado,precisa", [
    ("-10", False, True, False),
    ("-10.01", False, True, True),
    ("-1", True, True, True),
    ("0", False, False, True),
])
def test_limite_existente_e_alertas_pequenos_nunca_em_ordem(monkeypatch, valor, sem_desconto, consolidado, precisa):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    registro = _registro(alertas=())
    caixa = _caixa_com_quebra(registro, 10, 7295, valor, consolidado=consolidado)
    if sem_desconto:
        caixa = caixa.model_copy(update={"alertas": (Alerta(
            codigo="FALTA_SEM_DESCONTO", severidade=Severidade.VERMELHO,
            mensagem="falta", valor=abs(Decimal(valor)), referencia=10,
        ),)})
    registro = registro.model_copy(update={"fechamento": registro.fechamento.model_copy(update={"caixas": (caixa,)})})
    texto = resumo_diario.gerar_resumo_diario(DIA, [registro], [EMPRESA], set(), {EMPRESA: 0})
    assert ("⚠️ PRECISA DE VOCÊ" in texto) is precisa
    assert ("✅ EM ORDEM" in texto) is not precisa
    assert "R$ 0,00" not in texto


def test_principal_25_linhas_sem_perder_continuacao(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    registro = _registro(alertas=())
    caixas = tuple(_caixa_com_quebra(registro, n, 7295, f"-{n + 20}") for n in range(40))
    registro = registro.model_copy(update={"fechamento": registro.fechamento.model_copy(update={"caixas": caixas})})
    texto = resumo_diario.gerar_resumo_diario(DIA, [registro], [EMPRESA], set(), {EMPRESA: 12})
    principal = texto.split("📌 CONTINUAÇÃO DOS CAIXAS")[0].rstrip()
    assert len(principal.splitlines()) <= 25
    assert texto.count("— 09/10, 1º turno") == 40
    assert "(caixa " not in texto
    assert "Pendências abertas: 12" in texto


def test_sem_fontes_pendentes_ou_linhas_zero(monkeypatch):
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente, ResumoCasados

    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    adquirentes = (
        RecebimentoAdquirente(adquirente="PAGBANK", situacao="ok", casados=ResumoCasados(quantidade=0, total=Decimal(0))),
        RecebimentoAdquirente(adquirente="CIELO", situacao="pendente"),
        RecebimentoAdquirente(adquirente="MAIS_PAGAMENTOS", situacao="pendente"),
    )
    texto = resumo_diario.gerar_resumo_diario(DIA, [_registro(alertas=(), adquirentes=adquirentes)], [EMPRESA], set(), {EMPRESA: 0})
    assert all(v not in texto for v in ("R$ 0,00", "CIELO", "MAIS_PAGAMENTOS", "pendente", "PAGBANK"))


def test_duplicidade_e_individual_com_nome_codigo(monkeypatch):
    from src.modules.cash_reconciliation.domain.cartoes import AlertaCartao
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import C

    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    cartao = C(1, "50", "10:00", nsu="12345678").model_copy(update={
        "empresa_codigo": EMPRESA, "momento": resultado_sintetico().caixas[0].caixa.abertura.replace(year=2026, month=10, day=9),
        "frentista": 7, "nome_frentista": "Sintético",
    })
    adquirente = RecebimentoAdquirente(
        adquirente="PAGBANK", situacao="ok", a_menor=(cartao,),
        alertas=(AlertaCartao(codigo="CARTAO_COBRANCA_DUPLICADA", severidade="vermelho",
                             mensagem="Cobrança em duplicidade R$ 100,00", valor=Decimal("100"), cartoes=(cartao,)),),
    )
    texto = resumo_diario.gerar_resumo_diario(DIA, [_registro(alertas=(), adquirentes=(adquirente,))], [EMPRESA], set(), {EMPRESA: 0})
    principal = texto.split("💳 RECEBIMENTOS")[0]
    assert "Cobrança em duplicidade R$ 100,00 — Sintético (7)" in principal
    assert "Sintético (7)" in principal
    assert principal.index("R$ 100,00") < principal.index("R$ 50,00")


def test_divisao_preserva_entidades_e_unicode():
    texto = "🌙 LOGOS\n" + ("😀" * 1500 + " &amp; " + "x" * 2500)
    mensagens = resumo_diario.dividir_resumo(texto)
    assert len(mensagens) > 1
    assert all(len(m.encode("utf-16-le")) // 2 <= 4000 for m in mensagens)
    assert "&amp;" in "\n".join(mensagens)


def _resumo_recebimentos(*adquirentes):
    from .test_cartoes import DIA as DIA_CARTOES
    from .test_pendencias import criar_diario

    diario = criar_diario(DIA_CARTOES, 74014).model_copy(update={
        "recebimentos": ResultadoRecebimentos(
            empresa_codigo=74014, dia=DIA_CARTOES, adquirentes=tuple(adquirentes),
        ),
    })
    return resumo_diario.gerar_resumo_diario(DIA_CARTOES, [diario], [74014], set(), {74014: 0})


@pytest.mark.parametrize("portal,erp", [
    ("MASTERCARD", "MASTERCARD PAGSEGURO"),
    ("MAESTRO", "MAESTRO PAGSEGURO"),
    ("VISA ELECTRON", "VISA ELECTRON PAGSEGURO"),
    ("ELO", "ELO DEBITO PAGSEGURO"),
])
def test_nomes_bandeiras_diferentes_mesmo_casamento_sem_falsa_diferenca(portal, erp):
    from src.modules.cash_reconciliation.application.conciliar_cartoes import conciliar
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import T, C, DIA as DIA_CARTOES

    resultado = conciliar(74014, "PAGBANK", DIA_CARTOES,
                         [T("3427.53", "10:00", nsu="TESTE001").model_copy(update={"bandeira": portal})],
                         [C(1, "3427.53", "10:00", nsu="TESTE001", adm=erp)], [], set())
    adquirente = RecebimentoAdquirente.de_resultado(resultado)
    assert len(adquirente.formas) == 2
    texto = _resumo_recebimentos(adquirente)
    assert "PagBank: portal R$ 3.427,53 · webPosto R$ 3.427,53 · ✅ conferido (1 casados)" in texto
    assert "❌" not in texto and "diferença" not in texto
    assert "portal 0 / webPosto" not in texto
    agrupadas = resumo_diario._agrupar_bandeiras_planilha(list(adquirente.formas))
    assert len(agrupadas) == 1
    assert agrupadas[0].portal.total == agrupadas[0].webposto.total == Decimal("3427.53")
    assert agrupadas[0].diferenca == 0


@pytest.mark.parametrize("erp,esperado", [("150", "-50,00"), ("70", "30,00")])
def test_diferenca_real_casada_por_nsu_respeita_os_dois_totais(erp, esperado):
    from src.modules.cash_reconciliation.application.conciliar_cartoes import conciliar
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import T, C, DIA as DIA_CARTOES

    resultado = conciliar(74014, "PAGBANK", DIA_CARTOES,
                         [T("100", "10:00", nsu="TESTE001")],
                         [C(1, erp, "10:00", nsu="TESTE001")], [], set())
    adquirente = RecebimentoAdquirente.de_resultado(resultado)
    assert resumo_diario._totais_conciliados(adquirente) == (Decimal("100"), Decimal(erp))
    texto = _resumo_recebimentos(adquirente)
    assert f"❌ diferença R$ {esperado}" in texto
    assert "✅ conferido" not in texto


def test_sangrias_agrupadas_por_caixa_e_tipos_ordenados_por_valor(monkeypatch):
    monkeypatch.setattr(resumo_diario, "carregar_unidades", lambda: {EMPRESA: _unidade("VIP")})
    registro = _registro(alertas=())
    caixa = _caixa_com_quebra(registro, 10, 7295, "-182.51", consolidado=False)
    alertas = tuple(Alerta(
        codigo="SANGRIA_ALTERADA", severidade=Severidade.LARANJA,
        mensagem=f"Sangria de R$ {valor}", valor=Decimal(valor), referencia=n,
    ) for n, valor in enumerate(("432", "228", "6", "1"), start=1))
    caixa = caixa.model_copy(update={"alertas": (*alertas, alertas[0])})
    loja = _caixa_com_quebra(registro, 11, 21650, "-900")
    loja = loja.model_copy(update={"alertas": (Alerta(
        codigo="SANGRIA_ALTERADA", severidade=Severidade.LARANJA,
        mensagem="outra sangria", valor=Decimal("10"), referencia=99,
    ),)})
    registro = registro.model_copy(update={"fechamento": registro.fechamento.model_copy(update={"caixas": (caixa, loja)})})
    texto = resumo_diario.gerar_resumo_diario(DIA, [registro], [EMPRESA], set(), {EMPRESA: 0})
    assert texto.count("4 sangrias alteradas após o lançamento (R$ 667,00)") == 1
    assert texto.index("R$ 900,00") < texto.index("R$ 667,00") < texto.index("R$ 182,51") < texto.index("R$ 10,00")
    assert "VIP · Loja — 09/10, 1º turno" in texto
    assert "1 sangria alterada" in texto
    assert "R$ 432,00" not in texto


def test_itens_recebimentos_cinco_com_hora_nsu_mascarado_e_responsavel():
    from src.modules.cash_reconciliation.application.conciliar_cartoes import conciliar
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import T, DIA as DIA_CARTOES

    transacoes = [
        T(str(10+n), f"{10+n:02d}:00", nsu=f"1234567{n}",
          ident=f"t{n}").model_copy(update={"bandeira": "Pix"}) for n in range(7)
    ]
    resultado = conciliar(74014, "PAGBANK", DIA_CARTOES, transacoes, [], [], set())
    texto = _resumo_recebimentos(RecebimentoAdquirente.de_resultado(resultado))
    bloco = texto.split("💳 RECEBIMENTOS")[1]
    assert bloco.count("• A maior:") == 5
    assert "+2 itens — detalhe na planilha" in bloco
    assert "16:00" in bloco and "R$ 16,00" in bloco
    assert "frentista não identificado" in bloco
    assert "NSU 5676" in bloco and "12345676" not in bloco


def test_totais_incluem_pares_e_grupos_sem_conferir_so_por_saldo():
    from src.modules.cash_reconciliation.domain.cartoes import ParProvavel, GrupoProvavel
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente, ResumoCasados
    from .test_conferencia_recebimentos import inv
    from .test_cartoes import C

    fonte = RecebimentoAdquirente(
        adquirente="PREMMIA", situacao="ok", casados=ResumoCasados(quantidade=1, total=Decimal("100")),
        pares_provaveis=(ParProvavel(investigacao=inv("20"), cartao=C(1, "20", "10:01")),),
        grupos_provaveis=(GrupoProvavel(investigacoes=(inv("30", "a"), inv("40", "b")), cartao=C(2, "70", "10:02")),),
    )
    assert resumo_diario._totais_conciliados(fonte) == (Decimal("190"), Decimal("190"))
    assert "⚠️ totais iguais; itens a conferir" in _resumo_recebimentos(fonte)


def test_planilha_agrupa_bandeiras_e_explica_cartao_app():
    from openpyxl import Workbook
    from src.modules.cash_reconciliation.application.classificar_recebimentos import montar_conferencia
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente, ResumoCasados, ResumoForma
    from .test_cartoes import DIA as DIA_CARTOES
    from .test_pendencias import criar_diario

    formas = tuple(ResumoForma(
        forma=nome,
        portal=ResumoCasados(quantidade=qtd_p, total=Decimal(total_p)),
        webposto=ResumoCasados(quantidade=qtd_e, total=Decimal(total_e)),
        diferenca=Decimal(total_p)-Decimal(total_e),
    ) for nome,qtd_p,total_p,qtd_e,total_e in [
        ("MASTERCARD", 1, "100", 0, "0"),
        ("MASTERCARD PAGSEGURO", 0, "0", 1, "100"),
        ("Cartão APP", 0, "0", 2, "50"),
    ])
    fonte = RecebimentoAdquirente(adquirente="PAGBANK", situacao="ok",
        casados=ResumoCasados(quantidade=1, total=Decimal("100")), formas=formas)
    diario = criar_diario(DIA_CARTOES, 74014).model_copy(update={
        "recebimentos": ResultadoRecebimentos(empresa_codigo=74014, dia=DIA_CARTOES, adquirentes=(fonte,)),
    })
    conf = montar_conferencia(74014, DIA_CARTOES, DIA_CARTOES, [diario])
    wb = Workbook()
    resumo_diario.abas_conferencia_recebimentos(wb, [conf])
    linhas = list(wb["Formas 74014"].values)[1:]
    mastercard = next(l for l in linhas if l[1] == "Mastercard")
    assert mastercard[2:] == (1, 100.0, 1, 100.0, 0.0)
    assert any(l[1] == "Cartão APP (crédito sem bandeira discriminada)" for l in linhas)


@pytest.mark.parametrize("estornada", [False, True])
def test_total_duplicidade_nao_soma_casado_duas_vezes_nem_estorno(estornada):
    from src.modules.cash_reconciliation.application.conciliar_cartoes import conciliar
    from src.modules.cash_reconciliation.domain.recebimentos import RecebimentoAdquirente
    from .test_cartoes import T, C, DIA as DIA_CARTOES

    original = T("100", "10:00", nsu="TESTE001", ident="original").model_copy(update={"terminal": "TESTETERMINAL"})
    extra = T("100", "10:01", nsu="TESTE002", ident="extra").model_copy(update={"terminal": "TESTETERMINAL"})
    transacoes = [original, extra]
    if estornada:
        transacoes.append(extra.model_copy(update={"situacao": "estornada"}))
    resultado = conciliar(74014, "PAGBANK", DIA_CARTOES,
                         transacoes, [C(1, "100", "10:00", nsu="TESTE001")], [], set())
    fonte = RecebimentoAdquirente.de_resultado(resultado)
    assert resumo_diario._totais_conciliados(fonte) == (Decimal("100" if estornada else "200"), Decimal("100"))
    bloco = _resumo_recebimentos(fonte).split("💳 RECEBIMENTOS")[1]
    assert "Duplicidade: R$ 100,00" in bloco
    if not estornada:
        assert "❌ diferença R$ 100,00" in bloco
    else:
        assert "✅ conferido" in bloco
